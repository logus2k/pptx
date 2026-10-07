import io, sys
sys.path.insert(0, "/src/backend")
from app.docengine import ops, read, render
data = open("/src/fixtures/decks/simple.pptx", "rb").read()
prs = read.open_deck(data)
sid = read.outline(prs)[2]["slide_id"]
r = ops.add_chart(prs, sid, "column", ["2023", "2024", "2025"], [{"name": "Vendas", "values": [120, 150, 180]}], title="Vendas")
shid = r["shape_id"]
sh = next(x for x in prs.slides.get(sid).shapes if x.shape_id == shid)
print("before", read.chart_data(sh))
ops.edit_chart(prs, sid, shid, kind="column", categories=["2023", "2024", "2025"], series=[{"name": "Vendas", "values": [120, 150, 200]}])
print("same kind", read.chart_data(sh))
ops.edit_chart(prs, sid, shid, kind="bar")
print("bar", read.chart_data(sh))
ops.edit_chart(prs, sid, shid, kind="pie", title="")
print("pie no title", read.chart_data(sh))
buf = io.BytesIO(); prs.save(buf)
again = read.open_deck(buf.getvalue())
sh2 = next(x for x in again.slides.get(sid).shapes if x.shape_id == shid)
print("reopened", read.chart_data(sh2), "shape_id kept", sh2.shape_id == shid)
from pptx import Presentation
import zipfile
z = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
print("chart parts", [n for n in z.namelist() if "chart" in n or "embeddings" in n])
open("/src/tests/scenarios/out/chartkind.pptx", "wb").write(buf.getvalue())
import pypdfium2 as pdfium
pdf = pdfium.PdfDocument(render._convert(buf.getvalue(), 3))
pdf[2].render(scale=1).to_pil().save("/src/tests/scenarios/out/chartkind-3.png")
z = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
x = z.read("ppt/charts/chart1.xml").decode()
i = x.find("externalData")
print("externalData:", x[i - 3:i + 60] if i >= 0 else "MISSING")
print("rels:", z.read("ppt/charts/_rels/chart1.xml.rels").decode()[-200:])
import openpyxl
wb = openpyxl.load_workbook(io.BytesIO(z.read([n for n in z.namelist() if n.startswith("ppt/embeddings/")][0])))
print("workbook rows:", [[c.value for c in r] for r in wb.active.iter_rows()])
