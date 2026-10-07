"""Writes templates/default.pptx: the plain default template used until Banco CTT supplies its own (spec AD-3).

It is python-pptx's own default presentation (its masters, layouts and Office theme) set to 16:9 and with no slides.
Run: .venv/bin/python scripts/make_default_template.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

OUT = Path(__file__).resolve().parents[1] / "templates" / "default.pptx"

prs = Presentation()
old_width = prs.slide_width           # python-pptx's default is 4:3 (10 in x 7.5 in)
prs.slide_width = Emu(12192000)       # 13.333 in: 16:9
prs.slide_height = Emu(6858000)       # 7.5 in, unchanged
# The masters and layouts are positioned for 4:3: stretch every shape's horizontal geometry to the new width, or the
# placeholders would sit in the left three quarters of the slide. Only shapes with a position of their own (an a:xfrm):
# a layout placeholder without one inherits the master's, already stretched, and reading `left` would return that.
factor = prs.slide_width / old_width
for master in prs.slide_masters:
    for shapes in [master.shapes] + [layout.shapes for layout in master.slide_layouts]:
        for shape in shapes:
            if shape._element.xfrm is not None:
                shape.left = Emu(round(shape.left * factor))
                shape.width = Emu(round(shape.width * factor))
OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f"wrote {OUT}: {len(prs.slide_layouts)} layouts, {len(prs.slides)} slides")
