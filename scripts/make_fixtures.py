"""Writes fixtures/decks/: generated decks for the tests (technical design section 12). Real-world decks (from Banco CTT,
public samples, and files saved by PowerPoint, Keynote, Google Slides and LibreOffice: assumption A-12) are added to
fixtures/decks/real/ by hand; these generated ones cover what can be made with python-pptx.
Run: .venv/bin/python scripts/make_fixtures.py"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "fixtures" / "decks"
DEFAULT = ROOT / "templates" / "default.pptx"


def _deck() -> Presentation:
    return Presentation(DEFAULT)


def _picture() -> io.BytesIO:
    img = Image.new("RGB", (640, 400), (230, 236, 245))
    d = ImageDraw.Draw(img)
    d.rectangle([40, 40, 600, 360], outline=(40, 60, 90), width=6)
    d.ellipse([220, 120, 420, 280], fill=(224, 0, 36))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    buf.seek(0)
    return buf


def simple() -> Presentation:
    """Three slides: a title slide, a bulleted slide with notes, a two-column slide; Portuguese text with accents."""
    p = _deck()
    s = p.slides.add_slide(p.slide_layouts[0])
    s.shapes.title.text = "Proposta para o Cliente X"
    s.placeholders[1].text = "Apresentação de capacidades · 2026"
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Agenda"
    body = s.placeholders[1].text_frame
    body.text = "Contexto e objetivos"
    for line, level in (("Solução proposta", 0), ("Arquitetura", 1), ("Calendário", 1), ("Preços e condições", 0)):
        para = body.add_paragraph()
        para.text = line
        para.level = level
    s.notes_slide.notes_text_frame.text = "Falar primeiro do contexto; os preços ficam para o fim."
    s = p.slides.add_slide(p.slide_layouts[3])
    s.shapes.title.text = "Antes e depois"
    s.placeholders[1].text = "Processos manuais\nVários sistemas"
    s.placeholders[2].text = "Fluxo único\nUm só sistema"
    return p


def features() -> Presentation:
    """A table, a picture, a group of shapes, a hidden slide and speaker notes."""
    p = _deck()
    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = "Preços por escalão"
    rows, cols = 4, 3
    table = s.shapes.add_table(rows, cols, Inches(1), Inches(1.6), Inches(11), Inches(3)).table
    for c, h in enumerate(("Escalão", "Utilizadores", "Preço mensal")):
        table.cell(0, c).text = h
    for r, row in enumerate(
        (("Base", "até 50", "€ 400"), ("Pro", "até 500", "€ 2 500"), ("Empresa", "ilimitado", "sob consulta")), 1
    ):
        for c, v in enumerate(row):
            table.cell(r, c).text = v
    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = "Fotografia da linha de triagem"
    s.shapes.add_picture(_picture(), Inches(3), Inches(1.6), width=Inches(6))
    s.notes_slide.notes_text_frame.text = "Imagem de exemplo."
    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = "Componentes agrupados"
    group = s.shapes.add_group_shape()
    for i in range(3):
        box = group.shapes.add_textbox(Inches(1 + 4 * i), Inches(2), Inches(3.5), Inches(1))
        box.text_frame.text = f"Componente {i + 1}"
    hidden = p.slides.add_slide(p.slide_layouts[1])
    hidden.shapes.title.text = "Diapositivo oculto"
    hidden.placeholders[1].text = "Não aparece na apresentação"
    hidden._element.set("show", "0")
    return p


def long(count: int = 200) -> Presentation:
    """A 200-slide deck (the largest the spec allows): rendering time (SP-4)."""
    p = _deck()
    for i in range(count):
        s = p.slides.add_slide(p.slide_layouts[1])
        s.shapes.title.text = f"Diapositivo {i + 1}"
        s.placeholders[1].text = f"Conteúdo do diapositivo {i + 1}\nSegunda linha"
    return p


TIMING = (
    b'<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par>'
    b'<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"/></p:par></p:tnLst></p:timing>'
)
TRANSITION = (
    b'<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="med"><p:fade/></p:transition>'
)
EXT = (
    b'<p:extLst xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
    b'<p:ext uri="{BB962C8B-B14F-4D97-AF65-F5344CB8AC3E}">'
    b'<p14:creationId xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" val="1234567890"/></p:ext></p:extLst>'
)


def fidelity() -> Presentation:
    """What the engine must keep although it cannot edit it (spec section 6): a native chart, an animation timing
    tree, a transition, a slide extension (p14), and a custom XML part. Real files (SmartArt, video, comments) come
    from fixtures/decks/real/."""
    from lxml import etree
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.opc.package import Part
    from pptx.opc.packuri import PackURI

    p = _deck()
    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = "Vendas por trimestre"
    data = CategoryChartData()
    data.categories = ["T1", "T2", "T3", "T4"]
    data.add_series("2026", (12.5, 14.0, 13.2, 16.8))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1.5), Inches(1.6), Inches(10), Inches(5), data)
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Diapositivo animado"
    s.placeholders[1].text = "Primeiro ponto\nSegundo ponto"
    # p:sld's sequence: cSld, clrMapOvr, transition, timing, extLst
    for blob in (TRANSITION, TIMING, EXT):
        s._element.append(etree.fromstring(blob))
    custom = Part(
        PackURI("/customXml/item1.xml"),
        "application/xml",
        p.part.package,
        b'<?xml version="1.0" encoding="UTF-8"?><b:Sources xmlns:b="http://schemas.openxmlformats.org/officeDocument/2006/bibliography"/>',
    )
    p.part.relate_to(custom, RT.CUSTOM_XML)
    return p


def _save(p: Presentation, name: str) -> bytes:
    buf = io.BytesIO()
    p.save(buf)
    data = buf.getvalue()
    (OUT / name).write_bytes(data)
    return data


def _retyped(data: bytes, content_type: str, extra: dict[str, bytes] | None = None) -> bytes:
    """The package with its main part's content type changed (a .potx, a .pptm) and optional extra parts."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            blob = src.read(info.filename)
            if info.filename == "[Content_Types].xml":
                old = b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
                if old not in blob:
                    raise SystemExit("the package's main content type was not found")
                blob = blob.replace(old, content_type.encode())
            dst.writestr(info, blob)
        for name, blob in (extra or {}).items():
            dst.writestr(name, blob)
    return out.getvalue()


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    simple_bytes = _save(simple(), "simple.pptx")
    _save(features(), "features.pptx")
    _save(fidelity(), "fidelity.pptx")
    _save(long(), "long-200.pptx")
    # a template as Office saves one (.potx): the same package, another content type for its main part
    template = Presentation(DEFAULT)
    buf = io.BytesIO()
    template.save(buf)
    (OUT / "template.potx").write_bytes(
        _retyped(buf.getvalue(), "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml")
    )
    # refused uploads
    (OUT / "macro.pptm").write_bytes(
        _retyped(
            simple_bytes, "application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml", {"ppt/vbaProject.bin": b"\0" * 64}
        )
    )
    (OUT / "encrypted.pptx").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 4096)
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("ppt/media/huge.bin", b"\0" * (300 * 1024 * 1024))
    (OUT / "zip-bomb.pptx").write_bytes(bomb.getvalue())
    (OUT / "not-a-deck.pptx").write_bytes(b"%PDF-1.7 this is not a presentation")
    for f in sorted(OUT.iterdir()):
        print(f"{f.name}: {f.stat().st_size} bytes")


if __name__ == "__main__":
    main()
