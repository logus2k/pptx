"""The document engine (M2): the read model and every single-deck operation on the fixture decks. Each operation is
checked for its effect (through the read model, as the assistant sees it) and for fidelity (spec section 6): every
part of the package it did not target is unchanged in content (XML compared in canonical form, binaries by bytes),
and what it targets keeps what it does not change (animations, extensions)."""

from __future__ import annotations

import io
import shutil
import struct
import json
import zipfile
from pathlib import Path

import jsonschema
import pytest
from lxml import etree
from PIL import Image

from pptx.oxml.ns import qn as qn_

from app.docengine import ops, read

from .conftest import REPO

DECKS = REPO / "fixtures" / "decks"
SCHEMA = json.loads((REPO / "contracts" / "slide.schema.json").read_text())


def deck(name: str) -> bytes:
    return (DECKS / name).read_bytes()


def parts(data: bytes) -> dict[str, bytes]:
    """Each part in comparable form: XML canonicalised (attribute order, namespace prefixes), binaries as they are."""
    z = zipfile.ZipFile(io.BytesIO(data))
    out = {}
    for name in z.namelist():
        blob = z.read(name)
        if name.endswith((".xml", ".rels")):
            out[name] = etree.tostring(etree.fromstring(blob), method="c14n")
        else:
            out[name] = blob
    return out


def unchanged_except(before: bytes, after: bytes, touched: set[str], allowed_new: tuple[str, ...] = ()) -> None:
    """Every part of `before` not in `touched` is in `after`, unchanged; `after` adds only parts starting with
    `allowed_new`. [Content_Types].xml and presentation.xml(.rels) may change when slides or media come and go."""
    a, b = parts(before), parts(after)
    for name, blob in a.items():
        if name in touched:
            continue
        assert name in b, f"{name} disappeared"
        assert b[name] == blob, f"{name} changed"
    for name in set(b) - set(a):
        assert name.startswith(allowed_new), f"unexpected new part {name}"


def slide_part(data: bytes, slide_id: int) -> str:
    """The package path of a slide, by its ID."""
    prs = read.open_deck(data)
    return str(prs.slides.get(slide_id).part.partname).lstrip("/")


def first_slide(data: bytes, index: int = 0) -> int:
    return read.outline(read.open_deck(data))[index]["slide_id"]


def get(data: bytes, slide_id: int) -> dict:
    s = read.slide(read.open_deck(data), slide_id)
    jsonschema.validate(s, SCHEMA)
    return s


def text_of(shape: dict) -> list[str]:
    return ["".join(r["text"] for r in p["runs"]) for p in shape.get("paragraphs", [])]


# ── the read model ───────────────────────────────────────────────────
def test_read_model_on_every_fixture_slide_matches_the_contract():
    for name in ("simple.pptx", "features.pptx", "fidelity.pptx"):
        prs = read.open_deck(deck(name))
        for o in read.outline(prs):
            jsonschema.validate(read.slide(prs, o["slide_id"]), SCHEMA)


def test_read_model_marks_what_it_cannot_edit():
    data = deck("fidelity.pptx")
    chart_slide = get(data, first_slide(data, 0))
    chart = next(s for s in chart_slide["shapes"] if s["type"] == "chart")
    assert chart["locked"] is True and "paragraphs" not in chart
    features = deck("features.pptx")
    group = next(s for s in get(features, first_slide(features, 2))["shapes"] if s["type"] == "group")
    assert [text_of(c) for c in group["shapes"]] == [["Componente 1"], ["Componente 2"], ["Componente 3"]]
    assert get(features, first_slide(features, 3))["hidden"] is True


def test_line_breaks_and_levels_are_read():
    data = deck("simple.pptx")
    body = next(s for s in get(data, first_slide(data, 1))["shapes"] if s.get("placeholder", {}).get("idx") == 1)
    assert [(p["level"], "".join(r["text"] for r in p["runs"])) for p in body["paragraphs"]] == [
        (0, "Contexto e objetivos"),
        (0, "Solução proposta"),
        (1, "Arquitetura"),
        (1, "Calendário"),
        (0, "Preços e condições"),
    ]


# ── text ─────────────────────────────────────────────────────────────
def test_update_text_changes_only_that_shape_and_keeps_the_animation():
    data = deck("fidelity.pptx")
    sid = first_slide(data, 1)
    body = next(s for s in get(data, sid)["shapes"] if s.get("placeholder", {}).get("idx") == 1)
    new, result = ops.apply(
        data,
        "update_text",
        {
            "slide_id": sid,
            "shape_id": body["shape_id"],
            "paragraphs": [
                {"runs": [{"text": "Ponto revisto "}, {"text": "importante", "bold": True}]},
                {"level": 1, "runs": [{"text": "linha um\nlinha dois"}]},
            ],
        },
    )
    assert result["slides"] == [sid]
    after = next(s for s in get(new, sid)["shapes"] if s["shape_id"] == body["shape_id"])
    assert text_of(after) == ["Ponto revisto importante", "linha um\nlinha dois"]
    assert after["paragraphs"][0]["runs"][1]["bold"] is True and after["paragraphs"][1]["level"] == 1
    xml = zipfile.ZipFile(io.BytesIO(new)).read(slide_part(new, sid))
    assert b"p:timing" in xml and b"p14:creationId" in xml and b"p:transition" in xml
    unchanged_except(data, new, {slide_part(data, sid)})


def test_format_text_and_the_minimum_size():
    data = deck("simple.pptx")
    sid = first_slide(data, 1)
    body = next(s for s in get(data, sid)["shapes"] if s.get("placeholder", {}).get("idx") == 1)
    with pytest.raises(ops.OpError) as e:
        ops.apply(data, "format_text", {"slide_id": sid, "shape_id": body["shape_id"], "style": {"size_pt": 10}})
    assert e.value.code == "TEXT_TOO_SMALL"
    new, _ = ops.apply(
        data,
        "format_text",
        {
            "slide_id": sid,
            "shape_id": body["shape_id"],
            "style": {"size_pt": 20, "italic": True, "theme_color": "accent1"},
            "range": {"first_paragraph": 0, "last_paragraph": 0},
        },
    )
    p = next(s for s in get(new, sid)["shapes"] if s["shape_id"] == body["shape_id"])["paragraphs"]
    assert p[0]["runs"][0] == {"text": "Contexto e objetivos", "italic": True, "size_pt": 20.0, "theme_color": "accent1"}
    assert "size_pt" not in p[1]["runs"][0]
    unchanged_except(data, new, {slide_part(data, sid)})


def test_locked_and_missing_targets_are_refused_with_a_hint():
    data = deck("fidelity.pptx")
    sid = first_slide(data, 0)
    chart = next(s for s in get(data, sid)["shapes"] if s["type"] == "chart")
    for args, code in (
        ({"slide_id": sid, "shape_id": chart["shape_id"], "paragraphs": [{"runs": [{"text": "x"}]}]}, "LOCKED_ELEMENT"),
        ({"slide_id": sid, "shape_id": 999, "paragraphs": [{"runs": [{"text": "x"}]}]}, "SHAPE_NOT_FOUND"),
        ({"slide_id": 1, "shape_id": 2, "paragraphs": [{"runs": [{"text": "x"}]}]}, "SLIDE_NOT_FOUND"),
    ):
        with pytest.raises(ops.OpError) as e:
            ops.apply(data, "update_text", args)
        assert e.value.code == code and e.value.hint


# ── slides ───────────────────────────────────────────────────────────
def test_add_slide_at_a_position_with_placeholders():
    data = deck("simple.pptx")
    new, result = ops.apply(
        data,
        "add_slide",
        {
            "layout": "Title and Content",
            "position": 1,
            "placeholders": [
                {"type": "title", "text": "Garantia 2026"},
                {"idx": 1, "paragraphs": [{"runs": [{"text": "Cobertura de 2 anos"}]}]},
            ],
        },
    )
    outline = read.outline(read.open_deck(new))
    assert [o["title"] for o in outline] == ["Proposta para o Cliente X", "Garantia 2026", "Agenda", "Antes e depois"]
    assert outline[1]["slide_id"] == result["slides"][0] and outline[1]["summary"] == "Cobertura de 2 anos"
    with pytest.raises(ops.OpError) as e:
        ops.apply(data, "add_slide", {"layout": "Nope", "position": 0})
    assert e.value.code == "LAYOUT_NOT_FOUND" and "Title and Content" in e.value.hint
    unchanged_except(
        data,
        new,
        {"[Content_Types].xml", "ppt/presentation.xml", "ppt/_rels/presentation.xml.rels"},
        allowed_new=("ppt/slides/",),
    )


def test_delete_and_move_slides():
    data = deck("simple.pptx")
    ids = [o["slide_id"] for o in read.outline(read.open_deck(data))]
    moved, _ = ops.apply(data, "move_slide", {"slide_id": ids[2], "position": 0})
    assert [o["slide_id"] for o in read.outline(read.open_deck(moved))] == [ids[2], ids[0], ids[1]]
    gone, _ = ops.apply(data, "delete_slide", {"slide_id": ids[1]})
    assert [o["slide_id"] for o in read.outline(read.open_deck(gone))] == [ids[0], ids[2]]
    assert not any(n.endswith("slide2.xml") for n in zipfile.ZipFile(io.BytesIO(gone)).namelist()), "the slide's part is dropped"


def test_duplicate_slide_copies_chart_parts_and_animations():
    data = deck("fidelity.pptx")
    ids = [o["slide_id"] for o in read.outline(read.open_deck(data))]
    new, result = ops.apply(data, "duplicate_slide", {"slide_id": ids[0]})
    dup = result["slides"][0]
    assert [o["slide_id"] for o in read.outline(read.open_deck(new))] == [ids[0], dup, ids[1]]
    names = zipfile.ZipFile(io.BytesIO(new)).namelist()
    charts = [n for n in names if n.startswith("ppt/charts/chart") and n.endswith(".xml")]
    workbooks = [n for n in names if n.startswith("ppt/embeddings/")]
    assert len(charts) == 2 and len(workbooks) == 2, (charts, workbooks)  # its own chart and workbook, not shared
    prs = read.open_deck(new)
    rels = {r.reltype.rsplit("/", 1)[-1]: r.target_part.partname for r in prs.slides.get(dup).part.rels.values()}
    assert str(rels["chart"]) != str(
        {r.reltype.rsplit("/", 1)[-1]: r.target_part.partname for r in prs.slides.get(ids[0]).part.rels.values()}["chart"]
    )
    chart = next(s for s in get(new, dup)["shapes"] if s["type"] == "chart")
    assert chart["locked"] is True
    # the copied chart still finds its own workbook by the rId written in its XML
    z = zipfile.ZipFile(io.BytesIO(new))
    for c in charts:
        xml = etree.fromstring(z.read(c))
        ext = xml.find("{http://schemas.openxmlformats.org/drawingml/2006/chart}externalData")
        rid = ext.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        rels = etree.fromstring(z.read(c.replace("charts/", "charts/_rels/") + ".rels"))
        target = next(r.get("Target") for r in rels if r.get("Id") == rid)
        assert target.startswith("../embeddings/")
    assert len({etree.fromstring(z.read(c.replace("charts/", "charts/_rels/") + ".rels"))[0].get("Target") for c in charts}) == 2
    new2, r2 = ops.apply(data, "duplicate_slide", {"slide_id": ids[1]})
    xml = zipfile.ZipFile(io.BytesIO(new2)).read(slide_part(new2, r2["slides"][0]))
    assert b"p:timing" in xml and b"p:transition" in xml
    unchanged_except(
        data,
        new,
        {"[Content_Types].xml", "ppt/presentation.xml", "ppt/_rels/presentation.xml.rels"},
        allowed_new=("ppt/slides/", "ppt/charts/", "ppt/embeddings/"),
    )


def test_change_layout_keeps_content_and_reports_what_it_could_not_place():
    data = deck("simple.pptx")
    sid = first_slide(data, 2)  # Two Content: a title and two columns
    new, result = ops.apply(data, "change_layout", {"slide_id": sid, "layout": "Title and Content"})
    nid = result["slides"][0]
    outline = read.outline(read.open_deck(new))
    assert (
        outline[2]["slide_id"] == nid and outline[2]["layout"] == "Title and Content" and outline[2]["title"] == "Antes e depois"
    )
    texts = [t for s in get(new, nid)["shapes"] for t in text_of(s)]
    assert "Processos manuais" in texts  # the first column moved into the body
    assert result["unmatched"] == ["Content Placeholder 3"]  # the second column had no place: reported...
    assert "Fluxo único" in texts and "Um só sistema" in texts  # ...and kept on the slide as a text box, never dropped
    back, _ = ops.apply(data, "change_layout", {"slide_id": first_slide(data, 1), "layout": "Title Only"})
    assert ops.apply(data, "set_notes", {"slide_id": first_slide(data, 1), "text": "x"})  # sanity: notes still writable
    notes = read.slide(read.open_deck(back), read.outline(read.open_deck(back))[1]["slide_id"])["notes"]
    assert notes.startswith("Falar primeiro do contexto")  # the notes travel with the slide


# ── shapes, images, tables, notes ────────────────────────────────────
def test_add_move_delete_shape_and_bounds():
    data = deck("simple.pptx")
    sid = first_slide(data, 1)
    new, r = ops.apply(
        data,
        "add_shape",
        {
            "slide_id": sid,
            "kind": "text_box",
            "box": {"x": 0.6, "y": 0.8, "w": 0.3, "h": 0.1},
            "paragraphs": [{"runs": [{"text": "Nota"}]}],
        },
    )
    added = next(s for s in get(new, sid)["shapes"] if s["shape_id"] == r["shape_id"])
    assert added["type"] == "text_box" and text_of(added) == ["Nota"] and added["box"]["x"] == 0.6
    moved, _ = ops.apply(
        new, "move_resize_shape", {"slide_id": sid, "shape_id": r["shape_id"], "box": {"x": 0.1, "y": 0.1, "w": 0.2, "h": 0.1}}
    )
    assert next(s for s in get(moved, sid)["shapes"] if s["shape_id"] == r["shape_id"])["box"] == {
        "x": 0.1,
        "y": 0.1,
        "w": 0.2,
        "h": 0.1,
    }
    gone, _ = ops.apply(moved, "delete_shape", {"slide_id": sid, "shape_id": r["shape_id"]})
    assert all(s["shape_id"] != r["shape_id"] for s in get(gone, sid)["shapes"])
    with pytest.raises(ops.OpError) as e:
        ops.apply(data, "add_shape", {"slide_id": sid, "kind": "text_box", "box": {"x": 0.9, "y": 0.1, "w": 0.3, "h": 0.1}})
    assert e.value.code == "OUT_OF_BOUNDS"


def _png(w: int, h: int, colour=(0, 128, 255)) -> bytes:
    b = io.BytesIO()
    Image.new("RGB", (w, h), colour).save(b, "PNG")
    return b.getvalue()


def test_insert_image_in_a_box_keeps_proportions_and_webp_becomes_png():
    data = deck("simple.pptx")
    sid = first_slide(data, 1)
    webp = io.BytesIO()
    Image.new("RGB", (400, 200), (10, 200, 30)).save(webp, "WEBP")
    new, r = ops.apply(
        data,
        "insert_image",
        {"slide_id": sid, "image": webp.getvalue(), "target": {"box": {"x": 0.5, "y": 0.2, "w": 0.4, "h": 0.6}}},
    )
    pic = next(s for s in get(new, sid)["shapes"] if s["shape_id"] == r["shape_id"])
    W, H = 12192000, 6858000
    assert pic["type"] == "picture"
    assert abs((pic["emu"]["w"] / pic["emu"]["h"]) - 2.0) < 0.01  # the image's 2:1, not the box's
    assert pic["emu"]["w"] == round(0.4 * W) and pic["emu"]["y"] > round(0.2 * H)  # full width, centred vertically
    media = [n for n in zipfile.ZipFile(io.BytesIO(new)).namelist() if n.startswith("ppt/media/")]
    assert any(n.endswith(".png") for n in media)
    with pytest.raises(ops.OpError) as e:
        ops.apply(
            data,
            "insert_image",
            {"slide_id": sid, "image": b"not an image", "target": {"box": {"x": 0, "y": 0, "w": 0.5, "h": 0.5}}},
        )
    assert e.value.code == "BAD_IMAGE"
    with pytest.raises(ops.OpError) as e:
        ops.apply(data, "insert_image", {"slide_id": sid, "image": _png(10, 10), "target": {"placeholder_idx": 1}})
    assert e.value.code == "NO_PICTURE_PLACEHOLDER"


def test_replace_image_keeps_the_frame_and_alt_text():
    data = deck("features.pptx")
    sid = first_slide(data, 1)
    pic = next(s for s in get(data, sid)["shapes"] if s["type"] == "picture")
    with_alt, _ = ops.apply(data, "set_alt_text", {"slide_id": sid, "shape_id": pic["shape_id"], "text": "Linha de triagem"})
    new, _ = ops.apply(with_alt, "replace_image", {"slide_id": sid, "shape_id": pic["shape_id"], "image": _png(1000, 1000)})
    after = next(s for s in get(new, sid)["shapes"] if s["shape_id"] == pic["shape_id"])
    assert after["emu"] == pic["emu"] and after["alt_text"] == "Linha de triagem"
    prs = read.open_deck(new)
    sh = next(x for x in prs.slides.get(sid).shapes if x.shape_id == pic["shape_id"])
    frame = sh.width / sh.height  # 6:3.75 = 1.6; a square image: crop top and bottom
    assert sh.crop_left == 0 and sh.crop_top > 0 and abs(sh.crop_top - sh.crop_bottom) < 1e-6
    assert abs((1 - sh.crop_top - sh.crop_bottom) - 1 / frame) < 0.01


def test_edit_table_rows_columns_and_cells():
    data = deck("features.pptx")
    sid = first_slide(data, 0)
    table = next(s for s in get(data, sid)["shapes"] if s["type"] == "table")
    new, _ = ops.apply(
        data,
        "edit_table",
        {
            "slide_id": sid,
            "shape_id": table["shape_id"],
            "operations": [
                {"op": "set_cell", "row": 1, "col": 2, "text": "€ 450"},
                {"op": "add_row", "at": 4},
                {"op": "set_cell", "row": 4, "col": 0, "text": "Público"},
                {"op": "add_column", "at": 3},
                {"op": "set_cell", "row": 0, "col": 3, "text": "Suporte"},
                {"op": "delete_row", "row": 2},
            ],
        },
    )
    t = next(s for s in get(new, sid)["shapes"] if s["shape_id"] == table["shape_id"])["table"]
    assert (t["rows"], t["cols"]) == (4, 4)
    assert [row[0] for row in t["cells"]] == ["Escalão", "Base", "Empresa", "Público"]
    assert t["cells"][1][2] == "€ 450" and t["cells"][0][3] == "Suporte" and t["cells"][3][1] == ""
    prs = read.open_deck(new)
    sh = next(x for x in prs.slides.get(sid).shapes if x.shape_id == table["shape_id"])
    assert abs(sum(c.width for c in sh.table.columns) - sh.width) < 10  # the table keeps its width



def test_edit_table_new_row_and_column_with_their_cells():
    data = deck("features.pptx")
    sid = first_slide(data, 0)
    table = next(s for s in get(data, sid)["shapes"] if s["type"] == "table")
    before = table["table"]
    new, _ = ops.apply(
        data,
        "edit_table",
        {
            "slide_id": sid,
            "shape_id": table["shape_id"],
            "operations": [
                {"op": "add_row", "cells": ["Start", "até 10", "€ 100"]},  # at the end
                {"op": "add_row", "at": 1, "cells": ["Mini"]},  # fewer cells than columns: the rest stay empty
                {"op": "add_column", "at": 1, "cells": ["Nota"] + ["-"] * 99},  # more than rows: the extra ignored
            ],
        },
    )
    t = next(s for s in get(new, sid)["shapes"] if s["shape_id"] == table["shape_id"])["table"]
    assert (t["rows"], t["cols"]) == (before["rows"] + 2, before["cols"] + 1)
    assert t["cells"][-1][0] == "Start" and t["cells"][-1][2:4] == ["até 10", "€ 100"]
    assert t["cells"][1][0] == "Mini" and t["cells"][1][2] == ""
    assert [row[1] for row in t["cells"]] == ["Nota"] + ["-"] * (t["rows"] - 1)
    assert [row[0] for row in t["cells"]][2:-1] == [row[0] for row in before["cells"]][1:]

def test_set_notes():
    data = deck("simple.pptx")
    sid = first_slide(data, 0)
    new, _ = ops.apply(data, "set_notes", {"slide_id": sid, "text": "Fonte: KB 'Política de garantia' (WP-4)."})
    assert get(new, sid)["notes"] == "Fonte: KB 'Política de garantia' (WP-4)."


def test_the_outline_says_what_each_slide_holds_beyond_its_text():
    """A slide whose content is a table, a picture or a group is not blank in the outline: the model learns there is
    a price table on slide 1 without reading every slide."""
    rows = {o["index"] + 1: o["summary"] for o in read.outline(read.open_deck(deck("features.pptx")))}
    assert rows[1] == "table, 4 rows by 3 columns: Escalão, Utilizadores, Preço mensal"
    assert rows[2].startswith("picture")
    assert rows[3] == "Componente 1 · Componente 2 · Componente 3"
    agenda = {o["index"] + 1: o["summary"] for o in read.outline(read.open_deck(deck("simple.pptx")))}[2]
    assert agenda == "Contexto e objetivos / Solução proposta / Arquitetura / Calendário / Preços e condições"
    report = read.open_deck((REPO / "fixtures" / "eval" / "report.pptx").read_bytes())
    longest = max((o["summary"] for o in read.outline(report)), key=len)
    assert len(longest) == 200 and longest.endswith("…")


def test_slides_are_placed_next_to_another_slide_by_its_id():
    """after_slide_id / before_slide_id: no counting (measured: gemma-4 put "before the agenda" at position 0)."""
    data = deck("simple.pptx")
    a, b, c = [o["slide_id"] for o in read.outline(read.open_deck(data))]
    order = lambda d: [o["slide_id"] for o in read.outline(read.open_deck(d))]  # noqa: E731
    moved, _ = ops.apply(data, "move_slide", {"slide_id": c, "before_slide_id": b})
    assert order(moved) == [a, c, b]
    moved, _ = ops.apply(data, "move_slide", {"slide_id": a, "after_slide_id": c})
    assert order(moved) == [b, c, a]
    added, result = ops.apply(
        data,
        "add_slide",
        {
            "layout": "Title and Content",
            "after_slide_id": a,
            "placeholders": [{"type": "title", "text": "Nova"}, {"type": "body", "text": "Corpo"}],
        },
    )
    new = result["slides"][0]
    assert order(added) == [a, new, b, c]
    texts = [text_of(s) for s in get(added, new)["shapes"]]
    assert ["Nova"] in texts and ["Corpo"] in texts  # "body": the layout's content placeholder
    added, result = ops.apply(data, "add_slide", {"layout": "Title Slide", "placeholders": [{"type": "title", "text": "Fim"}]})
    assert order(added)[-1] == result["slides"][0]  # no place given: at the end
    assert ["Fim"] in [text_of(s) for s in get(added, result["slides"][0])["shapes"]]  # a title slide's centred title
    dup, result = ops.apply(data, "duplicate_slide", {"slide_id": a, "before_slide_id": c})
    assert order(dup) == [a, b, result["slides"][0], c]
    with pytest.raises(ops.OpError) as e:
        ops.apply(data, "move_slide", {"slide_id": a, "after_slide_id": a})
    assert e.value.code == "SLIDE_NOT_FOUND"
    with pytest.raises(ops.OpError):
        ops.apply(data, "move_slide", {"slide_id": a})


def test_a_table_cell_without_its_column_is_refused_not_a_crash():
    data = deck("features.pptx")
    sid = first_slide(data)
    with pytest.raises(ops.OpError) as e:
        ops.apply(data, "edit_table", {"slide_id": sid, "shape_id": 3, "operations": [{"op": "set_cell", "row": 2, "text": "x"}]})
    assert e.value.code == "BAD_ARGUMENTS" and "col" in str(e.value)


def test_the_executor_takes_nulls_as_left_out_and_enum_values_in_any_case():
    from app.agent.tools import Executor, ToolError

    ex = Executor.__new__(Executor)
    from app.agent.tools import load_schemas

    ex.schemas = load_schemas()
    args = ex.parse(
        "update_text",
        {"slide_id": 256, "shape_id": 2, "paragraphs": [{"runs": [{"text": "x", "theme_color": None, "bold": None}]}]},
    )
    assert args["paragraphs"][0]["runs"][0] == {"text": "x"}
    assert (
        ex.parse("add_shape", {"slide_id": 256, "kind": "TEXT_BOX", "box": {"x": 0, "y": 0, "w": 0.1, "h": 0.1}})["kind"]
        == "text_box"
    )
    with pytest.raises(ToolError) as e:
        ex.parse(
            "edit_table", {"slide_id": 1, "shape_id": 3, "operations": [{"op": "set_text", "row": 0, "col": 0, "text": "x"}]}
        )
    assert e.value.code == "BAD_ARGUMENTS"  # not a case away from a listed value: refused
    with pytest.raises(ToolError) as e:
        ex.parse("update_text", {"slide_id": None, "shape_id": 2, "paragraphs": [{"runs": [{"text": "x"}]}]})
    assert "slide_id" in str(e.value)  # a required field sent as null: refused by name


def test_edit_paragraphs_changes_only_the_named_ones():
    """A bullet added, changed or removed by its index; every other paragraph's XML is untouched (measured: rewriting
    the whole list with update_text, gemma-4 dropped bullets and flattened levels)."""
    data = deck("simple.pptx")
    agenda = first_slide(data, 1)
    before = [etree.tostring(p) for p in _paragraph_elements(data, agenda, 3)]
    new, _ = ops.apply(
        data,
        "edit_paragraphs",
        {
            "slide_id": agenda,
            "shape_id": 3,
            "operations": [
                {"op": "insert", "after": 3, "text": "Garantia"},  # after Calendário (level 1): its level
                {"op": "insert", "text": "Perguntas", "level": 0},  # at the end
                {"op": "insert", "after": -1, "text": "Abertura"},  # first
                {"op": "set", "index": 1, "text": "A solução"},
                {"op": "delete", "index": 2},  # Arquitetura
            ],
        },
    )
    shape = next(x for x in get(new, agenda)["shapes"] if x["shape_id"] == 3)
    rows = [("".join(r["text"] for r in p["runs"]), p.get("level", 0)) for p in shape["paragraphs"]]
    assert rows == [
        ("Abertura", 0),
        ("Contexto e objetivos", 0),
        ("A solução", 0),
        ("Calendário", 1),
        ("Garantia", 1),
        ("Preços e condições", 0),
        ("Perguntas", 0),
    ]
    after = [etree.tostring(p) for p in _paragraph_elements(new, agenda, 3)]
    assert after[1] == before[0] and after[3] == before[3] and after[5] == before[4]  # untouched: the same XML
    unchanged_except(data, new, {slide_part(data, agenda)})
    lines, _ = ops.apply(
        data,
        "edit_paragraphs",
        {"slide_id": agenda, "shape_id": 3, "operations": [{"op": "set", "index": 4, "text": "Preços\nCondições"}]},
    )
    assert text_of(next(x for x in get(lines, agenda)["shapes"] if x["shape_id"] == 3))[-2:] == ["Preços", "Condições"]
    with pytest.raises(ops.OpError) as e:
        ops.apply(data, "edit_paragraphs", {"slide_id": agenda, "shape_id": 3, "operations": [{"op": "delete", "index": 9}]})
    assert e.value.code == "BAD_PARAGRAPH" and "0-4" in str(e.value)


def _paragraph_elements(data: bytes, slide_id: int, shape_id: int):
    prs = read.open_deck(data)
    s = next(x for x in prs.slides if x.slide_id == slide_id)
    sh = next(x for x in s.shapes if x.shape_id == shape_id)
    return sh.text_frame._txBody.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}p")


def test_a_plain_paragraph_with_line_breaks_is_several_paragraphs():
    """{"text": "a\\nb"}: two paragraphs (measured: gemma-4 sent a four-bullet list that way)."""
    from app.agent.tools import Executor, load_schemas

    ex = Executor.__new__(Executor)
    ex.schemas = load_schemas()
    args = ex.parse(
        "update_text",
        {"slide_id": 1, "shape_id": 3, "paragraphs": [{"text": "Instagram\nTikTok", "level": 1}, {"text": "Podcasts"}]},
    )
    assert args["paragraphs"] == [
        {"level": 1, "runs": [{"text": "Instagram"}]},
        {"level": 1, "runs": [{"text": "TikTok"}]},
        {"runs": [{"text": "Podcasts"}]},
    ]
    keep = ex.parse("update_text", {"slide_id": 1, "shape_id": 3, "paragraphs": [{"runs": [{"text": "line\nbreak"}]}]})
    assert keep["paragraphs"] == [{"runs": [{"text": "line\nbreak"}]}]  # inside runs: a line break, as before


def test_reading_a_deck_never_changes_it():
    """The read model only reads: every slide of every fixture deck, read, serialises exactly as before (measured: a
    run's colour read through python-pptx wrote an empty <a:solidFill/> into it, and edits saved them)."""
    for name in ("simple.pptx", "features.pptx", "fidelity.pptx"):
        prs = read.open_deck(deck(name))
        before = [etree.tostring(s._element) for s in prs.slides]
        for o in read.outline(prs):
            read.slide(prs, o["slide_id"])
        assert [etree.tostring(s._element) for s in prs.slides] == before, name


def test_an_edit_leaves_the_other_runs_of_its_shape_as_they_were():
    data = deck("simple.pptx")
    agenda = first_slide(data, 1)
    before = [etree.tostring(p) for p in _paragraph_elements(data, agenda, 3)]
    new, _ = ops.apply(
        data, "edit_paragraphs", {"slide_id": agenda, "shape_id": 3, "operations": [{"op": "set", "index": 0, "text": "X"}]}
    )
    after = [etree.tostring(p) for p in _paragraph_elements(new, agenda, 3)]
    assert after[1:] == before[1:]
    assert b"solidFill" not in b"".join(after)


def test_a_slide_by_its_number_and_a_layout_by_its_name_written_another_way():
    """Slide numbers (1 to 255: never an ID, ECMA-376 starts IDs at 256) are positions in the deck as it is now; the
    recorded call holds the ID. A layout named "title_and_content" is "Title and Content"."""
    from app.agent.tools import Executor

    data = deck("simple.pptx")
    ids = [o["slide_id"] for o in read.outline(read.open_deck(data))]
    assert Executor._positions({"slide_id": 3, "after_slide_id": 1, "shape_id": 2}, data) == {
        "slide_id": ids[2],
        "after_slide_id": ids[0],
        "shape_id": 2,
    }
    assert Executor._positions({"slide_id": ids[1]}, data) == {"slide_id": ids[1]}  # an ID stays
    from app.agent.tools import ToolError

    with pytest.raises(ToolError) as e:
        Executor._positions({"slide_id": 9}, data)
    assert "1 to 3" in str(e.value)
    new, result = ops.apply(data, "add_slide", {"layout": "title_and_content", "placeholders": [{"type": "title", "text": "T"}]})
    assert get(new, result["slides"][0])["layout"] == "Title and Content"


# ── slides from another deck (spec PJ-6, PJ-7) ──────────────────────
def _names(data: bytes) -> list[str]:
    return zipfile.ZipFile(io.BytesIO(data)).namelist()


def test_slides_copied_from_another_deck_bring_their_content_and_parts():
    target, source = deck("simple.pptx"), deck("features.pptx")
    src_ids = [o["slide_id"] for o in read.outline(read.open_deck(source))]
    new, result = ops.apply(
        target,
        "copy_slides",
        {"source": source, "from_deck_id": "x", "slide_ids": src_ids, "after_slide_id": first_slide(target)},
    )
    names = _names(new)
    assert len(names) == len(set(names)), "no part name twice"
    order = [o["slide_id"] for o in read.outline(read.open_deck(new))]
    assert order[1:5] == result["slides"] and len(order) == 7  # after the first slide, in their order
    table, picture, group, hidden = (get(new, sid) for sid in result["slides"])
    assert next(s for s in table["shapes"] if s["type"] == "table")["table"]["cells"][2] == ["Pro", "até 500", "€ 2 500"]
    assert any(s["type"] == "picture" for s in picture["shapes"]) and picture["notes"] == "Imagem de exemplo."
    assert [text_of(x) for x in next(s for s in group["shapes"] if s["type"] == "group")["shapes"]] == [
        ["Componente 1"],
        ["Componente 2"],
        ["Componente 3"],
    ]
    assert hidden["hidden"] and text_of(next(s for s in hidden["shapes"] if s["shape_id"] != 2)) == [
        "Não aparece na apresentação"
    ]
    assert table["layout"] == "Title Only" and text_of(
        next(s for s in table["shapes"] if s.get("placeholder", {}).get("type") == "title")
    ) == ["Preços por escalão"]
    # the picture's bytes travelled: an image part in the target that the source's picture had
    src_media = {n: zipfile.ZipFile(io.BytesIO(source)).read(n) for n in _names(source) if n.startswith("ppt/media/")}
    new_media = [zipfile.ZipFile(io.BytesIO(new)).read(n) for n in names if n.startswith("ppt/media/")]
    assert any(b in new_media for b in src_media.values())
    assert result["unmatched"] == []


def test_a_copied_chart_brings_its_own_chart_and_workbook():
    target, source = deck("simple.pptx"), deck("fidelity.pptx")
    chart_slide = read.outline(read.open_deck(source))[0]["slide_id"]
    new, result = ops.apply(target, "copy_slides", {"source": source, "slide_ids": [chart_slide]})
    names = _names(new)
    assert len(names) == len(set(names))
    assert any(n.startswith("ppt/charts/chart") for n in names) and any(n.startswith("ppt/embeddings/") for n in names)
    assert any(s["type"] == "chart" for s in get(new, result["slides"][0])["shapes"])


def test_text_with_no_place_in_the_target_layouts_is_kept_and_reported():
    source = deck("simple.pptx")
    prs = read.open_deck((REPO / "templates" / "default.pptx").read_bytes())  # no slides yet
    master = prs.slide_masters[0]
    # a template with no layout of two bodies ("Content with Caption" would take both columns: an exact match)
    for layout in [lay for lay in master.slide_layouts if lay.name in ("Two Content", "Comparison", "Content with Caption")]:
        master.slide_layouts.remove(layout)
    prs.slides.add_slide(master.slide_layouts[0])
    buf = io.BytesIO()
    prs.save(buf)
    target = buf.getvalue()
    two_columns = read.outline(read.open_deck(source))[2]["slide_id"]
    new, result = ops.apply(target, "copy_slides", {"source": source, "slide_ids": [two_columns]})
    slide = get(new, result["slides"][0])
    assert slide["layout"] == "Title and Content"  # the nearest: a title and one body
    texts = [text_of(s) for s in slide["shapes"]]
    assert ["Processos manuais", "Vários sistemas"] in texts and ["Fluxo único", "Um só sistema"] in texts
    assert len(result["unmatched"]) == 1 and "Fluxo único" in result["unmatched"][0] and "text box" in result["unmatched"][0]


@pytest.mark.skipif(
    shutil.which("soffice") is None or not Path("/usr/lib/libreoffice/share/registry/impress.xcd").exists(),
    reason="LibreOffice Impress is not installed here (run make check)",
)
def test_a_deck_with_copied_slides_renders():
    import asyncio
    import tempfile

    from app.docengine import render

    new, _ = ops.apply(deck("simple.pptx"), "copy_slides", {"source": deck("features.pptx"), "slide_ids": [256, 257, 258]})
    cache = Path(tempfile.mkdtemp())
    keys = asyncio.run(render.ensure(new, cache))
    assert len(keys) == 6 and all((cache / k["key"] / "preview.png").exists() for k in keys)


def _other_template() -> bytes:
    """python-pptx's own default: 4:3 and the Office theme - unlike the 16:9 default template here."""
    from pptx import Presentation

    buf = io.BytesIO()
    Presentation().save(buf)
    return buf.getvalue()


def test_a_deck_on_another_template_keeps_its_slides_and_their_content():
    data = deck("features.pptx")
    before = read.outline(read.open_deck(data))
    new, result = ops.apply(data, "change_template", {"template": _other_template()})
    prs = read.open_deck(new)
    assert (prs.slide_width, prs.slide_height) == (9144000, 6858000)  # the template's 4:3
    after = read.outline(prs)
    assert [o["slide_id"] for o in after] == [o["slide_id"] for o in before] == result["slides"]  # the same slides
    assert [o["title"] for o in after] == [o["title"] for o in before]
    assert after[3]["hidden"]
    table = next(s for s in get(new, before[0]["slide_id"])["shapes"] if s["type"] == "table")
    assert table["table"]["cells"][2] == ["Pro", "até 500", "€ 2 500"]
    assert table["box"]["x"] >= 0 and table["box"]["x"] + table["box"]["w"] <= 1.001  # scaled into the narrower slide
    frame = next(sh for sh in next(x for x in prs.slides if x.slide_id == before[0]["slide_id"]).shapes if sh.has_table)
    assert sum(c.width for c in frame.table.columns) <= frame.width + 10  # its columns too (seen: wider than the slide)
    assert result["unmatched"] == []
    names = _names(new)
    assert len(names) == len(set(names))


def test_the_overflow_estimate_agrees_with_what_libreoffice_draws():
    """Calibrated on rendered slides (looked at, 2026-10-06): the report's results slide as it is fits; seven
    one-line bullets fit; nine pass the bottom of the box; fourteen two-line bullets overflow by far."""
    data = (REPO / "fixtures" / "eval" / "report.pptx").read_bytes()

    def body(n: int, words: str) -> bool:
        new, _ = ops.apply(
            data,
            "update_text",
            {"slide_id": 258, "shape_id": 3, "paragraphs": [{"runs": [{"text": f"Ponto {i + 1}: {words}"}]} for i in range(n)]},
        )
        return next(s for s in get(new, 258)["shapes"] if s["shape_id"] == 3)["overflow"]

    assert next(s for s in get(data, 258)["shapes"] if s["shape_id"] == 3)["overflow"] is False
    assert body(7, "uma frase longa sobre os resultados do trimestre") is False
    assert body(9, "uma frase longa sobre os resultados do trimestre") is True
    assert body(14, "uma frase longa sobre os resultados do trimestre com bastante detalhe") is True


def test_set_texts_keeps_untouched_paragraphs_exactly():
    data = deck("simple.pptx")
    sid = first_slide(data, 1)  # the agenda: five paragraphs
    body = next(s for s in get(data, sid)["shapes"] if len(s.get("paragraphs") or []) >= 5)
    old = text_of(body)
    lines = [old[0], "Novo ponto", old[1].upper(), *old[3:]]  # one inserted, one changed, one deleted (old[2])

    def xml_of(d: bytes) -> list[bytes]:
        sh = next(x for x in read.open_deck(d).slides.get(sid).shapes if x.shape_id == body["shape_id"])
        return [etree.tostring(p) for p in sh.text_frame._txBody.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}p")]

    new, _ = ops.apply(data, "set_texts", {"slide_id": sid, "shape_id": body["shape_id"], "texts": lines})
    assert text_of(next(s for s in get(new, sid)["shapes"] if s["shape_id"] == body["shape_id"])) == lines
    before, after = xml_of(data), xml_of(new)
    assert after[0] == before[0] and after[3:] == before[3:]  # untouched: byte for byte
    same, _ = ops.apply(data, "set_texts", {"slide_id": sid, "shape_id": body["shape_id"], "texts": old})
    assert xml_of(same) == before  # nothing typed differently: nothing changes
    with pytest.raises(ops.OpError):
        ops.apply(data, "set_texts", {"slide_id": 9999, "shape_id": body["shape_id"], "texts": ["x"]})


def test_a_slide_s_render_key_follows_the_installed_fonts(monkeypatch):
    from app.docengine import render

    data = deck("simple.pptx")
    monkeypatch.setattr(render, "_fonts", "a stand-in font")
    before = [k["key"] for k in render.keys(data)]
    monkeypatch.setattr(render, "_fonts", "the template's own font")
    after = [k["key"] for k in render.keys(data)]
    assert len(before) == len(after) and not set(before) & set(after)  # every slide is rendered again
    assert [k["key"] for k in render.keys(data)] == after  # and the same fonts give the same keys


def test_add_slide_idx_0_and_1_mean_title_and_body_on_a_layout_numbered_otherwise():
    data = deck("simple.pptx")
    prs = read.open_deck(data)
    layout = next(lo for m in prs.slide_masters for lo in m.slide_layouts if lo.name == "Title and Content")
    body = next(ph for ph in layout.placeholders if ph.placeholder_format.idx == 1)
    body.element.nvSpPr.nvPr.find(qn_ph()).set("idx", "16")  # as a template numbers its body (Banco CTT's: 16)
    renumbered = io.BytesIO()
    prs.save(renumbered)
    new, result = ops.apply(
        renumbered.getvalue(),
        "add_slide",
        {
            "layout": "Title and Content",
            "placeholders": [{"idx": 0, "text": "Boas práticas"}, {"idx": 1, "text": "Primeira\nSegunda"}],
        },
    )
    s = get(new, result["slides"][0])
    texts = {sh["placeholder"]["idx"]: text_of(sh) for sh in s["shapes"] if sh.get("placeholder")}
    assert texts[0] == ["Boas práticas"] and texts[16] == ["Primeira", "Segunda"]


def qn_ph():
    from pptx.oxml.ns import qn

    return qn("p:ph")


def _emf(*records: bytes) -> bytes:
    """A small EMF: its header, the records, its end (MS-EMF layouts)."""
    body = b"".join(records) + struct.pack("<IIIII", 14, 20, 0, 16, 20)
    header = bytearray(108)
    struct.pack_into("<II", header, 0, 1, 108)
    struct.pack_into("<iiii", header, 8, 0, 0, 100, 20)  # rclBounds (px)
    struct.pack_into("<iiii", header, 24, 0, 0, 2600, 520)  # rclFrame (0.01 mm)
    struct.pack_into("<I", header, 40, 0x464D4520)  # " EMF"
    struct.pack_into("<I", header, 48, 108 + len(body))
    return bytes(header) + body


def test_emf_pictures_are_drawn_from_their_records_and_not_clipped_at_their_frame():
    from app.docengine import emf

    comment = b"GDIC" + struct.pack("<I", emf.MULTIFORMATS) + b"\0" * 24  # an alternative copy (a PDF, in Banco CTT's logo)
    raw = _emf(
        struct.pack("<III", 0x11, 12, 8),  # SETMAPMODE anisotropic
        struct.pack("<IIii", 0x0A, 16, 0, 0),  # SETWINDOWORGEX
        struct.pack("<IIii", 0x09, 16, 1000, 200),  # SETWINDOWEXTEX: the declared window 0-1000 x 0-200
        struct.pack("<IIii", 0x0B, 16, 100, 20),  # SETVIEWPORTEXTEX
        struct.pack("<III", 0x46, 12 + len(comment), len(comment)) + comment,
        struct.pack("<II4iI4h", 0x57, 36, 0, 0, 0, 0, 2, 0, 0, 1000, 210),  # POLYLINE16 down to y 210: past the window
    )
    assert emf.is_emf(raw)
    out, (left, top, right, bottom) = emf.prepare(raw)
    assert b"GDIC" not in out and b"SLID" in out  # the alternative copy is skipped: the records are drawn
    assert (left, top) == (2 / 1000, 2 / 200) and right == 2 / 1000 and bottom == 12 / 200  # grown to the line + 2
    frame = struct.unpack_from("<iiii", out, 24)
    assert frame[3] == round(520 * (1 + top + bottom)) and frame[2] == round(2600 * (1 + left + right))
    plain = _emf(struct.pack("<III", 0x11, 12, 1))  # MM_TEXT, no window: nothing to grow
    assert emf.prepare(plain) == (plain, (0.0, 0.0, 0.0, 0.0))
    assert emf.prepare(b"\x89PNG not an emf") == (b"\x89PNG not an emf", (0.0, 0.0, 0.0, 0.0))


def test_a_diagram_grows_with_copies_of_its_boxes_and_arrows_attached_to_them():
    data, blank = ops.apply(deck("simple.pptx"), "add_slide", {"layout": "Title Only"})  # a diagram's slide
    sid = blank["slides"][0]
    first = {"x": 0.1, "y": 0.7, "w": 0.2, "h": 0.1}
    args = {"slide_id": sid, "kind": "rounded_rectangle", "box": first, "paragraphs": [{"runs": [{"text": "NEW"}]}]}
    data, made = ops.apply(data, "add_shape", args)
    box = made["shape_id"]
    data, copy_ = ops.apply(data, "duplicate_shape", {"slide_id": sid, "shape_id": box, "box": {"x": 0.6, "y": 0.7},
                                                      "paragraphs": [{"runs": [{"text": "DONE"}]}]})  # fmt: skip
    new = copy_["shape_id"]
    prs = read.open_deck(data)
    s = prs.slides.get(sid)
    a, b = (next(x for x in s.shapes if x.shape_id == i) for i in (box, new))
    strip = lambda el: etree.tostring(el.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}spPr/{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom"))  # noqa: E731
    assert strip(a._element) == strip(b._element) and (b.width, b.height) == (a.width, a.height)  # same box, same size
    assert abs(b.left / prs.slide_width - 0.6) < 0.001 and b.text_frame.text == "DONE" and a.text_frame.text == "NEW"
    data, line = ops.apply(data, "connect_shapes", {"slide_id": sid, "from_shape_id": box, "to_shape_id": new})
    data, styled = ops.apply(data, "connect_shapes", {"slide_id": sid, "from_shape_id": new, "to_shape_id": box,
                                                      "like_connector_id": line["shape_id"], "arrow": "both"})  # fmt: skip
    shapes = {x["shape_id"]: x for x in get(data, sid)["shapes"]}
    assert shapes[line["shape_id"]]["type"] == "connector" and shapes[line["shape_id"]]["connects"] == [box, new]
    assert shapes[styled["shape_id"]]["connects"] == [new, box]
    el = next(x for x in read.open_deck(data).slides.get(sid).shapes if x.shape_id == line["shape_id"])._element
    a_ = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    assert el.find(f".//{a_}stCxn").get("idx") == "3" and el.find(f".//{a_}endCxn").get("idx") == "1"  # right -> left
    assert el.find(f".//{a_}prstGeom").get("prst") == "straightConnector1"  # in line: straight
    assert el.find(f".//{a_}tailEnd") is not None and el.find(f".//{a_}headEnd") is None
    back = next(x for x in read.open_deck(data).slides.get(sid).shapes if x.shape_id == styled["shape_id"])._element
    assert back.find(f".//{a_}xfrm").get("flipH") == "1" and back.find(f".//{a_}headEnd") is not None  # leftwards, both
    with pytest.raises(ops.OpError):
        ops.apply(data, "connect_shapes", {"slide_id": sid, "from_shape_id": box, "to_shape_id": box})
    title = next(x for x in get(data, sid)["shapes"] if (x.get("placeholder") or {}).get("idx") == 0)["shape_id"]
    with pytest.raises(ops.OpError):  # a placeholder is not copied
        ops.apply(data, "duplicate_shape", {"slide_id": sid, "shape_id": title})


def test_shapes_inside_a_group_are_not_copied_or_connected_on_their_own():
    data = deck("features.pptx")
    sid = first_slide(data, 2)
    group = next(s for s in get(data, sid)["shapes"] if s["type"] == "group")
    member = group["shapes"][0]["shape_id"]
    with pytest.raises(ops.OpError) as e:
        ops.apply(data, "duplicate_shape", {"slide_id": sid, "shape_id": member})
    assert e.value.code == "IN_GROUP"


def test_a_copy_goes_where_it_covers_nothing_and_an_arrow_looks_like_the_diagram_s():
    data, blank = ops.apply(deck("simple.pptx"), "add_slide", {"layout": "Blank"})  # a diagram's slide
    sid = blank["slides"][0]
    boxes = []
    for x in (0.1, 0.35):  # two boxes in a row: the copy of the first must not land on the second
        args = {"slide_id": sid, "kind": "rectangle", "box": {"x": x, "y": 0.75, "w": 0.2, "h": 0.1}}
        data, made = ops.apply(data, "add_shape", args)
        boxes.append(made["shape_id"])
    data, copy_ = ops.apply(data, "duplicate_shape", {"slide_id": sid, "shape_id": boxes[0]})
    assert copy_["covers"] == []
    prs = read.open_deck(data)
    s = prs.slides.get(sid)
    others = [(x.shape_id, x.left, x.top, x.width, x.height) for x in s.shapes if x.shape_id != copy_["shape_id"]]
    new = next(x for x in s.shapes if x.shape_id == copy_["shape_id"])
    assert ops._overlapping((new.left, new.top, new.width, new.height), others) == []
    data, over = ops.apply(data, "duplicate_shape", {"slide_id": sid, "shape_id": boxes[0], "box": {"x": 0.35, "y": 0.75}})
    assert over["moved"] and over["covers"] == []  # asked to go on the second box: put at the nearest free place
    moved = next(x for x in read.open_deck(data).slides.get(sid).shapes if x.shape_id == over["shape_id"])
    assert abs(moved.top / read.open_deck(data).slide_height - 0.75) < 0.35  # near where it was asked for
    link = {"slide_id": sid, "from_shape_id": boxes[0], "to_shape_id": boxes[1], "kind": "elbow"}
    data, first = ops.apply(data, "connect_shapes", link)
    a_ = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    el = lambda d, i: next(x for x in read.open_deck(d).slides.get(sid).shapes if x.shape_id == i)._element  # noqa: E731
    prs = read.open_deck(data)  # the diagram's own line, as a template has it: 2.5 pt
    target = next(x for x in prs.slides.get(sid).shapes if x.shape_id == first["shape_id"])._element
    target.find(f".//{a_}ln").set("w", "31750")
    edited = io.BytesIO()
    prs.save(edited)
    link = {"slide_id": sid, "from_shape_id": boxes[1], "to_shape_id": copy_["shape_id"]}
    data, second = ops.apply(edited.getvalue(), "connect_shapes", link)
    assert el(data, second["shape_id"]).find(f".//{a_}ln").get("w") == "31750"  # like the slide's connectors


def test_fit_text_shrinks_to_the_largest_size_that_fits_and_never_below_14_pt():
    from app.docengine import textfit

    data = deck("simple.pptx")
    sid = first_slide(data, 1)  # the agenda: a body of 5 points
    body = next(s for s in get(data, sid)["shapes"] if len(s.get("paragraphs") or []) >= 5)
    ops_ = [{"op": "insert", "text": f"Mais um ponto da agenda, o número {i}"} for i in range(5)]
    data, _ = ops.apply(data, "edit_paragraphs", {"slide_id": sid, "shape_id": body["shape_id"], "operations": ops_})
    assert next(s for s in get(data, sid)["shapes"] if s["shape_id"] == body["shape_id"])["overflow"]  # ten points: too long
    fitted, res = ops.apply(data, "fit_text", {"slide_id": sid, "shape_id": body["shape_id"]})
    shape = next(s for s in get(fitted, sid)["shapes"] if s["shape_id"] == body["shape_id"])
    assert 0.5 < res["scale"] < 1 and shape["overflow"] is False
    sh = next(x for x in read.open_deck(fitted).slides.get(sid).shapes if x.shape_id == body["shape_id"])
    assert textfit.smallest_size(sh) >= 14  # never below the body minimum
    many = [{"op": "insert", "text": "Ponto extra que não cabe de forma nenhuma na caixa"} for _ in range(25)]
    data, _ = ops.apply(data, "edit_paragraphs", {"slide_id": sid, "shape_id": body["shape_id"], "operations": many})
    with pytest.raises(ops.OpError) as e:
        ops.apply(data, "fit_text", {"slide_id": sid, "shape_id": body["shape_id"]})
    assert e.value.code == "CANNOT_FIT"


def test_fit_text_grows_the_box_into_free_space_when_shrinking_is_not_allowed():
    data, blank = ops.apply(deck("simple.pptx"), "add_slide", {"layout": "Blank"})
    sid = blank["slides"][0]
    lines = [{"runs": [{"text": f"Ponto número {i} de uma lista longa", "size_pt": 12}]} for i in range(10)]
    place = {"x": 0.1, "y": 0.1, "w": 0.4, "h": 0.15}
    # a rectangle, text centred, no autofit (a text box made by add_shape grows with its text: never overflows)
    data, box = ops.apply(data, "add_shape", {"slide_id": sid, "kind": "rectangle", "box": place, "paragraphs": lines})
    shape = lambda d: next(s for s in get(d, sid)["shapes"] if s["shape_id"] == box["shape_id"])  # noqa: E731
    assert shape(data)["overflow"]  # ten 12 pt lines in a short box
    grown, res = ops.apply(data, "fit_text", {"slide_id": sid, "shape_id": box["shape_id"]})
    assert res["scale"] == 1.0 and res["grown_pt"] > 0 and shape(grown)["overflow"] is False  # 12 pt: not shrunk, grown
    under = {"x": 0.1, "y": 0.26, "w": 0.4, "h": 0.6}
    blocked, _ = ops.apply(data, "add_shape", {"slide_id": sid, "kind": "rectangle", "box": under})
    with pytest.raises(ops.OpError) as e:  # a box right under it: no room to grow
        ops.apply(blocked, "fit_text", {"slide_id": sid, "shape_id": box["shape_id"]})
    assert e.value.code == "CANNOT_FIT"


def test_text_is_measured_run_by_run_with_kerning_letter_spacing_and_its_letters_descent():
    """Measured against LibreOffice's renders (docs/decisions.md): a paragraph measured in its first run's bold wrapped
    a line more; widths without the font's kerning were 3% wide; a 71 pt title's descenders passed its box."""
    from app.docengine import textfit as tf

    face = "Liberation Sans"  # in the base image: these hold without the licensed fonts
    plain = (12.0, face, False, False, 0.0)
    bold = (12.0, face, True, False, 0.0)
    words = "uma frase com algumas palavras para medir"
    room = tf._width(words, plain, face, face) + 0.5
    assert tf._lines([("Nota: ", *bold), (words, *plain)], room + tf._width("Nota: ", bold, face, face), True, face, face)[0] == 1
    assert tf._lines([("Nota: " + words, *bold)], room + tf._width("Nota: ", bold, face, face), True, face, face)[0] == 2
    tight = (12.0, face, False, False, -0.5)  # letter spacing (spc) counts on every character
    assert tf._width(words, tight, face, face) == pytest.approx(tf._width(words, plain, face, face) - 0.5 * len(words))
    assert tf.width_pt("AV", face, 40) < tf.width_pt("A", face, 40) + tf.width_pt("V", face, 40) - 1  # kerned (raqm)
    assert tf._descent([("ABC", *plain)], face, face) < 0.5 < tf._descent([("gpy", *plain)], face, face)


def test_an_empty_line_is_as_tall_as_its_own_mark():
    from app.docengine import textfit as tf

    data, blank = ops.apply(deck("simple.pptx"), "add_slide", {"layout": "Blank"})
    sid = blank["slides"][0]
    lines = [{"runs": [{"text": "Primeira linha", "size_pt": 14}]}] + [{"runs": [{"text": ""}]} for _ in range(6)]
    place = {"x": 0.1, "y": 0.1, "w": 0.4, "h": 0.5}
    data, box = ops.apply(data, "add_shape", {"slide_id": sid, "kind": "rectangle", "box": place, "paragraphs": lines})

    def extent(size):
        prs = read.open_deck(data)
        sh = next(x for x in prs.slides.get(sid).shapes if x.shape_id == box["shape_id"])
        a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        for p in sh._element.txBody.findall(f"{a}p")[1:]:
            end = p.find(f"{a}endParaRPr")
            if end is None:
                end = etree.SubElement(p, f"{a}endParaRPr")
            end.set("sz", str(size * 100))
        return tf.measure(sh)["top"]  # centred: the taller the empty lines, the higher the text starts

    assert extent(8) - extent(28) == pytest.approx(6 * 20 * tf.PITCH / 2, abs=0.5)


def test_content_goes_where_the_layout_has_its_place_and_no_empty_placeholder_is_left():
    """Banco CTT's layouts: a heading typed "body", text boxes in columns with their own headings and agenda numbers
    (each layout rendered and looked at); the model, given the boxes, put a list into a subtitle box, or into one
    agenda box, or nowhere. The application places it."""
    ctt = (REPO / "templates" / "bancoctt.pptx").read_bytes()

    def texts(data, sid):
        from pptx import Presentation

        s = Presentation(io.BytesIO(data)).slides.get(sid)
        return {ph.placeholder_format.idx: ph.text_frame.text for ph in s.placeholders if ph.has_text_frame}

    topics = ["Regime e finalidades", "Fases do processo", "Servicing", "Crédito obras", "Crédito bonificado"]
    data, made = ops.apply(ctt, "add_slide", {"layout": "1_Texto", "content": {"title": "Índice", "points": topics}})
    got = texts(data, made["slides"][0])
    assert got == {0: "Índice", 17: topics[0], 18: topics[1], 19: topics[2], 20: topics[3], 21: topics[4]}  # [11] removed
    data, made = ops.apply(ctt, "add_slide", {"layout": "3_Agenda", "content": {"title": "Agenda", "points": topics[:3]}})
    got = texts(data, made["slides"][0])
    assert got == {0: "Agenda", 29: topics[0], 47: "01", 13: topics[1], 48: "02", 21: topics[2], 49: "03"}
    columns = [{"heading": f"Fase {i}", "text": f"O que acontece na fase {i}."} for i in range(1, 5)]
    content = {"title": "Fases", "subtitle": "Do pedido à escritura", "points": columns}
    data, made = ops.apply(ctt, "add_slide", {"layout": "10_Texto", "content": content})
    got = texts(data, made["slides"][0])
    assert got[19] == "Fases" and got[20] == "Do pedido à escritura"
    assert (got[30], got[12]) == ("Fase 1", "O que acontece na fase 1.")
    assert (got[36], got[35]) == ("Fase 4", "O que acontece na fase 4.") and len(got) == 10
    many = [f"Ponto {i}" for i in range(8)]  # more points than boxes: one list, in the box with the most room
    data, made = ops.apply(ctt, "add_slide", {"layout": "1_Texto", "content": {"title": "Muitos", "points": many}})
    got = texts(data, made["slides"][0])
    assert list(got.values()).count("\n".join(many)) == 1 and len(got) == 2
    sid = made["slides"][0]
    data, _ = ops.apply(data, "fill_slide", {"slide_id": sid, "content": {"title": "Outro", "points": topics[:2]}})
    assert texts(data, sid) == {0: "Outro", 17: topics[0], 18: topics[1]}  # what it said is replaced
    cover, made = ops.apply(ctt, "add_slide", {"layout": "2_Capa S/Imagem", "content": {"title": "Novo"}})
    sid = made["slides"][0]  # a cover has no place for points: the slide moves to the layout that fits them, same ID
    data, res = ops.apply(cover, "fill_slide", {"slide_id": sid, "content": {"title": "Índice", "points": topics}})
    assert res["layout"] == "1_Texto" and res["slides"] == [sid]
    assert texts(data, sid) == {0: "Índice", 17: topics[0], 18: topics[1], 19: topics[2], 20: topics[3], 21: topics[4]}
    _, res = ops.apply(ctt, "add_slide", {"layout": "Mensagem_Final", "content": {"title": "x", "subtitle": "y"}})
    assert res["left_out"] == ["subtitle"]  # a heading alone: the slide is made, the subtitle reported


def test_placed_content_is_fitted_to_its_box():
    ctt = (REPO / "templates" / "bancoctt.pptx").read_bytes()
    long_title = "Índice dos tópicos de crédito à habitação do Banco CTT"  # 40 pt in a narrow title box
    data, made = ops.apply(ctt, "add_slide", {"layout": "1_Texto", "content": {"title": long_title, "points": ["a", "b"]}})
    title = next(sh for sh in get(data, made["slides"][0])["shapes"] if (sh.get("placeholder") or {}).get("idx") == 0)
    assert title["overflow"] is False


def test_add_slides_makes_a_slide_for_each_item_on_the_layout_its_content_fits():
    ctt = (REPO / "templates" / "bancoctt.pptx").read_bytes()
    items = [
        {"title": "Contexto"}, {"title": "Objetivos", "points": ["Um", "Dois", "Três"]}, {"title": "Capa", "subtitle": "Outubro"},
    ]  # fmt: skip
    data, res = ops.apply(ctt, "add_slides", {"slides": items})
    prs = read.open_deck(data)
    assert [o["slide_id"] for o in read.outline(prs)] == res["new_slide_ids"] and len(res["new_slide_ids"]) == 3
    layouts = [o["layout"] for o in read.outline(prs)]
    assert layouts[1] != layouts[0] and layouts[2] in ("1_Capa C/ Imagem", "2_Capa S/Imagem")  # 3 points; a cover
    texts = [" / ".join(" ".join(r.get("text", "") for r in p.get("runs", [])) for sh in read.slide(prs, sid)["shapes"]
                        for p in (sh.get("paragraphs") or [])) for sid in res["new_slide_ids"]]  # fmt: skip
    assert texts[0].strip(" /") == "Contexto" and all(w in texts[1] for w in ("Objetivos", "Um", "Três"))
    assert "Outubro" in texts[2]


def test_a_deck_reaches_nothing_outside_it_when_rendered():
    """Security review H1, reproduced: a picture linked to http://127.0.0.1:8765/ was fetched with GET while
    LibreOffice converted the deck (and a file:// link drew a file from the disk). The converted copy keeps no external
    relationship but its hyperlinks; the deck itself is not changed."""
    from app.docengine import external

    src = io.BytesIO(deck("simple.pptx"))
    out = io.BytesIO()
    with zipfile.ZipFile(src) as z, zipfile.ZipFile(out, "w") as w:
        for info in z.infolist():
            data = z.read(info)
            if info.filename == "ppt/slides/_rels/slide1.xml.rels":
                data = data.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rId90" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"'
                    b' Target="http://127.0.0.1:8765/x.png" TargetMode="External"/>'
                    b'<Relationship Id="rId91" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"'
                    b' Target="https://www.bancoctt.pt/" TargetMode="External"/></Relationships>',
                )
            w.writestr(info, data)
    linked = out.getvalue()
    clean, removed = external.stripped(linked)
    rels = zipfile.ZipFile(io.BytesIO(clean)).read("ppt/slides/_rels/slide1.xml.rels")
    assert removed == 1 and b"127.0.0.1" not in rels and b"https://www.bancoctt.pt/" in rels
    assert external.stripped(deck("simple.pptx")) == (deck("simple.pptx"), 0)  # nothing external: the same bytes


def test_text_alone_is_not_put_on_a_layout_made_for_a_table():
    ctt = (REPO / "templates" / "bancoctt.pptx").read_bytes()
    data, res = ops.apply(ctt, "add_slide", {"layout": "7_Tabela", "content": {"title": "Capa", "subtitle": "Outubro"}})
    # the cover with no photo place, its heading on the red the layout draws
    assert res["instead_of"] == "7_Tabela" and res["layout"] == "2_Capa S/Imagem"
    assert read.outline(read.open_deck(data))[0]["layout"] == "2_Capa S/Imagem"


def test_a_generated_deck_has_a_readable_cover_and_each_slides_points_as_one_list():
    """Rendered and looked at (spec NL-12): on Banco CTT's template the cover went on "2_Capa S/Imagem", a white title on
    a white page, and each slide's points one in each of "3_Texto"'s boxes; on the default one, a list on "Vertical
    Title and Text". Then on "1_Capa C/ Imagem" without its photo, the logo's red "ctt" on the red page (the owner's
    screenshot): the template's image-free cover now draws Banco CTT's own cover artwork (from DevAI v1.1's cover), its
    heading on the red panel, read as such."""
    from app.docengine import layouts

    slides = [{"role": "content", "title": "Regime", "points": ["Primeiro ponto.", "Segundo ponto.", "Terceiro ponto."]},
              {"role": "section", "title": "Módulo", "points": ["O objetivo do módulo."]}]  # fmt: skip
    ctt = read.open_deck((REPO / "templates" / "bancoctt.pptx").read_bytes())
    W, H = ctt.slide_width, ctt.slide_height
    assert [x.name for x in ctt.slide_layouts if not layouts.readable(x, W, H)] == []
    bare = read.open_deck((REPO / "templates" / "bancoctt.pptx").read_bytes())
    plain_cover = next(x for x in bare.slide_layouts if x.name == "2_Capa S/Imagem")
    for sh in [sh for sh in plain_cover.shapes if not sh.is_placeholder]:
        sh._element.getparent().remove(sh._element)
    assert not layouts.readable(plain_cover, W, H)  # without the artwork under it: white on the white page
    made = ops.add_outline(ctt, slides, cover={"title": "Crédito Habitação Jovem", "subtitle": "Formação"})
    cover, content, section = (ctt.slides.get(i) for i in made["slides"])
    assert cover.slide_layout.name == "2_Capa S/Imagem"
    assert len([sh for sh in cover.slide_layout.shapes if not sh.is_placeholder]) == 7  # its artwork, logo, note
    assert content.slide_layout.name == "12_Texto" and section.slide_layout.name == "2_Separador S/Imagem"
    texts = [ph.text_frame.text for ph in content.placeholders if ph.text_frame.text.strip()]
    assert texts == ["Regime", "Primeiro ponto.\nSegundo ponto.\nTerceiro ponto."]  # one list, in one box
    from pptx import Presentation

    plain = Presentation()
    made = ops.add_outline(plain, slides, cover={"title": "Crédito Habitação Jovem"})
    assert [plain.slides.get(i).slide_layout.name for i in made["slides"]] == ["Title Slide", "Title and Content", "Title Slide"]


def test_native_charts_are_made_read_and_changed():
    """Spec NL-5: editable PowerPoint charts (their data in the file) from the person's figures; the read model gives
    the data edit_chart takes (rendered and looked at: a column chart in the layout's chart place, a pie, a line)."""
    ctt = (REPO / "templates" / "bancoctt.pptx").read_bytes()
    data, made = ops.apply(ctt, "add_slide", {"layout": "4_Gráficos", "content": {"title": "Crédito"}})
    sid = made["slides"][0]
    series = [{"name": "2025", "values": [12, 15, 14, 18]}, {"name": "2026", "values": [14, 17, 19, 22]}]
    args = {"slide_id": sid, "kind": "column", "title": "Crédito (M€)", "categories": ["T1", "T2", "T3", "T4"], "series": series}
    data, res = ops.apply(data, "add_chart", args)
    chart = next(sh for sh in get(data, sid)["shapes"] if sh["shape_id"] == res["shape_id"])
    assert chart["type"] == "chart" and chart["chart"] == {
        "kind": "column", "title": "Crédito (M€)", "categories": ["T1", "T2", "T3", "T4"],
        "series": [{"name": "2025", "values": [12.0, 15.0, 14.0, 18.0]}, {"name": "2026", "values": [14.0, 17.0, 19.0, 22.0]}],
    }  # fmt: skip
    graphic = next(sh for sh in read.open_deck(data).slides.get(sid).shapes if sh.shape_id == res["shape_id"])
    assert graphic.is_placeholder  # in the layout's chart place
    data, _ = ops.apply(data, "edit_chart", {"slide_id": sid, "shape_id": res["shape_id"], "categories": ["T1"],
                                             "series": [{"name": "2026", "values": [14]}], "title": ""})  # fmt: skip
    chart = next(sh for sh in get(data, sid)["shapes"] if sh["shape_id"] == res["shape_id"])
    assert chart["chart"]["categories"] == ["T1"] and chart["chart"]["title"] is None
    on = next(sh for sh in read.open_deck(data).slides.get(sid).shapes if sh.shape_id == res["shape_id"])
    assert on.chart.has_legend is False  # one series left: no legend
    data, made = ops.apply(data, "add_slide", {"layout": "8_Texto", "content": {"title": "Distribuição"}})
    pie_slide = made["slides"][0]
    data, pie = ops.apply(data, "add_chart", {"slide_id": pie_slide, "kind": "pie", "categories": ["HPP", "HPS"],
                                              "series": [{"name": "Peso", "values": [70, 30]}]})  # fmt: skip
    shapes = get(data, pie_slide)["shapes"]
    left = [s for s in shapes if s.get("placeholder") and not any(p.get("runs") for p in s.get("paragraphs") or [])]
    assert pie["covers"] == [] and not left  # no empty text box left under it
    for bad in ({"series": [{"name": "x", "values": [1]}]}, {"kind": "pie", "series": series}):  # 1 value for 4; a pie of 2
        with pytest.raises(ops.OpError) as e:
            ops.apply(data, "add_chart", {**args, "slide_id": pie_slide, **bad})
        assert e.value.code == "BAD_CHART"
    text_id = next(sh["shape_id"] for sh in get(data, sid)["shapes"] if sh["type"] != "chart")
    with pytest.raises(ops.OpError) as e:
        ops.apply(data, "edit_chart", {"slide_id": sid, "shape_id": text_id, "title": "x"})
    assert e.value.code == "NOT_A_CHART"


def test_a_diagram_from_a_description_is_laid_out_in_steps_and_styled_like_the_decks():
    """M8. Measured on the Squad Model deck (rendered and looked at): its arrows are free lines, so its boxes are found
    by their look (the most repeated, with connector lines on the slide), the nearest such slide's."""
    ctt = (REPO / "templates" / "bancoctt.pptx").read_bytes()
    nodes = [{"text": t} for t in ("Simulação", "Proposta", "Avaliação", "Aprovação", "Recusa")]
    edges = [{"from": 0, "to": 1}, {"from": 1, "to": 2}, {"from": 2, "to": 3}, {"from": 2, "to": 4}]
    data, made = ops.apply(ctt, "add_slide", {"layout": "8_Texto", "content": {"title": "Processo"}})
    sid = made["slides"][0]
    data, res = ops.apply(data, "draw_diagram", {"slide_id": sid, "nodes": nodes, "edges": edges})
    assert res["styled_from"] == "the theme" and len(res["shape_ids"]) == 5 and len(res["connector_ids"]) == 4
    shapes = {sh["shape_id"]: sh for sh in get(data, sid)["shapes"]}
    x = [shapes[i]["box"]["x"] for i in res["shape_ids"]]
    assert x[0] < x[1] < x[2] < x[3] and x[3] == x[4]  # steps left to right; the branch side by side
    assert all(shapes[c].get("connects") for c in res["connector_ids"])  # attached at both ends
    # the deck's diagram outlined in red: the next slide's diagram copies that look
    prs = read.open_deck(data)
    for i in res["shape_ids"]:
        box = next(x for x in prs.slides.get(sid).shapes if x.shape_id == i)
        ln = etree.SubElement(box._element.find(qn_("p:spPr")), qn_("a:ln"), w="38100")
        etree.SubElement(etree.SubElement(ln, qn_("a:solidFill")), qn_("a:srgbClr"), val="E00024")
    out = io.BytesIO()
    prs.save(out)
    after = {"layout": "8_Texto", "after_slide_id": sid, "content": {"title": "Outro"}}
    data, made = ops.apply(out.getvalue(), "add_slide", after)
    data, res = ops.apply(data, "draw_diagram", {"slide_id": made["slides"][0], "nodes": nodes[:2], "edges": edges[:1]})
    assert res["styled_from"].startswith("shape ")
    new = next(x for x in read.open_deck(data).slides.get(made["slides"][0]).shapes if x.shape_id == res["shape_ids"][0])
    assert b'val="E00024"' in etree.tostring(new._element)  # the deck's red outline, copied
    with pytest.raises(ops.OpError) as e:
        ops.apply(data, "draw_diagram", {"slide_id": sid, "nodes": nodes, "edges": [{"from": 0, "to": 9}]})
    assert e.value.code == "BAD_DIAGRAM"


def test_a_chart_named_with_its_kind_is_unchanged_and_another_kind_is_drawn_anew_in_its_frame():
    """The model sends edit_chart the chart whole, as add_chart takes it (measured: 2 runs in 2 named its kind, the
    same): that kind changes nothing; another redraws it in the same frame, keeping its data, its title and the link to
    its workbook (PowerPoint's Edit Data; measured: the new chart's XML had lost it)."""
    simple = (DECKS / "simple.pptx").read_bytes()
    data, made = ops.apply(simple, "add_slide", {"layout": "Title Only", "content": {"title": "Resultados"}})
    sid = made["slides"][0]
    cats, rows = ["2023", "2024", "2025"], [{"name": "Vendas", "values": [120, 150, 180]}]
    whole = {"slide_id": sid, "kind": "column", "title": "Vendas", "categories": cats, "series": rows}
    data, res = ops.apply(data, "add_chart", whole)
    shid = res["shape_id"]
    chart = lambda d: next(sh for sh in get(d, sid)["shapes"] if sh["shape_id"] == shid)["chart"]  # noqa: E731
    rows[0]["values"][2] = 200
    data, _ = ops.apply(data, "edit_chart", {**whole, "shape_id": shid})
    now = [{"name": "Vendas", "values": [120.0, 150.0, 200.0]}]
    assert chart(data) == {"kind": "column", "title": "Vendas", "categories": cats, "series": now}
    data, _ = ops.apply(data, "edit_chart", {"slide_id": sid, "shape_id": shid, "kind": "bar"})
    assert chart(data) == {"kind": "bar", "title": "Vendas", "categories": cats, "series": now}
    data, _ = ops.apply(data, "edit_chart", {"slide_id": sid, "shape_id": shid, "kind": "line", "title": ""})
    assert chart(data)["kind"] == "line" and chart(data)["title"] is None
    z = zipfile.ZipFile(io.BytesIO(data))
    part = next(n for n in z.namelist() if n.startswith("ppt/charts/chart") and n.endswith(".xml"))
    assert b"<c:externalData" in z.read(part)  # still linked to its workbook
    book = zipfile.ZipFile(io.BytesIO(z.read(next(n for n in z.namelist() if n.startswith("ppt/embeddings/")))))
    sheet = book.read("xl/worksheets/sheet1.xml")
    assert b"<v>200</v>" in sheet and b"<v>180</v>" not in sheet
    with pytest.raises(ops.OpError) as e:  # a pie has one series
        two = [*rows, {"name": "Custos", "values": [1, 2, 3]}]
        ops.apply(data, "edit_chart", {"slide_id": sid, "shape_id": shid, "kind": "pie", "categories": cats, "series": two})
    assert e.value.code == "BAD_CHART"


def test_a_diagram_without_arrows_is_a_process_and_leaves_no_empty_text_box_under_it():
    """Measured: "a diagram of the process: Pedido, Análise, Aprovação, Entrega" came with no edges, and four boxes
    stood joined by nothing; and drawn on a title-and-content slide, its empty text box stayed under the diagram."""
    simple = (DECKS / "simple.pptx").read_bytes()
    data, made = ops.apply(simple, "add_slide", {"layout": "Title and Content", "content": {"title": "Processo"}})
    sid = made["slides"][0]
    nodes = [{"text": t} for t in ("Pedido", "Análise", "Aprovação", "Entrega")]
    data, res = ops.apply(data, "draw_diagram", {"slide_id": sid, "nodes": nodes})
    assert len(res["connector_ids"]) == 3
    shapes = {sh["shape_id"]: sh for sh in get(data, sid)["shapes"]}
    x = [shapes[i]["box"]["x"] for i in res["shape_ids"]]
    assert x == sorted(x) and len(set(x)) == 4  # a step each, left to right
    left = [s for s in shapes.values() if s.get("placeholder") and not any(p.get("runs") for p in s.get("paragraphs") or [])]
    assert not left  # the layout's empty text box under the diagram went
    data, res = ops.apply(data, "draw_diagram", {"slide_id": sid, "nodes": nodes[:2], "edges": []})
    assert res["connector_ids"] == []  # no arrows, when that is what is asked
    assert ops.blank_layout(read.open_deck((DECKS / "simple.pptx").read_bytes())).name == "Blank"
    assert ops.blank_layout(read.open_deck((REPO / "templates" / "bancoctt.pptx").read_bytes())).name == "Descanso"


def test_a_chart_over_a_slides_list_says_it_covers_it():
    """Measured: on "Title and Content" with a point written, the chart was drawn over the list and covers said
    nothing (a box over half the slide was taken for a background)."""
    simple = (DECKS / "simple.pptx").read_bytes()
    listed = {"layout": "Title and Content", "content": {"title": "Vendas", "points": ["Vendas por ano"]}}
    data, made = ops.apply(simple, "add_slide", listed)
    sid = made["slides"][0]
    body = next(sh["shape_id"] for sh in get(data, sid)["shapes"] if (sh.get("placeholder") or {}).get("type") == "object")
    chart = {"slide_id": sid, "kind": "column", "categories": ["2023", "2024"], "series": [{"name": "Vendas", "values": [1, 2]}]}
    _, res = ops.apply(data, "add_chart", chart)
    assert res["covers"] == [body]


def test_a_new_slide_with_a_chart_or_a_diagram_has_it_under_its_title_in_place_of_its_list():
    """Measured: for "a new slide with a diagram of the process", the model put the boxes in add_slide first (8 turns in
    8); for a chart, it wrote a filler point and drew the chart over it (4 runs in 4). The figure is content: drawn where
    the list would go, nothing left under it, a filler point left out and said."""
    simple = (DECKS / "simple.pptx").read_bytes()
    nodes = [{"text": t} for t in ("Pedido", "Análise", "Aprovação", "Entrega")]
    content = {"title": "Processo de crédito", "diagram": {"nodes": nodes}}
    data, res = ops.apply(simple, "add_slide", {"layout": "Blank", "content": content})
    sid = res["slides"][0]
    assert res["figure"] == "diagram" and len(res["connector_ids"]) == 3 and res["instead_of"] == "Blank"
    shapes = get(data, sid)["shapes"]
    title = next(sh for sh in shapes if (sh.get("placeholder") or {}).get("type") == "title")
    assert title["paragraphs"][0]["runs"][0]["text"] == "Processo de crédito"
    assert not [sh for sh in shapes if sh.get("placeholder") and not any(p.get("runs") for p in sh.get("paragraphs") or [])]
    boxes = [sh for sh in shapes if sh["shape_id"] in res["shape_ids"]]
    assert all(b["box"]["y"] >= title["box"]["y"] + title["box"]["h"] - 0.01 for b in boxes)  # under the title

    chart = {"kind": "column", "categories": ["2023", "2024", "2025"], "series": [{"name": "Vendas", "values": [120, 150, 180]}]}
    content = {"title": "Vendas anuais", "points": ["Vendas por ano"], "chart": chart}
    data, res = ops.apply(simple, "add_slide", {"layout": "Title and Content", "content": content})
    sid = res["slides"][0]
    assert res["figure"] == "chart" and res["left_out"] == ["points (the chart takes their place)"] and "covers" not in res
    shapes = get(data, sid)["shapes"]
    made = next(sh for sh in shapes if sh["shape_id"] == res["shape_id"])
    assert made["chart"]["categories"] == ["2023", "2024", "2025"] and made["chart"]["title"] is None
    assert [sh["type"] for sh in shapes] == ["placeholder", "chart"]  # the title and the chart: nothing under it

    ctt = (REPO / "templates" / "bancoctt.pptx").read_bytes()  # a layout with a chart place: the chart is put in it
    data, res = ops.apply(ctt, "add_slide", {"layout": "4_Gráficos", "content": {"title": "Crédito", "chart": chart}})
    graphic = next(x for x in read.open_deck(data).slides.get(res["slides"][0]).shapes if x.shape_id == res["shape_id"])
    assert graphic.is_placeholder
    shapes = get(data, res["slides"][0])["shapes"]
    texts = [sh for sh in shapes if sh.get("placeholder") and "paragraphs" in sh]
    empty = [sh for sh in texts if not any(p.get("runs") for p in sh["paragraphs"])]
    assert not empty
    # the chart sent whole to edit_chart, the slide's title as its own (measured): said once, by the slide
    data, _ = ops.apply(data, "edit_chart", {"slide_id": res["slides"][0], "shape_id": res["shape_id"], "title": "crédito "})
    assert next(sh for sh in get(data, res["slides"][0])["shapes"] if sh["shape_id"] == res["shape_id"])["chart"]["title"] is None


def test_a_slide_with_a_subtitle_and_points_goes_on_a_layout_with_both():
    """Measured: a cover sent with a subtitle and one point went on a layout with no subtitle place; the subtitle was
    left out (and the reply claimed it)."""
    ctt = (REPO / "templates" / "bancoctt.pptx").read_bytes()
    content = {"title": "Apresentação", "subtitle": "Crédito Habitação Jovem", "points": ["Para a equipa comercial"]}
    chosen = ops._layout_for_content(read.open_deck(ctt), content)  # as fill_slide on a slide not there yet chooses
    data, res = ops.apply(ctt, "add_slide", {"layout": chosen.name, "content": content, "position": 0})
    sid = (res["slides"] if isinstance(res, dict) else res)[0]
    assert not (isinstance(res, dict) and res.get("left_out"))
    text = [p["runs"][0]["text"] for sh in get(data, sid)["shapes"] for p in sh.get("paragraphs") or [] if p.get("runs")]
    assert {"Apresentação", "Crédito Habitação Jovem", "Para a equipa comercial"} <= set(text)


def test_points_sent_with_a_layout_that_has_no_place_for_them_go_on_one_that_has():
    """As the model sent it: an index for an empty cover, with the cover's own layout named (refused NO_PLACE, then the
    title alone written, 2 runs in 3)."""
    ctt = (REPO / "templates" / "bancoctt.pptx").read_bytes()
    data, made = ops.apply(ctt, "add_slide", {"layout": "1_Capa C/ Imagem"})
    sid = (made["slides"] if isinstance(made, dict) else made)[0]
    topics = ["Regime e finalidades", "Fases do processo", "Servicing", "Crédito obras"]
    index = {"title": "Índice", "points": topics}
    data, res = ops.apply(data, "fill_slide", {"slide_id": sid, "layout": "1_Capa C/ Imagem", "content": index})
    lines = [p["runs"][0]["text"] for sh in get(data, sid)["shapes"] for p in sh.get("paragraphs") or [] if p.get("runs")]
    assert res.get("layout") and res["layout"] != "1_Capa C/ Imagem" and set(topics) <= set(lines)
    data, res = ops.apply(ctt, "add_slide", {"layout": "1_Capa C/ Imagem", "content": {"title": "Índice", "points": topics}})
    assert res["instead_of"] == "1_Capa C/ Imagem"


def test_a_list_too_long_for_its_slide_goes_on_over_as_many_as_it_needs():
    """Measured: ten criteria from the knowledge base (1183 characters) ran past the slide's bottom, 3 runs in 3."""
    from app.docengine import textfit

    simple = (DECKS / "simple.pptx").read_bytes()
    points = [f"Critério {n}: o(s) mutuário(s) do contrato cumprem a condição número {n} do regime, como a portaria "
              f"o descreve para o crédito à habitação, com prazos e valores definidos" for n in range(1, 11)]
    content = {"title": "Condições", "points": points}
    data, res = ops.apply(simple, "add_slide", {"layout": "Title and Content", "content": content})
    assert res["continued"] >= 1 and len(res["slides"]) == res["continued"] + 1
    prs = read.open_deck(data)
    said = []
    for sid in res["slides"]:
        s = prs.slides.get(sid)
        assert s.shapes.title.text == "Condições"
        assert not any(ph.has_text_frame and textfit.overflows(ph) for ph in s.placeholders)
        said += [p.text for ph in s.placeholders if ph != s.shapes.title for p in ph.text_frame.paragraphs if p.text]
    assert said == points  # every point, in order
    assert len(prs.slides) == 3 + len(res["slides"])


def test_a_covers_short_points_are_its_subtitles_lines():
    """Measured (cover-for-a-new-deck): the model filled a new deck's cover with the audience and the date as points,
    and the cover moved to "4_Texto" / "6_Texto", its lines scattered in boxes, 2 runs in 2. On a deck's first slide on
    a cover they are the subtitle's lines; a point repeating the title is left out."""
    ctt = read.open_deck((REPO / "templates" / "bancoctt.pptx").read_bytes())
    cover = ops.cover_layout(ctt).name
    first = ops._ids(ops.add_slide(ctt, cover))[0]
    second = ops._ids(ops.add_slide(ctt, cover))[0]
    sent = {"points": [{"text": "Crédito Habitação Jovem"}, {"text": "Para a Equipa Comercial"}, {"text": "Outubro de 2026"}],
            "subtitle": "Crédito Habitação Jovem", "title": "Crédito Habitação Jovem"}  # fmt: skip
    got = ops.fill_slide(ctt, first, sent)
    assert "layout" not in got and ctt.slides.get(first).slide_layout.name == cover
    one = read.slide(ctt, first)
    texts = {sh["placeholder"]["idx"]: ["".join(r["text"] for r in p["runs"]) for p in sh["paragraphs"]]
             for sh in one["shapes"] if sh.get("paragraphs")}  # fmt: skip
    from app.docengine import layouts

    where = layouts.slots(ctt.slides.get(first).slide_layout, ctt.slide_width, ctt.slide_height)
    assert texts[where["heading"]] == ["Crédito Habitação Jovem"]
    assert texts[where["subtitle"]] == ["Para a Equipa Comercial", "Outubro de 2026"]  # its lines, the title's echo out
    moved = ops.fill_slide(ctt, second, sent)  # not the first slide: a list, on a layout with its boxes
    assert moved.get("layout") and moved["layout"] != cover


def test_a_new_first_slide_with_a_covers_content_goes_on_the_cover():
    """Measured (cover-for-a-new-deck, after the subtitle fold): on an empty deck the slide was made from its content,
    which chose "6_Texto" for a title and three points before the cover was ever considered, 2 runs in 2."""
    ctt = read.open_deck((REPO / "templates" / "bancoctt.pptx").read_bytes())
    sent = {"title": "Crédito Habitação Jovem", "points": ["Para a Equipa Comercial", "Outubro de 2026"]}
    assert ops.layout_for_new(ctt, sent, first=True).name == ops.cover_layout(ctt).name
    assert ops.layout_for_new(ctt, sent, first=False).name != ops.cover_layout(ctt).name  # elsewhere: its points' layout
    long = {"title": "Índice", "points": [f"Tema {n}" for n in range(5)]}  # an index is no cover, even first
    assert ops.layout_for_new(ctt, long, first=True).name != ops.cover_layout(ctt).name


def test_a_cover_whose_lines_overflow_stays_one_slide():
    """Measured (cover-for-a-new-deck, 1 run in 2): a subtitle and three points, folded into four subtitle lines,
    overflowed the box; add_slide went on over a second slide and the date landed on "4_Texto"."""
    ctt = read.open_deck((REPO / "templates" / "bancoctt.pptx").read_bytes())
    sent = {"points": ["Apresentação sobre o Crédito Habitação Jovem", "Para: Equipa Comercial", "Outubro de 2026"],
            "subtitle": "Acesso facilitado à habitação própria permanente para jovens.",
            "title": "Crédito Habitação Jovem com Garantia do Estado"}  # fmt: skip
    lay = ops.layout_for_new(ctt, sent, first=True)
    made = ops._ids(ops.add_slide(ctt, lay.name, content=sent))
    assert len(made) == 1 and len(ctt.slides) == 1 and ctt.slides[0].slide_layout.name == ops.cover_layout(ctt).name
    text = "\n".join(ph.text_frame.text for ph in ctt.slides[0].placeholders)
    assert "Outubro de 2026" in text and "Para: Equipa Comercial" in text
    # the four lines did not fit even at 14 pt (measured: CANNOT_FIT, and the model gave up): the model's own subtitle
    # sentence went to the notes, the cover's short lines stay on it, nothing past its box
    assert "Acesso facilitado" not in text and "Acesso facilitado" in ctt.slides[0].notes_slide.notes_text_frame.text
    assert not ops._overflowing(ctt, made[0])
    again = read.open_deck((REPO / "templates" / "bancoctt.pptx").read_bytes())  # the same through fill_slide
    sid = ops._ids(ops.add_slide(again, ops.cover_layout(again).name))[0]
    got = ops.fill_slide(again, sid, sent)
    assert got.get("moved_to_notes") == "subtitle" and not ops._overflowing(again, sid) and len(again.slides) == 1
    assert "Outubro de 2026" in "\n".join(ph.text_frame.text for ph in again.slides[0].placeholders)


def test_a_list_goes_one_point_a_box_only_over_boxes_alike():
    """Looked at (first-use-index-from-kb, 1 run in 4): an index's four entries spread over "2_Texto"'s bold banner,
    large body and two small columns read as fragments; over "1_Texto"'s five equal rows they read as an index."""
    ctt = read.open_deck((REPO / "templates" / "bancoctt.pptx").read_bytes())
    index = {"title": "Índice", "points": ["Oferta", "Fases do processo", "Servicing", "Regime bonificado"]}
    assert ops._layout_for_content(ctt, index).name != "2_Texto"
    sid = ops._ids(ops.add_slide(ctt, "2_Texto", content=index))[0]
    one = read.slide(ctt, sid)
    filled = [sh for sh in one["shapes"] if sh.get("paragraphs") and sh["placeholder"]["idx"] != 19]  # 19: its heading
    assert len(filled) == 1 and len(filled[0]["paragraphs"]) == 4  # one list, in its large body
    rows = ops._ids(ops.add_slide(ctt, "1_Texto", content=index))[0]  # rows alike: one entry a row, as designed
    assert len([sh for sh in read.slide(ctt, rows)["shapes"] if sh.get("paragraphs")]) == 5  # the heading and four rows


def test_fill_slide_keeps_the_slides_title_when_none_is_sent():
    data = deck("simple.pptx")
    prs = read.open_deck(data)
    second = read.outline(prs)[1]
    ops.fill_slide(prs, second["slide_id"], {"points": ["Um", "Dois"]})
    assert read.outline(prs)[1]["title"] == second["title"]
    ops.fill_slide(prs, second["slide_id"], {"points": ["Um", "Dois", "Três"]}, layout="Two Content")  # relaid too
    assert read.outline(prs)[1]["title"] == second["title"]


def test_a_list_avoids_a_layout_whose_picture_place_it_would_leave_empty():
    """Looked at (first-use-index-from-kb, 2 runs in 8): an index on "1_Agenda C/Imagem", its picture place a grey block;
    "2_Agenda S/Imagem" is the same without it."""
    ctt = read.open_deck((REPO / "templates" / "bancoctt.pptx").read_bytes())
    for n in range(2, 9):
        lay = ops._layout_for_content(ctt, {"title": "Índice", "points": [f"Tema {i}" for i in range(n)]})
        roles = {x["role"] for x in layouts_of(ctt, lay)}
        assert "picture" not in roles, (n, lay.name)


def layouts_of(prs, lay):
    from app.docengine import layouts

    return layouts.placeholders(lay, prs.slide_width, prs.slide_height)


def test_a_generated_slides_list_is_bulleted_at_the_templates_size():
    """The user (2026-10-08): "add bullets and keep the font size" for generated decks' lists (Banco CTT's body text has
    no bullets of its own)."""
    ctt = read.open_deck((REPO / "templates" / "bancoctt.pptx").read_bytes())
    made = ops.add_outline(ctt, [{"role": "content", "title": "Regime", "points": ["Primeiro.", "Segundo.", "Terceiro."]}])
    s = ctt.slides.get(made["slides"][0])
    body = next(ph for ph in s.placeholders if ph.text_frame.text.startswith("Primeiro"))
    paras = body.text_frame._txBody.findall(qn_("a:p"))
    assert len(paras) == 3
    for p in paras:
        ppr = p.find(qn_("a:pPr"))
        tags = [etree.QName(c).localname for c in ppr]
        assert ppr.find(qn_("a:buChar")).get("char") == "•" and "buNone" not in tags
        assert tags.index("buChar") < tags.index("defRPr") if "defRPr" in tags else True  # schema order
    data = ops.save(ctt)
    assert read.open_deck(data).slides.get(made["slides"][0])  # it reopens


def test_a_rounded_boxs_text_is_measured_inside_its_corners():
    """Looked at (a generated training's diagram): "Processamento", 114.3 pt, judged to fit a node 121.4 pt wide inside
    its insets, broken by the renderer as "Processamen" / "to": the rounded corners take 5.3 pt a side."""
    from app.docengine import textfit

    ctt = read.open_deck((REPO / "templates" / "bancoctt.pptx").read_bytes())
    item = {"role": "content", "title": "Processo de Financiamento para Obras e Construção",
            "points": ["O processo de desembolso segue 4 passos: Pedido, Avaliação, Processamento e Comunicação."],
            "design": {"form": "diagram", "diagram": {"nodes": ["Pedido", "Avaliação", "Processamento", "Comunicação"],
                                                      "direction": "right"}}}  # fmt: skip
    s = ctt.slides.get(ops.add_outline(ctt, [item])["slides"][0])
    node = next(sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text == "Processamento")
    own = node.text_frame._txBody.find(qn_("a:bodyPr"))
    sides = sum(int(own.get(k) or textfit.DEFAULT_INS[k]) for k in ("lIns", "rIns"))
    corners = 2 * min(node.width, node.height) * 16667 / 100000 * 0.29289  # DrawingML's roundRect text rectangle
    inside = (node.width - sides - corners) / textfit.EMU_PT
    size = textfit.smallest_size(node)  # at its scale
    major, minor = textfit._theme_fonts(node)
    word = textfit._width("Processamento", (size, "+mn-lt", False, False, 0.0), major, minor)
    assert word <= inside * textfit.WORD_ROOM, (word, inside)  # its size fitted to the text inside its corners
    # written into its runs, as every renderer then shows it (LibreOffice drew a box's fontScale at full size when the
    # text's height fitted), and every box of the diagram at that one size
    boxes = [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text in item["design"]["diagram"]["nodes"]]
    assert len(boxes) == 4 and len({textfit.smallest_size(sh) for sh in boxes}) == 1
    for sh in boxes:
        fit = sh.text_frame._txBody.find(qn_("a:bodyPr")).find(qn_("a:normAutofit"))
        assert fit is None or fit.get("fontScale") is None
        assert all(r.find(qn_("a:rPr")).get("sz") for r in sh.text_frame._txBody.iter(qn_("a:r")))


def test_a_word_wider_than_its_box_is_fitted_not_broken():
    """Looked at (a generated training's section dividers): "Complementos" and "Documentação" set as "Complemento" / "s"
    in "2_Separador S/Imagem"'s narrow title; the box's height held the lines, so the text was never fitted."""
    from app.docengine import textfit

    ctt = read.open_deck((REPO / "templates" / "bancoctt.pptx").read_bytes())
    sections = [{"role": "section", "title": "Casos Específicos e Complementos", "points": ["Abordar casos."]},
                {"role": "section", "title": "Processo de Candidatura e Documentação", "points": ["Fases."]}]  # fmt: skip
    made = ops.add_outline(ctt, sections)
    for sid in made["slides"]:
        s = ctt.slides.get(sid)
        title = next(ph for ph in s.placeholders if ph.placeholder_format.idx == 0)
        m = textfit.measure(title)
        assert not m["broken"] and not textfit.overflows(title)
        assert m["font_scale"] < 1  # fitted: made smaller so its longest word fits the line
