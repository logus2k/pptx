"""The assistant (M3) through the real server and socket.io, with the scripted model: spec scenarios S1, S5, S6, S7,
S10, plus waiting, failure and limit paths. What is checked is what lands on disk and what the model was sent."""

from __future__ import annotations

import io
import json
import threading
import time

import pytest
import requests
import socketio
from PIL import Image

from app.docengine import read

from .conftest import REPO, SECRET

DECKS = REPO / "fixtures" / "decks"
ANA = "ana@example.com"


def h(email=ANA):
    return {"X-Slides-Proxy-Secret": SECRET, "X-Auth-Request-Email": email}


class Chat:
    """A socket.io client following one conversation; it collects the events and waits for a turn's end."""

    def __init__(self, server: str, pid: str, cid: str, email: str = ANA) -> None:
        self.pid, self.cid = pid, cid
        self.events: list[tuple[str, dict]] = []
        self.ended = threading.Event()
        self.sio = socketio.Client()
        for name in (
            "message",
            "assistant_delta",
            "assistant_message",
            "tool_progress",
            "question",
            "plan",
            "instructions_proposed",
            "instructions_changed",
            "proposal_updated",
            "deck_changed",
            "turn_started",
            "turn_ended",
            "assistant_error",
        ):
            self.sio.on(name, self._handler(name))
        self.sio.connect(server, headers=h(email), transports=["websocket"])
        joined = self.sio.call("join_conversation", {"project_id": pid, "conversation_id": cid}, timeout=10)
        assert joined["ok"], joined

    def _handler(self, name):
        def on(data):
            self.events.append((name, data))
            if name == "turn_ended":
                self.ended.set()

        return on

    def send(self, text: str, **extra) -> dict:
        self.ended.clear()
        r = self.sio.call(
            "user_message", {"project_id": self.pid, "conversation_id": self.cid, "text": text, **extra}, timeout=10
        )
        assert r["ok"], r
        assert self.ended.wait(30), "the turn did not end"
        return self.last("turn_ended")

    def answer(self, **decision) -> dict:
        self.ended.clear()
        r = self.sio.call("answer", {"project_id": self.pid, "conversation_id": self.cid, **decision}, timeout=10)
        assert r["ok"], r
        assert self.ended.wait(30), "the turn did not end"
        return self.last("turn_ended")

    def decide(self, proposal_id: str, accept: bool, slides=None) -> dict:
        r = self.sio.call(
            "proposal_decision",
            {"project_id": self.pid, "conversation_id": self.cid, "proposal_id": proposal_id, "accept": accept, "slides": slides},
            timeout=30,
        )
        assert r["ok"], r
        return r

    def last(self, name: str) -> dict:
        return next(d for n, d in reversed(self.events) if n == name)

    def close(self) -> None:
        self.sio.disconnect()


def setup(server, deck="simple.pptx"):
    pid = requests.post(f"{server}/api/projects", json={"name": "P"}, headers=h(), timeout=10).json()["id"]
    did = None
    if deck:
        with open(DECKS / deck, "rb") as f:
            did = requests.post(f"{server}/api/projects/{pid}/decks", files={"file": (deck, f)}, headers=h(), timeout=60).json()[
                "id"
            ]
    cid = requests.post(f"{server}/api/projects/{pid}/conversations", json={"deck_id": did}, headers=h(), timeout=10).json()["id"]
    return pid, did, cid


def deck_bytes(server, pid, did, version=None) -> bytes:
    params = {"version": version} if version else {}
    return requests.get(f"{server}/api/projects/{pid}/decks/{did}/download", params=params, headers=h(), timeout=30).content


def outline(data: bytes) -> list[dict]:
    return read.outline(read.open_deck(data))


def body_id(data: bytes, sid: int) -> int:
    return next(
        s["shape_id"] for s in read.slide(read.open_deck(data), sid)["shapes"] if s.get("placeholder", {}).get("idx") == 1
    )


def test_s1_plan_then_per_slide_acceptance(server, fake_model):
    """S1: tighten several slides; the plan is approved; the person accepts some slides and rejects one."""
    pid, did, cid = setup(server, "long-200.pptx")
    data = deck_bytes(server, pid, did)
    ids = [o["slide_id"] for o in outline(data)][3:7]  # "slides 4 to 7"
    fake_model.script = [{"tools": [("propose_plan", {"steps": [f"Tighten slide {i + 4}" for i in range(4)]})]}]
    chat = Chat(server, pid, cid)
    assert chat.send("tighten all the bullet points on slides 4 to 7, keep the meaning")["status"] == "waiting"
    assert chat.last("plan")["steps"][0] == "Tighten slide 4"
    fake_model.script = [
        {
            "tools": [
                (
                    "update_text",
                    {"slide_id": sid, "shape_id": body_id(data, sid), "paragraphs": [{"runs": [{"text": f"Curto {n}"}]}]},
                )
                for n, sid in enumerate(ids)
            ]
        },
        {"text": "I tightened slides 4 to 7. Review them before accepting."},
    ]
    assert chat.answer(approve=True)["status"] == "done"
    proposal = chat.last("proposal_updated")
    assert proposal["status"] == "pending"
    assert [s["slide_id"] for s in proposal["decks"][did]["slides"]] == ids and all(
        s["state"] == "changed" for s in proposal["decks"][did]["slides"]
    )
    # the saved deck has not changed: nothing changes until accepted (NL-9)
    assert deck_bytes(server, pid, did) == data
    chat.decide(proposal["id"], True, {did: ids[:3]})
    after = deck_bytes(server, pid, did)
    texts = {o["slide_id"]: o["summary"] for o in outline(after)}
    assert [texts[i] for i in ids[:3]] == ["Curto 0", "Curto 1", "Curto 2"]
    assert texts[ids[3]] == f"Conteúdo do diapositivo {7} / Segunda linha"  # the rejected slide is as it was
    versions = requests.get(f"{server}/api/projects/{pid}/decks/{did}/versions", headers=h(), timeout=10).json()
    assert versions["current_version"] == 2 and versions["versions"][0]["source"] == "assistant"
    assert chat.last("proposal_updated")["status"] == "partially_accepted"
    # the model saw the deck's outline and the conversation (the context), and the plan's approval as the tool's result
    sent = fake_model.sent[-1]
    assert "Active deck, slide by slide" in sent[0]["content"] and "<data>" in sent[0]["content"]
    assert any(m["role"] == "tool" and '"approved": true' in m["content"] for m in sent)
    chat.close()


def _png(colour) -> bytes:
    b = io.BytesIO()
    Image.new("RGB", (800, 500), colour).save(b, "PNG")
    return b.getvalue()


def test_s5_replace_a_photo_with_an_uploaded_one_and_alt_text(server, fake_model):
    pid, did, cid = setup(server, "features.pptx")
    asset = requests.post(
        f"{server}/api/projects/{pid}/assets",
        data={"kind": "image"},
        files={"file": ("foto.png", _png((200, 30, 30)))},
        headers=h(),
        timeout=30,
    ).json()
    data = deck_bytes(server, pid, did)
    sid = outline(data)[1]["slide_id"]
    pic = next(s for s in read.slide(read.open_deck(data), sid)["shapes"] if s["type"] == "picture")
    fake_model.script = [
        {
            "tools": [
                ("replace_image", {"slide_id": sid, "shape_id": pic["shape_id"], "image": {"asset_id": asset["id"]}}),
                ("set_alt_text", {"slide_id": sid, "shape_id": pic["shape_id"], "text": "A red photo."}),
            ]
        },
        {"text": "I replaced the image on slide 2 and wrote its alt text."},
    ]
    chat = Chat(server, pid, cid)
    chat.send(
        "put this on slide 2 instead of the stock image, and write alt text for it",
        attachments=[asset["id"]],
        selection={"slide_id": sid},
    )
    # the person's attachment reached the model by its asset id
    first = next(
        sent for sent in fake_model.sent if sent[0]["role"] == "system"
    )  # the assistant's call (the router's comes first)
    assert asset["id"] in next(m["content"] for m in first if m["role"] == "user")
    p = chat.last("proposal_updated")
    chat.decide(p["id"], True)
    after = read.slide(read.open_deck(deck_bytes(server, pid, did)), sid)
    new = next(s for s in after["shapes"] if s["shape_id"] == pic["shape_id"])
    assert new["alt_text"] == "A red photo." and new["emu"] == pic["emu"]
    chat.close()


def test_s6_slides_from_a_dictated_outline(server, fake_model):
    pid, _, cid = setup(server, deck=None)
    points = ["Contexto", "Objetivos", "Solução", "Arquitetura", "Calendário", "Equipa", "Preços", "Próximos passos"]
    fake_model.script = [
        {"tools": [("create_deck", {"title": "Proposta"})]},
        {"tools": [("propose_plan", {"steps": [f"Slide: {p}" for p in points]})]},
    ]
    chat = Chat(server, pid, cid)
    assert chat.send("create eight slides: " + ", ".join(points))["status"] == "waiting"
    fake_model.script = [
        {
            "tools": [
                (
                    "add_slide",
                    {
                        "layout": "Title and Content",
                        "position": i,
                        "placeholders": [{"type": "title", "text": p}, {"idx": 1, "text": f"Sobre {p.lower()}"}],
                    },
                )
                for i, p in enumerate(points)
            ]
        },
        {"text": "I created eight slides from your outline."},
    ]
    chat.answer(approve=True)
    p = chat.last("proposal_updated")
    did = next(iter(p["decks"]))
    assert len(p["decks"][did]["slides"]) == 8 and all(s["state"] == "added" for s in p["decks"][did]["slides"])
    chat.decide(p["id"], True)
    assert [o["title"] for o in outline(deck_bytes(server, pid, did))] == points
    chat.close()


def test_s7_undo_the_last_two_changes_then_download(server, fake_model):
    pid, did, cid = setup(server)
    data = deck_bytes(server, pid, did)
    first = outline(data)[0]["slide_id"]
    chat = Chat(server, pid, cid)
    for text in ("one", "two"):
        fake_model.script = [{"tools": [("set_notes", {"slide_id": first, "text": text})]}, {"text": "Done."}]
        chat.send(f"notes {text}")
        chat.decide(chat.last("proposal_updated")["id"], True)
    fake_model.script = [{"tools": [("undo", {"steps": 2})]}, {"text": "I undid the last two changes."}]
    chat.send("undo the last two changes")
    final = deck_bytes(server, pid, did)
    assert read.slide(read.open_deck(final), first)["notes"] == read.slide(read.open_deck(data), first)["notes"]
    versions = requests.get(f"{server}/api/projects/{pid}/decks/{did}/versions", headers=h(), timeout=10).json()
    assert versions["current_version"] == 5 and [v["source"] for v in versions["versions"][:2]] == ["undo", "undo"]
    assert any(n == "deck_changed" for n, _ in chat.events)
    chat.close()


def test_s10_project_instructions_are_proposed_confirmed_and_used(server, fake_model):
    pid, _, cid = setup(server)
    rule = "Always call the product 'Atlas Platform', never 'Atlas'."
    fake_model.script = [{"tools": [("update_instructions", {"text": rule})]}]
    chat = Chat(server, pid, cid)
    assert chat.send("in this project, always call the product 'Atlas Platform', never 'Atlas'")["status"] == "waiting"
    assert chat.last("instructions_proposed")["text"] == rule
    assert (
        requests.get(f"{server}/api/projects/{pid}", headers=h(), timeout=10).json()["instructions"] == ""
    )  # not before confirming
    fake_model.script = [{"text": "Noted: I will always say Atlas Platform."}]
    chat.answer(accept=True)
    assert requests.get(f"{server}/api/projects/{pid}", headers=h(), timeout=10).json()["instructions"] == rule
    # a later conversation: the rule is in the model's context
    cid2 = requests.post(f"{server}/api/projects/{pid}/conversations", json={}, headers=h(), timeout=10).json()["id"]
    chat2 = Chat(server, pid, cid2)
    fake_model.script = [{"text": "Hello."}]
    chat2.send("hello")
    assert rule in fake_model.sent[-1][0]["content"]
    chat.close()
    chat2.close()


def test_a_question_waits_and_the_answer_becomes_the_tool_result(server, fake_model):
    pid, _, cid = setup(server)
    fake_model.script = [{"tools": [("ask_user", {"question": "Which SLA numbers?", "options": ["Support", "Uptime"]})]}]
    chat = Chat(server, pid, cid)
    assert chat.send("update the SLA numbers")["status"] == "waiting"
    assert chat.last("question")["options"] == ["Support", "Uptime"]
    fake_model.script = [{"text": "Thanks, using the uptime SLA."}]
    chat.send("Uptime")  # typed instead of clicked: it answers the question
    assert any(m["role"] == "tool" and "Uptime" in m["content"] for m in fake_model.sent[-1])
    chat.close()



def test_an_unclear_request_answered_in_words_becomes_a_question(server, fake_model):
    pid, _, cid = setup(server)
    fake_model.routes = [{"intent": "Change something", "kind": "unclear", "groups": []}]
    fake_model.script = [{"text": "What would you like to change, and on which slide?"}]  # no tool called
    chat = Chat(server, pid, cid)
    assert chat.send("Change it.")["status"] == "waiting"
    assert chat.last("question")["question"] == "What would you like to change, and on which slide?"
    assert "Now\nThe request is too vague" in fake_model.sent[-1][0]["content"]
    assert fake_model.offered[-1] == ["ask_user"]  # too vague: asking is the only step offered
    chat.close()

def test_bad_arguments_and_locked_shapes_come_back_as_tool_errors(server, fake_model):
    pid, did, cid = setup(server, "fidelity.pptx")
    data = deck_bytes(server, pid, did)
    sid = outline(data)[0]["slide_id"]
    chart = next(s for s in read.slide(read.open_deck(data), sid)["shapes"] if s["type"] == "chart")
    fake_model.script = [
        {
            "tools": [
                ("update_text", {"slide_id": sid}),
                ("update_text", {"slide_id": sid, "shape_id": chart["shape_id"], "paragraphs": [{"runs": [{"text": "x"}]}]}),
            ]
        },
        {"text": "I cannot change that chart."},
    ]
    chat = Chat(server, pid, cid)
    chat.send("change the chart's text")
    results = [m["content"] for m in fake_model.sent[-1] if m["role"] == "tool"]
    assert "BAD_ARGUMENTS" in results[0] and "LOCKED_ELEMENT" in results[1]
    assert not any(n == "proposal_updated" for n, _ in chat.events)  # nothing was changed: no proposal
    chat.close()


def test_a_turn_stops_at_the_tool_call_limit(server, fake_model):
    pid, _, cid = setup(server)
    fake_model.script = [{"tools": [("get_deck_outline", {})] * 5}] * 9
    chat = Chat(server, pid, cid)
    assert chat.send("loop")["status"] == "limit"
    assert "more steps than one turn allows" in chat.last("assistant_message")["content"]
    chat.close()


def test_a_waiting_proposal_blocks_new_edits_until_decided(server, fake_model):
    pid, did, cid = setup(server)
    first = outline(deck_bytes(server, pid, did))[0]["slide_id"]
    fake_model.script = [{"tools": [("set_notes", {"slide_id": first, "text": "a"})]}, {"text": "Done."}]
    chat = Chat(server, pid, cid)
    chat.send("notes a")
    fake_model.script = [{"tools": [("set_notes", {"slide_id": first, "text": "b"})]}, {"text": "Settle the first."}]
    chat.send("notes b")
    assert "PROPOSAL_WAITING" in [m for m in fake_model.sent[-1] if m["role"] == "tool"][-1]["content"]
    chat.close()


def test_the_turn_streams_and_ends_once(server, fake_model):
    """A first answer without a tool is held back and the model asked once more, with the reason (it may have said it
    changed the deck without doing it); only that second answer is streamed and kept."""
    from app.agent.loop import NUDGE

    pid, _, cid = setup(server)
    fake_model.routes = [{"intent": "A greeting", "kind": "question", "groups": []}]  # as the router reads "olá"
    fake_model.script = [{"text": "Mudei o título."}, {"text": "Olá! Como posso ajudar?"}]
    chat = Chat(server, pid, cid)
    chat.send("olá")
    time.sleep(0.2)
    assert [n for n, _ in chat.events].count("turn_ended") == 1
    assert chat.last("assistant_message")["speakable"] == "Olá! Como posso ajudar?"
    assert "".join(d["text"] for n, d in chat.events if n == "assistant_delta") == "Olá! Como posso ajudar?"
    retry = fake_model.sent[-1]
    assert retry[-2] == {"role": "assistant", "content": "Mudei o título."} and retry[-1]["content"] == NUDGE
    stored = requests.get(f"{server}/api/projects/{pid}/conversations/{cid}", headers=h(), timeout=10).json()
    kept = json.dumps(stored["messages"], ensure_ascii=False)
    assert "Olá! Como posso ajudar?" in kept and "Mudei o título." not in kept  # the held answer never joins it
    chat.close()


def test_a_long_message_is_refused_never_cut(server, fake_model):
    pid, _, cid = setup(server, deck=None)
    chat = Chat(server, pid, cid)
    try:
        r = chat.sio.call("user_message", {"project_id": pid, "conversation_id": cid, "text": "x" * 8001}, timeout=10)
        assert not r["ok"] and "8001 characters" in r["error"]
        assert fake_model.sent == []  # the model was never called
    finally:
        chat.close()


def test_a_conversation_is_untitled_until_its_first_message_names_it(server, fake_model):
    pid, _, cid = setup(server, deck=None)
    listed = lambda: requests.get(f"{server}/api/projects/{pid}/conversations", headers=h(), timeout=10).json()["conversations"]  # noqa: E731
    assert [c["title"] for c in listed()] == [""]  # the page shows its own, translated, label
    fake_model.script = [{"text": "Done."}]
    chat = Chat(server, pid, cid)
    try:
        chat.send("  Make   a deck about Q3 results  ")
    finally:
        chat.close()
    assert [c["title"] for c in listed()] == ["Make a deck about Q3 results"]


def _summariser(window: int):
    """The agent's _summarise alone, over the scripted model with a small window."""
    from types import SimpleNamespace

    from app.agent.loop import Agent

    from .fakes import FakeModels

    models = FakeModels(window=window)
    agent = Agent.__new__(Agent)
    agent.app = SimpleNamespace(models=models)
    return agent, models


def _prompts(models) -> list[str]:
    return [m[0]["content"] for m in models.sent]


def test_the_summary_gets_every_line_whole_in_parts_that_fit():
    import asyncio

    agent, models = _summariser(window=1000)  # a part holds about 600 tokens (1500 characters)
    lines = [f'(tool get_slide: {"{"}"slide_id": {i}, "text": "{"w" * 600}", "end": "END{i}"{"}"})' for i in range(6)]
    models.summaries = ["S1", "S2", "S3"]
    out = asyncio.run(agent._summarise({"id": "fake"}, "S0", lines))
    prompts = _prompts(models)
    assert len(prompts) > 1  # more than one part was needed
    for line in lines:  # each line went whole, in exactly one part
        assert sum(line in p for p in prompts) == 1
    assert prompts[0].startswith("Summary so far:\nS0")
    assert prompts[1].startswith("Summary so far:\nS1")  # each part carries the summary before it
    assert all(models.count(p, None) <= 600 for p in prompts)
    assert out == f"S{len(prompts)}"


def test_only_a_line_larger_than_a_part_is_shortened_and_the_summariser_is_told():
    import asyncio

    agent, models = _summariser(window=1000)
    huge = "(tool get_deck_outline: " + "o" * 5000 + "END)"
    asyncio.run(agent._summarise({"id": "fake"}, "", ["Person: shorten slide 2", huge, "Assistant: done"]))
    text = "\n".join(_prompts(models))
    assert "Person: shorten slide 2" in text and "Assistant: done" in text
    assert "END)" not in text and "did not fit and were left out" in text
    assert all(models.count(p, None) <= 600 for p in _prompts(models))


def test_a_long_conversation_folds_into_the_summary_with_its_tool_results_whole(server, fake_model):
    pid, did, cid = setup(server)
    s2 = outline(deck_bytes(server, pid, did))[1]["slide_id"]
    chat = Chat(server, pid, cid)
    try:
        fake_model.script = [{"tools": [["get_slide", {"slide_id": s2}]]}, {"text": "Read."}]
        chat.send("read slide 2")
        tool_result = next(m["content"] for sent in fake_model.sent for m in sent if m["role"] == "tool")
        for i in range(7):  # about 22,000 tokens: more than the history's room
            fake_model.script = [{"text": "ok"}]
            chat.send(f"{i} " + "words " * 1300)
    finally:
        chat.close()
    prompts = [m[0]["content"] for m in fake_model.sent if m[0]["content"].startswith("Summary so far:")]
    assert prompts, "the conversation was never folded"
    assert len(tool_result) > 300 and any(tool_result in p for p in prompts)
    assert "Person: read slide 2" in prompts[0]
    last = fake_model.sent[-1]
    assert "Summary." in last[0]["content"]  # the summary is in the next call's context
    assert not any(m.get("content") == "read slide 2" for m in last)  # and the folded turn is not


def test_a_model_too_small_is_refused_on_exact_counts_only():
    """The fixed part (prompt and tools) over a quarter of the window: refused when counted exactly; on an estimate
    (the tokenizer down) the turn goes on with the estimate's smaller budgets."""
    import pytest

    from app.agent import context, tools

    from .fakes import FakeModels

    defs = tools.definitions(tools.load_schemas())
    model = {"id": "fake", "provider": "local"}
    with pytest.raises(context.ModelTooSmall):
        context.budget(FakeModels(window=8192, estimating=False), model, defs)
    # by the fake's count every tool is 30-50% of windows from 31 800 to 52 900 (measured 2026-10-07)
    room, answer = context.budget(FakeModels(window=40000, estimating=True), model, defs)  # over 30%, under 50%
    assert answer == int(40000 * context.ANSWER_SHARE) and room > 0
    with pytest.raises(context.ModelTooSmall):  # an estimate past half the window: refused all the same
        context.budget(FakeModels(window=8192, estimating=True), model, defs)


def test_a_large_deck_keeps_every_slide_in_the_map_and_details_the_selection():
    """200 slides in a small budget: each slide keeps its line (its ID: never guessed - measured: gemma-4 edited
    slide 70 for "slide 120" when slide 120 was missing from its context), the selected slide's neighbourhood is
    in full with shape IDs, and the whole stays within the budget."""
    from app.agent import context

    from .fakes import FakeModels

    m = FakeModels()
    count = lambda s: m.count(s, None)  # noqa: E731
    prs = read.open_deck((DECKS / "long-200.pptx").read_bytes())
    ids = [o["slide_id"] for o in read.outline(prs)]
    text, brief = context.deck_map(prs, ids[119], 4000, count)
    assert count(text) <= 4000
    assert all(f"slide {i + 1} · id {sid}" in text for i, sid in enumerate(ids))
    assert f"slide 120 · id {ids[119]} · Title and Content\n  «Diapositivo 120» (shape 2, title)" in text
    assert "[0] Conteúdo do diapositivo 1\n" not in text  # far from the selection: its line only
    assert "\n    [0] Conteúdo do diapositivo 120\n    [1] Segunda linha\n" in text  # near it: one paragraph a line
    assert 0 < brief < 200
    small, brief = context.deck_map(read.open_deck((DECKS / "simple.pptx").read_bytes()), None, 4000, count)
    right = "shape 4 (body, right), 2 paragraphs:\n    [0] Fluxo único\n    [1] Um só sistema"  # with text: its plain role
    assert brief == 0 and right in small


def test_when_every_edit_fails_the_reply_carries_the_applications_notice(server, fake_model):
    """The model may say it changed the deck after its calls failed (measured: "I've added a new slide" after two
    refused add_slide calls); the stored reply then carries what the application knows."""
    pid, did, cid = setup(server)
    s2 = outline(deck_bytes(server, pid, did))[1]["slide_id"]
    fake_model.script = [
        {"tools": [["update_text", {"slide_id": s2, "shape_id": 999, "paragraphs": [{"text": "x"}]}]]},
        {"text": "Vou tentar de novo."},
        {"text": "Vou alterar o diapositivo 2."},  # asked once more (FAILED_NUDGE), it still makes no call
        {"text": "Pronto, alterei o diapositivo 2."},  # asked a last time (STILL_NUDGE): still none
    ]
    chat = Chat(server, pid, cid)
    try:
        chat.send("muda o diapositivo 2")
    finally:
        chat.close()
    asked = [m["content"] for sent in fake_model.sent for m in sent if m["role"] == "user"]
    assert any("Nothing has changed yet: your edits failed" in c for c in asked)
    assert "Still nothing has changed" in fake_model.sent[-1][-1]["content"]  # and a last time
    reply = chat.last("assistant_message")
    assert reply["content"] == "Pronto, alterei o diapositivo 2." and reply["notice"] == "nothing_changed"
    fake_model.script = [
        {
            "tools": [
                [
                    "update_text",
                    {
                        "slide_id": s2,
                        "shape_id": body_id(deck_bytes(server, pid, did), s2),
                        "paragraphs": [{"text": "Novo", "type": "body"}],
                    },
                ]
            ]
        },
        {"text": "Feito."},
    ]
    chat = Chat(server, pid, cid)
    try:
        chat.send("agora com texto simples")
    finally:
        chat.close()
    reply = chat.last("assistant_message")
    assert "notice" not in reply  # the edit worked: a plain paragraph, a stray key dropped


def test_router_sketch_holds_every_slide_whole():
    from app.agent import router

    shapes = [
        {"shape_id": 2, "type": "placeholder", "paragraphs": [{"runs": [{"text": "Results"}]}]},
        {"shape_id": 3, "type": "placeholder", "paragraphs": [{"runs": [{"text": f"Point {i} " + "x" * 40}]} for i in range(6)]},
        {"shape_id": 4, "type": "table", "table": {"cells": [["Tier", "Price"], ["Pro", "2 500"]]}},
        {"shape_id": 5, "type": "picture", "alt_text": "A chart"},
        {"shape_id": 6, "type": "group", "shapes": [{"shape_id": 7, "paragraphs": [{"runs": [{"text": "Inner"}]}]}]},
    ]
    lines = router.slide_lines(shapes, 2)
    assert "Results" not in lines  # the title is shown apart
    assert lines == [*(f"Point {i} " + "x" * 40 for i in range(6)), "Tier | Price", "Pro | 2 500", "picture (A chart)", "Inner"]
    slides = [{"index": i, "title": f"T{i}", "lines": lines, "kinds": {"table"}} for i in range(3)]
    full = router.sketch(slides, "slide 2")
    assert full.count("    Point 5 ") == 3 and "Slide 1, title: T0 [table]\n    Point 0" in full
    assert full.endswith("Selected by the person: slide 2")
    part = router.sketch(slides, "", whole={1})
    assert part.count("Point 5") == 1 and "Slide 1, title: T0 [table]\nSlide 2" in part


def test_router_answer_parsed_with_text_tools_for_every_change():
    from app.agent import router

    groups, intent, kind = router.parse('{"intent": "Delete slide 4", "kind": "change", "groups": ["structure"]}')
    assert groups == {"structure", "text"} and intent == "Delete slide 4" and kind == "change"
    assert router.parse('{"intent": "How many?", "kind": "question", "groups": []}')[0] == set()
    assert router.parse('{"intent": "Change it", "kind": "unclear", "groups": []}')[2] == "unclear"
    assert router.parse("not json")[0] is None  # no route: every tool


def test_now_lists_the_edits_done_this_turn():
    from app.agent.loop import _done

    assert _done("duplicate_slide", '{"slide_id": 3}', {"new_slide_ids": [300]}) == "duplicate_slide on slide 3, new slide ID 300"
    assert _done("add_slide", {"after_slide_id": 5}, {"new_slide_ids": [300]}) == "add_slide, new slide ID 300"  # by its ID
    assert _done("set_notes", "not json", {}) == "set_notes"


def test_a_slide_the_router_names_by_a_title_the_person_did_not_say_is_the_whole_slide(server, fake_model):
    """Measured (onboarding-13): "translate the last slide", the router's where its title "Questions?", told the title
    names the slide, the model translated the title alone 3 runs in 3. Said by the person, the title still names it."""
    pid, did, cid = setup(server)
    title = outline(deck_bytes(server, pid, did))[1]["title"]
    chat = Chat(server, pid, cid)
    try:
        for said in ("Translate the second slide", f"Translate {title}"):
            route = {"intent": "Translate the slide", "where": f'2: "{title}"', "kind": "change", "groups": ["text"]}
            fake_model.routes = [route]
            fake_model.script = [{"text": "Done."}, {"text": "Done."}, {"text": "Done."}]
            chat.send(said)
            now = next(m[0]["content"] for m in fake_model.sent if m and m[0]["role"] == "system" and "## Now" in m[0]["content"])
            fake_model.sent.clear()
            if said.endswith("slide"):
                assert "(where: slide 2)" in now and "is the title of slide" not in now
            else:
                assert f"«{title}» is the title of slide 2" in now
    finally:
        chat.close()


def test_a_cover_asked_for_on_an_empty_deck_keeps_its_lines_on_the_cover(server, fake_model):
    """The assistant's path, with the calls measured in cover-for-a-new-deck: fill_slide on "slide 1" of an empty deck,
    the audience and the date as points; the slide is made on the cover, those its subtitle's lines."""
    pid = requests.post(f"{server}/api/projects", json={"name": "P"}, headers=h(), timeout=10).json()["id"]
    tpl = {"title": "Apresentação", "template": {"kind": "admin", "id": "default"}}
    did = requests.post(f"{server}/api/projects/{pid}/decks", json=tpl, headers=h(), timeout=60).json()["id"]
    cid = requests.post(f"{server}/api/projects/{pid}/conversations", json={"deck_id": did}, headers=h(), timeout=10).json()["id"]
    sent = {"slide_id": 1, "content": {"title": "Crédito Habitação Jovem", "subtitle": "Crédito Habitação Jovem",
                                       "points": ["Para a Equipa Comercial", "Outubro de 2026"]}}  # fmt: skip
    fake_model.routes = [{"intent": "A cover", "kind": "change", "groups": ["structure", "text"]}]
    fake_model.script = [{"tools": [("fill_slide", sent)]}, {"text": "Fiz a capa."}]
    chat = Chat(server, pid, cid)
    try:
        chat.send("Cria a capa de uma apresentação sobre o Crédito Habitação Jovem, para a equipa comercial, outubro de 2026")
        chat.decide(chat.last("proposal_updated")["id"], True)
    finally:
        chat.close()
    prs = read.open_deck(deck_bytes(server, pid, did))
    assert len(prs.slides) == 1 and prs.slides[0].slide_layout.name == "Title Slide"
    texts = [sh.text_frame.text for sh in prs.slides[0].placeholders]
    assert texts == ["Crédito Habitação Jovem", "Para a Equipa Comercial\nOutubro de 2026"]


def test_points_without_a_title_get_one_from_the_request_in_the_same_call(server, fake_model):
    """Measured (first-use-index-from-kb): an index sent as its points alone went without a heading (1 run in 4);
    refused, the model gave up or asked the person (3 runs in 3); told to write the title after, it wrote it to a slide
    that was not there (3 runs in 3). Titled from the request and the points, in the same call."""
    pid, did, cid = setup(server)
    points = {"points": ["Oferta", "Fases do processo", "Servicing"]}
    fake_model.routes = [{"intent": "An index", "kind": "change", "groups": ["structure", "text"]}]
    fake_model.script = [{"tools": [("add_slide", {"content": points, "layout": "Title and Content"})]},
                         {"text": "Fiz o índice."}]  # fmt: skip
    fake_model.titles = ["Índice: Crédito à Habitação"]
    chat = Chat(server, pid, cid)
    try:
        chat.send("escreve um índice com base nos tópicos de crédito à habitação")
        chat.decide(chat.last("proposal_updated")["id"], True)
    finally:
        chat.close()
    asked = fake_model.titled[0][0]["content"]
    assert "escreve um índice" in asked and "- Fases do processo" in asked  # the request and the points
    assert "Índice: Crédito à Habitação" in [o["title"] for o in outline(deck_bytes(server, pid, did))]


def test_a_slide_whose_title_alone_nobody_said_is_filled_by_the_next_request(server, fake_model):
    """Measured (first-use-index-from-kb, 4 runs in 8): "Cria um slide" made "Novo Slide"; the next request's index went
    on a second slide. A slide holding only a title in none of the person's words is the one to fill."""
    pid = requests.post(f"{server}/api/projects", json={"name": "P"}, headers=h(), timeout=10).json()["id"]
    tpl = {"title": "Apresentação", "template": {"kind": "admin", "id": "default"}}
    did = requests.post(f"{server}/api/projects/{pid}/decks", json=tpl, headers=h(), timeout=60).json()["id"]
    cid = requests.post(f"{server}/api/projects/{pid}/conversations", json={"deck_id": did}, headers=h(), timeout=10).json()["id"]
    index = {"title": "Índice", "points": ["Oferta", "Fases do processo", "Servicing"]}
    fake_model.routes = [{"intent": "A slide", "kind": "change", "groups": ["structure"]},
                         {"intent": "An index", "where": "1", "kind": "change", "groups": ["structure", "text"]}]  # fmt: skip
    fake_model.script = [{"tools": [("add_slide", {"content": {"title": "Novo Slide"}, "layout": "Title Only"})]},
                         {"text": "Criei um slide."},
                         {"tools": [("add_slide", {"content": index, "after_slide_id": 1})]},
                         {"text": "Escrevi o índice."}]  # fmt: skip
    chat = Chat(server, pid, cid)
    try:
        chat.send("Cria um slide")
        chat.decide(chat.last("proposal_updated")["id"], True)
        chat.send("Escreve um índice com base nos tópicos de crédito à habitação")
        chat.decide(chat.last("proposal_updated")["id"], True)
    finally:
        chat.close()
    out = outline(deck_bytes(server, pid, did))
    assert len(out) == 1 and out[0]["title"] == "Índice"  # the index on the placeholder slide, not a second one


def test_an_edit_made_twice_in_a_turn_or_a_paragraph_inserted_again_changes_nothing(server, fake_model):
    """Measured: told its answer said something was still to do, the model made the same change again (an index on two
    slides) or inserted its long point a second time (a box twice as full)."""
    pid, did, cid = setup(server)
    data = deck_bytes(server, pid, did)
    sid = outline(data)[1]["slide_id"]
    shape = body_id(data, sid)
    words = read.slide(read.open_deck(data), sid)
    first_line = "".join(r["text"] for r in next(x for x in words["shapes"] if x["shape_id"] == shape)["paragraphs"][0]["runs"])
    notes = {"slide_id": sid, "text": "O que dizer."}
    again = {"slide_id": sid, "shape_id": shape, "operations": [{"op": "insert", "after": 0, "text": first_line}]}
    fake_model.routes = [{"intent": "Notes", "kind": "change", "groups": ["notes", "text"]}]
    fake_model.script = [{"tools": [("set_notes", notes)]}, {"tools": [("set_notes", notes)]},
                         {"tools": [("edit_paragraphs", again)]}, {"text": "Feito."}]  # fmt: skip
    chat = Chat(server, pid, cid)
    try:
        chat.send("escreve as notas do segundo diapositivo")
    finally:
        chat.close()
    from .test_kb import tool_results

    assert tool_results(fake_model, "set_notes")[-1]["error"]["code"] == "ALREADY_MADE"
    assert tool_results(fake_model, "edit_paragraphs")[-1]["error"]["code"] == "SAME_TEXT"


def test_a_turn_resumed_after_an_approved_deletion_goes_on_with_the_same_request(server, fake_model):
    pid, did, cid = setup(server)
    second = outline(deck_bytes(server, pid, did))[1]["slide_id"]
    fake_model.routes = [{"intent": "Delete the agenda", "kind": "change", "groups": ["structure"]}]
    fake_model.script = [{"tools": [("delete_slide", {"slide_id": second})]}]
    chat = Chat(server, pid, cid)
    assert chat.send("delete the agenda")["status"] == "waiting"  # the application's plan for a deletion
    routed = sum(1 for m in fake_model.sent if m[0]["role"] == "user")  # the router's calls: one user message
    fake_model.script = [{"text": "I deleted the agenda."}]
    chat.answer(approve=True)
    assert sum(1 for m in fake_model.sent if m[0]["role"] == "user") == routed  # not routed again
    now = fake_model.sent[-1][0]["content"]
    assert "The person wants: Delete the agenda" in now and "Done in this turn: " in now and "(done)" in now
    done = now.split("Done in this turn: ")[1].split("\n")[0]
    assert "«Agenda»" in done and "258" not in done  # the slide as the approved plan named it, not its ID
    assert not any("called no tool, so nothing has happened" in str(m.get("content")) for m in fake_model.sent[-1])
    # the same route, not every tool: no memory tools (the knowledge tools come with every change: loop._route)
    assert fake_model.offered[-1] and "delete_slide" in fake_model.offered[-1] and "remember" not in fake_model.offered[-1]
    p = chat.last("proposal_updated")
    assert [s["state"] for s in p["decks"][did]["slides"]] == ["removed"]  # one slide, once
    chat.close()


def test_a_slide_number_keeps_its_meaning_for_the_whole_request(server, fake_model):
    pid, did, cid = setup(server)
    ids = [o["slide_id"] for o in outline(deck_bytes(server, pid, did))]
    fake_model.routes = [{"intent": "Delete slide 2", "kind": "change", "groups": ["structure"]}]
    fake_model.script = [{"tools": [("delete_slide", {"slide_id": 2})]}]
    chat = Chat(server, pid, cid)
    assert chat.send("delete slide 2")["status"] == "waiting"
    # after the deletion, the model asks for "slide 2" again: in the request's numbering it is gone, not slide 3
    fake_model.script = [{"tools": [("delete_slide", {"slide_id": 2})]}, {"text": "Slide 2 is deleted."}]
    chat.answer(approve=True)
    results = [m for m in fake_model.sent[-1] if m["role"] == "tool"]
    assert "SLIDE_DELETED" in results[-1]["content"]
    now = fake_model.sent[-1][0]["content"]
    assert "Deleted in this request: slide 2" in now and f"slide 3 · id {ids[2]}" in now  # numbers did not shift
    p = chat.last("proposal_updated")
    assert [(s["slide_id"], s["state"]) for s in p["decks"][did]["slides"]] == [(ids[1], "removed")]  # one slide
    chat.close()


def test_the_layouts_are_sent_only_when_a_slide_can_be_made_or_relaid(server, fake_model):
    # measured: a text-only turn took a layout example's "idx 17" for the paragraph index to edit
    pid, _, cid = setup(server)
    chat = Chat(server, pid, cid)
    fake_model.routes = [{"intent": "Fix a typo", "kind": "change", "groups": ["text"]}]
    fake_model.script = [{"text": "Nothing to fix."}]
    chat.send("fix the typo")
    assert "## Layouts of this deck" not in fake_model.sent[-1][0]["content"]
    fake_model.routes = [{"intent": "Add a slide", "kind": "change", "groups": ["structure"]}]
    fake_model.script = [{"text": "Which layout?"}]
    chat.send("add a slide")
    assert "## Layouts of this deck" in fake_model.sent[-1][0]["content"]
    chat.close()


def test_the_new_text_a_request_quotes_does_not_choose_the_slide():
    from app.agent import router

    req = "Nos destaques, junta à coluna do que correu bem: «Margem acima do objetivo»."
    assert router.unquoted(req) == "Nos destaques, junta à coluna do que correu bem: ."
    slides = [{"index": 2, "title": "Resultados", "lines": ["Margem operacional de 21%, acima do objetivo de 19%"]},
              {"index": 4, "title": "Destaques", "lines": ["Lançamento da app nova"]}]  # fmt: skip
    assert router.phrase_slide(req, slides) == (3, "acima do objetivo")  # the quote alone would choose slide 3
    assert router.phrase_slide(router.unquoted(req), slides) is None


def test_the_words_a_request_names_are_found_in_the_deck():
    from app.agent import router

    assert router.quoted("Corrige 'Responsabilities' e «Dev.AI 2.0» e “x”") == ["Responsabilities", "Dev.AI 2.0"]
    answer = '{"intent": "x", "where": "Slide 6: body, left: HR: hr@example.com", "kind": "change", "groups": []}'
    assert router.where_text(answer) == "HR: hr@example.com" and router.where_slide(answer) == 6
    assert router.where_slide('{"where": "16"}') == 16 and router.where_slide('{"where": ""}') is None
    paragraphs = [(6, 3, "[1]", "HR:  hr@example.com"), (6, 4, "[1]", "Your buddy"), (4, 9, "row 2, column 1", "Teams")]
    assert router.locate(paragraphs, "hr: HR@example.com") == [(6, 3, "[1]")]  # spaces and case aside
    assert router.locate(paragraphs, "teams") == [(4, 9, "row 2, column 1")]
    assert router.locate(paragraphs, "no") == []  # too short to mean anything


def test_layouts_are_described_by_what_their_placeholders_are_for():
    """Measured: Banco CTT's layouts type their heading "body" and are named "10_Texto"; in an empty deck the model
    chose four columns for "create a slide" and wrote nothing. Each layout as drawn (rendered and looked at)."""
    from pptx import Presentation

    from app.agent import context
    from app.docengine import layouts, read

    text = context.layouts_text(read.open_deck((REPO / "fixtures" / "eval" / "report.pptx").read_bytes()))
    assert "- Title and Content: heading [0] (32 pt); text large, 32 pt: [1] (as slide 2)" in text
    assert "- Two Content: heading [0] (32 pt); 2 side by side" in text  # its columns are not a subtitle
    assert "- Blank: no placeholders" in text
    ctt = Presentation(str(REPO / "templates" / "bancoctt.pptx"))
    by = {lay.name: layouts.describe(lay, ctt.slide_width, ctt.slide_height) for lay in ctt.slide_layouts}
    assert by["10_Texto"] == (
        "heading [19] (24 pt); subtitle [20]; 4 side by side, text, 18 pt bold: [30] [32] [34] [36]; "
        "4 side by side, text large, 12 pt: [12] [31] [33] [35]"
    )
    assert by["2_Capa S/Imagem"] == "heading [11] (48 pt); subtitle [12]"
    assert by["3_Agenda"].startswith("heading [0] (60 pt); 4 side by side, number, 36 pt bold: [47] [48] [49] [50]")
    lay = next(x for x in ctt.slide_layouts if x.name == "1_Tabela")
    size = (ctt.slide_width, ctt.slide_height)
    assert (layouts.heading(lay, *size), layouts.subtitle(lay, *size)) == (26, 25)


def test_a_title_goes_to_the_heading_of_a_layout_whose_heading_is_typed_body():
    from pptx import Presentation

    from app.docengine import ops

    ctt = io.BytesIO((REPO / "templates" / "bancoctt.pptx").read_bytes())
    spec = {"layout": "10_Texto", "placeholders": [{"type": "title", "text": "Índice"}, {"type": "subtitle", "text": "Crédito"}]}
    data, made = ops.apply(ctt.getvalue(), "add_slide", spec)
    s = Presentation(io.BytesIO(data)).slides.get(made["slides"][0])
    texts = {ph.placeholder_format.idx: ph.text_frame.text for ph in s.placeholders if ph.has_text_frame}
    assert texts[19] == "Índice" and texts[20] == "Crédito"


def test_every_tool_accepts_paragraphs_as_the_model_writes_them():
    """{"text": ...} is made into runs before validation (_plain_paragraphs): a contract that took text only refused
    every correct duplicate_shape call (measured on a real diagram request)."""
    from app.agent.tools import Executor

    ex = Executor.__new__(Executor)
    from app.agent.tools import load_schemas

    ex.schemas = load_schemas()
    for name, schema in ex.schemas.items():
        if "paragraphs" in schema.get("properties", {}):
            args = {k: 1 for k in schema.get("required", []) if k != "paragraphs"}
            args.update(slide_id=1, paragraphs=[{"text": "One"}, {"text": "Two", "level": 1}])
            if "shape_id" in schema.get("required", []):
                args["shape_id"] = 2
            if "kind" in schema.get("required", []):
                args.update(kind="text_box", box={"x": 0, "y": 0, "w": 0.5, "h": 0.5})
            if name == "add_slide":
                continue  # its paragraphs are inside placeholders
            ex.parse(name, args)  # raises on a refusal


def test_lines_the_router_ran_together_are_found_one_by_one():
    from app.agent import router

    paragraphs = [(30, 8, "[5]", "Sugerimos criar um PBI para agrupar os Incidentes atendidos no quarter:"),
                  (30, 8, "[6]", "Ticketing - 2026 Q1 Incidentes"), (30, 8, "[7]", "BPRB")]  # fmt: skip
    where = "Sugerimos criar um PBI para agrupar os Incidentes atendidos no quarter: Ticketing - 2026 Q1 Incidentes"
    assert router.locate(paragraphs, where) == []  # in no single paragraph
    assert [p[2] for p in router.contained(paragraphs, where)] == ["[5]", "[6]"]  # each line it holds (BPRB: too short)


def test_a_slide_the_request_quotes_word_for_word_is_found():
    from app.agent import router

    titles = ["4.3 Incidentes – BINCs (Bugs Snow)", "4.4 Problemas – BPRBs (PBIs Snow)", "4.5 Service Requests"]  # noqa: RUF001
    slides = [{"index": i, "title": "04. Gestão de Ticketing", "lines": [line]} for i, line in enumerate(titles)]
    asked = "Duplica o diapositivo 4.3 Incidentes – BINCs."  # noqa: RUF001
    assert router.phrase_slide(asked, slides) == (1, "4.3 Incidentes – BINCs")  # noqa: RUF001
    assert router.phrase_slide("Duplica o diapositivo de ticketing.", slides) is None  # every slide shares the title
    # a line of one slide whole, another slide's line close to it (devai-09: "KPIs Enabling 1/2" and "2/2")
    kpis = [{"index": 18, "title": "", "lines": ["KPIs Enabling 1/2", "Dev.AI"]},
            {"index": 19, "title": "", "lines": ["KPIs Enabling 2/2", "Dev.AI"]}]  # fmt: skip
    assert router.phrase_slide("Duplica o diapositivo KPIs Enabling 2/2.", kpis) == (20, "KPIs Enabling 2/2")
    assert router.phrase_slide("Apaga o diapositivo Dev.AI", kpis) is None  # on both, and short


def test_text_left_too_long_is_fixed_before_the_turn_ends(server, fake_model):
    pid, did, cid = setup(server)
    s1 = outline(deck_bytes(server, pid, did))[0]["slide_id"]
    shapes = read.slide(read.open_deck(deck_bytes(server, pid, did)), s1)["shapes"]
    title = next(s for s in shapes if (s.get("placeholder") or {}).get("idx") == 0)
    long = " ".join(["Uma frase longa que não cabe num título"] * 12)
    fake_model.script = [
        {"tools": [("update_text", {"slide_id": s1, "shape_id": title["shape_id"], "paragraphs": [{"text": long}]})]},
        {"text": "Mudei o título."},  # the turn would end here, the title far too long
        {"tools": [("update_text", {"slide_id": s1, "shape_id": title["shape_id"], "paragraphs": [{"text": "Curto"}]})]},
        {"text": "Encurtei o título para caber."},
    ]
    chat = Chat(server, pid, cid)
    chat.send("põe este texto no título")
    asked = [m[-1]["content"] for m in fake_model.sent if m and m[-1]["role"] == "user" and "does not fit" in m[-1]["content"]]
    assert len(asked) == 1 and f"slide 1, shape {title['shape_id']}" in asked[0]
    assert chat.last("assistant_message")["content"] == "Encurtei o título para caber."
    chat.close()


def test_with_no_route_and_every_tool_too_big_the_usual_groups_are_offered(server, fake_model):
    from app.agent import router

    pid, _, cid = setup(server)
    # by the fake's count (2.5 characters a token), the usual groups fit 30% of windows above 39 600, and the tools this
    # project is offered do not fit 42 000 (measured 2026-10-07; the real model: 32.9% and 24.6% of 32 768)
    fake_model._window, fake_model._estimating = 42000, False
    fake_model.routes = [{"intent": "-"}]  # an answer with no groups: no route
    fake_model.script = [{"text": "Feito."}]
    chat = Chat(server, pid, cid)
    chat.send("faz qualquer coisa")
    offered = set(fake_model.offered[-1])
    assert chat.last("turn_ended")["status"] == "done"  # not refused
    usual = router.tools_of(router.FALLBACK | {"knowledge"})  # the knowledge tools with them: there is a knowledge base
    assert offered == usual & offered and "update_text" in offered and "kb_search" in offered and "remember" not in offered
    chat.close()


def test_a_new_slide_has_a_place_in_an_empty_deck_and_after_slide_0_is_the_first():
    """Measured: a new deck from a template (no slides) refused "add a slide" five times: "after slide 0" and "after
    slide 1" were "not found" in a deck of slides 1 to 0."""
    from pptx import Presentation

    from app.agent.tools import Executor
    from app.docengine import ops

    empty = io.BytesIO()
    Presentation().save(empty)
    for after in (0, 1, 3):
        assert Executor._positions({"layout": "Blank", "after_slide_id": after}, empty.getvalue()) == {"layout": "Blank"}
    assert Executor._positions({"before_slide_id": 1}, empty.getvalue()) == {}
    first_args = Executor._positions({"layout": "Blank", "after_slide_id": 0}, empty.getvalue())
    data, made = ops.apply(empty.getvalue(), "add_slide", first_args)
    first = made["slides"][0]
    data, second = ops.apply(data, "add_slide", Executor._positions({"layout": "Blank", "after_slide_id": 0}, data))
    assert [o["slide_id"] for o in outline(data)] == [second["slides"][0], first]  # "after slide 0": at the start
    with pytest.raises(Exception, match="slides 1 to 2"):
        Executor._positions({"slide_id": 3}, data)  # a slide that is not there is still refused


def test_tools_that_do_not_fit_are_left_out_the_least_needed_first(server, fake_model):
    """Measured: a request needing slides, a table, a picture and the knowledge base was offered tools taking 34.1% of
    the window with the prompt, and the turn was refused; now the least needed are left out until they fit."""
    from app.agent.loop import EXPENDABLE

    pid, _, cid = setup(server)
    fake_model._window, fake_model._estimating = 42000, False
    fake_model.routes = [{"intent": "x", "kind": "change", "groups": ["structure", "table", "images", "decks", "knowledge"]}]
    fake_model.script = [{"text": "Feito."}, {"text": "Feito."}]
    chat = Chat(server, pid, cid)
    chat.send("faz uma coisa grande")
    chat.close()
    offered = set(fake_model.offered[-1])
    assert chat.last("turn_ended")["status"] == "done"  # not refused
    assert "update_text" in offered and "add_slide" in offered and "kb_search" in offered  # the core of it kept
    assert "kb_read_document" not in offered and EXPENDABLE[0] == "kb_read_document"  # the first left out


def test_the_assistant_does_not_change_a_deck_another_member_has_open(server, fake_model):
    """Spec PJ-13: a deck held by another member's open editor is not changed by anyone else's assistant."""
    from .test_kb import tool_results

    pid, did, cid = setup(server)
    rui = "rui@example.com"
    requests.put(f"{server}/api/projects/{pid}/members", json={"email": rui, "role": "editor"}, headers=h(), timeout=10)
    assert requests.post(f"{server}/api/projects/{pid}/decks/{did}/lease", headers=h(rui), timeout=10).status_code == 200
    data = deck_bytes(server, pid, did)
    sid = outline(data)[1]["slide_id"]
    fake_model.script = [
        {"tools": [("update_text", {"slide_id": sid, "shape_id": body_id(data, sid), "paragraphs": [{"text": "x"}]})]},
        {"text": "Não consegui."},
        {"text": "Não consegui."},
    ]
    chat = Chat(server, pid, cid)
    chat.send("muda a agenda")
    chat.close()
    assert tool_results(fake_model, "update_text")[0]["error"]["code"] == "DECK_LEASED"
    assert deck_bytes(server, pid, did) == data


def _figure_turn(server, fake_model, steps, say="acrescenta no fim um diapositivo com isto"):
    """One turn on the three-slide deck with the scripted calls; its proposal accepted. The deck's slides after, and
    every tool result the model was sent."""
    from .test_kb import tool_results

    pid, did, cid = setup(server)
    fake_model.routes = [{"intent": "A new slide with a figure", "kind": "change", "groups": ["structure", "table"]}]
    fake_model.script = [{"tools": [step]} for step in steps] + [{"text": "Feito."}]
    chat = Chat(server, pid, cid)
    try:
        chat.send(say)
        proposal = next((d for n, d in reversed(chat.events) if n == "proposal_updated"), None)
        if proposal:
            chat.decide(proposal["id"], True)
    finally:
        chat.close()
    prs = read.open_deck(deck_bytes(server, pid, did))
    results = {name: tool_results(fake_model, name) for name in ("add_slide", "add_slides", "draw_diagram", "add_chart")}
    return prs, results


def test_a_figure_on_the_slide_after_the_last_makes_that_slide_first(server, fake_model):
    """Measured: the model drew a diagram on "slide 4" of a three-slide deck, was told it is not there, and ended the
    turn saying it was done. The slide is added at the end, then the diagram drawn on it (on the blank layout: a
    diagram has no title); a chart's title becomes its slide's heading."""
    nodes = [{"text": t} for t in ("Pedido", "Análise", "Aprovação", "Entrega")]
    prs, results = _figure_turn(server, fake_model, [("draw_diagram", {"slide_id": 4, "nodes": nodes})])
    assert len(prs.slides) == 4
    last = list(prs.slides)[-1]
    assert last.slide_layout.name == "Blank" and not list(last.placeholders)
    boxes = [sh.text_frame.text for sh in last.shapes if sh.has_text_frame and sh.text_frame.text]
    assert boxes == ["Pedido", "Análise", "Aprovação", "Entrega"]
    assert sum(sh._element.tag.endswith("cxnSp") for sh in last.shapes) == 3
    assert "There was no slide 4" in results["draw_diagram"][-1]["note"]

    series = [{"name": "Vendas", "values": [120, 150, 180]}]
    chart = {"slide_id": 4, "kind": "column", "title": "Vendas por ano", "categories": ["2023", "2024", "2025"], "series": series}
    prs, _ = _figure_turn(server, fake_model, [("add_chart", chart)])
    last = list(prs.slides)[-1]
    assert len(prs.slides) == 4 and last.shapes.title.text == "Vendas por ano"
    graphic = next(sh for sh in last.shapes if getattr(sh, "has_chart", False))
    assert not graphic.chart.has_title  # the slide's heading says it once
    empty = [ph for ph in last.placeholders if ph.has_text_frame and not ph.text_frame.text.strip()]
    assert not empty


def test_a_figure_on_the_slide_this_turn_added_goes_on_it_and_a_refused_one_adds_nothing(server, fake_model):
    nodes = [{"text": t} for t in ("A", "B")]
    prs, _ = _figure_turn(server, fake_model, [("add_slide", {"layout": "Title Only", "content": {"title": "Fluxo"}}),
                                               ("draw_diagram", {"slide_id": 4, "nodes": nodes})])  # fmt: skip
    assert len(prs.slides) == 4  # one new slide, not two
    last = list(prs.slides)[-1]
    assert last.shapes.title.text == "Fluxo" and sum(sh._element.tag.endswith("cxnSp") for sh in last.shapes) == 1
    bad = {"slide_id": 4, "kind": "column", "categories": ["2023", "2024"], "series": [{"name": "x", "values": [1]}]}
    prs, results = _figure_turn(server, fake_model, [("add_chart", bad)])
    assert len(prs.slides) == 3 and results["add_chart"][-1]["error"]["code"] == "BAD_CHART"


def test_a_new_slide_with_a_diagram_in_one_call_names_its_boxes(server, fake_model):
    nodes = [{"text": t} for t in ("Pedido", "Análise", "Aprovação")]
    add = {"after_slide_id": 3, "layout": "Title and Content", "content": {"title": "Processo", "diagram": {"nodes": nodes}}}
    prs, results = _figure_turn(server, fake_model, [("add_slide", add)])
    note = results["add_slide"][-1]["note"]
    assert "Its diagram: boxes" in note and "arrows" in note and "arrows none" not in note
    last = list(prs.slides)[-1]
    assert last.shapes.title.text == "Processo" and sum(sh._element.tag.endswith("cxnSp") for sh in last.shapes) == 2


def test_the_layout_and_the_charts_title_put_inside_content_are_taken_where_they_belong(server, fake_model):
    """The calls as the model made them (6 runs in 6; each refused, then sent again unchanged up to five times)."""
    chart = {"categories": ["2023", "2024", "2025"], "kind": "column", "number_format": "###0",
             "series": [{"name": "Vendas", "values": [120, 150, 180]}], "title": "Vendas por Ano"}  # fmt: skip
    add = {"after_slide_id": 3, "content": {"chart": chart, "layout": "Title and Content"}, "layout": "Title and Content"}
    prs, results = _figure_turn(server, fake_model, [("add_slide", add)])
    last = list(prs.slides)[-1]
    assert len(prs.slides) == 4 and last.shapes.title.text == "Vendas por Ano"
    graphic = next(sh for sh in last.shapes if getattr(sh, "has_chart", False))
    assert [list(s.values) for s in graphic.chart.plots[0].series] == [[120.0, 150.0, 180.0]]
    assert "Its chart is shape" in results["add_slide"][-1]["note"]
    nodes = [{"text": t} for t in ("Pedido", "Análise")]
    content = {"diagram": {"nodes": nodes}, "layout": "Blank", "title": "Processo"}
    add = {"after_slide_id": 3, "content": content, "layout": "Blank"}
    prs, results = _figure_turn(server, fake_model, [("add_slide", add)])
    assert list(prs.slides)[-1].shapes.title.text == "Processo"
    assert "'Blank' has no place for its title" in results["add_slide"][-1]["note"]


def test_the_slides_title_beside_content_is_taken_into_it(server, fake_model):
    """As the model sent it (2 runs in 4 of "a slide with the conditions of the product", from the knowledge base)."""
    add = {"after_slide_id": 1, "content": {"points": ["Destinado a jovens até 35 anos.", "Financiamento até 100%."]},
           "layout": "Title and Content", "title": "Condições do Crédito Habitação Jovem", "subtitle": "Resumo"}  # fmt: skip
    said = "Acrescenta um diapositivo: Destinado a jovens até 35 anos. Financiamento até 100%."  # dictated
    prs, results = _figure_turn(server, fake_model, [("add_slide", add)], say=said)
    assert results["add_slide"][-1].get("ok"), results["add_slide"][-1]
    new = list(prs.slides)[1]
    assert new.shapes.title.text == "Condições do Crédito Habitação Jovem"
    assert "Destinado a jovens até 35 anos." in [p.text for sh in new.placeholders for p in sh.text_frame.paragraphs]


def test_a_point_that_only_repeats_the_title_is_left_out_and_said(server, fake_model):
    """As the model sent it: on a cover (the title again as its point), and on a slide written without reading the
    knowledge base (7 runs in 14). Refused, covers failed 4 runs in 4: left out, and said."""
    cover = {"content": {"points": [{"text": "Crédito Habitação Jovem"}], "subtitle": "Apresentação",
                         "title": "Crédito Habitação Jovem"}, "layout": "Title Slide"}  # fmt: skip
    prs, results = _figure_turn(server, fake_model, [("add_slide", cover)])
    assert len(prs.slides) == 4 and "only repeated the title was left out" in results["add_slide"][-1]["note"]
    texts = [sh.text_frame.text for sh in list(prs.slides)[-1].placeholders if sh.has_text_frame]
    assert texts.count("Crédito Habitação Jovem") == 1
    echo = {"content": {"points": ["Condições do Crédito Habitação Jovem"], "title": "Condições do crédito habitação jovem"},
            "layout": "Title and Content"}  # fmt: skip
    _, results = _figure_turn(server, fake_model, [("add_slide", echo)])
    assert results["add_slide"][-1]["error"]["code"] == "TITLE_ALONE"  # the title again, then nothing: write it
    assert "kb_search" in results["add_slide"][-1]["error"]["hint"]
    fine = {"content": {"title": "Condições", "points": ["Até 35 anos", "Financiamento até 100%"]}, "layout": "Title and Content"}
    prs, results = _figure_turn(server, fake_model, [("add_slide", fine)], say="Condições: até 35 anos, financiamento até 100%")
    assert "left out" not in results["add_slide"][-1]["note"]


def test_charts_and_diagrams_are_offered_only_when_asked_for(server, fake_model):
    """Measured: with add_slide's chart and diagram, and draw_diagram, offered for "a slide with the conditions of the
    product", the model wrote it without reading the knowledge base 7 runs in 14; without both, it read it 6 in 6."""
    pid, _, cid = setup(server)
    chat = Chat(server, pid, cid)
    try:
        for groups, figures in ((["structure", "text"], False), (["structure", "table"], True)):
            fake_model.routes = [{"intent": "A slide", "kind": "change", "groups": groups}]
            fake_model.script = [{"text": "Ok."}]
            chat.send("acrescenta um diapositivo")
            add = next(d for d in fake_model.definitions if d["function"]["name"] == "add_slide")["function"]
            content = add["parameters"]["properties"]["content"]["properties"]
            names = [d["function"]["name"] for d in fake_model.definitions]
            assert ("chart" in content and "diagram" in content) == figures
            assert ("draw_diagram" in names) == figures and ("chart or diagram" in add["description"]) == figures
    finally:
        chat.close()


def test_points_too_long_are_placed_and_named_with_their_lengths(server, fake_model):
    """Measured: refused, the message quoted the whole passage and the model sent the same points again until the slide
    was lost (3 runs in 6). Placed, fitted, and named with their lengths for shortening; past 1000 characters, refused."""
    long_a, long_b = "Critérios de elegibilidade: " + "x" * 210, "O regime " + "y" * 320
    points = ["Até 35 anos", long_a, {"heading": "Garantia", "text": long_b}]
    add = {"after_slide_id": 1, "layout": "Title and Content", "content": {"title": "Condições", "points": points}}
    prs, results = _figure_turn(server, fake_model, [("add_slide", add)], say="Condições: até 35 anos " + long_a)
    assert len(prs.slides) == 4 and results["add_slide"][-1].get("ok")
    note = results["add_slide"][-1]["note"]
    assert f"Points 2, 3 ({len(long_a)}, {len(long_b)} characters)" in note and "xxxx" not in note and "set_notes" in note
    huge = {**add, "content": {"title": "C", "points": ["z" * 1200]}}
    _, results = _figure_turn(server, fake_model, [("add_slide", huge)])
    assert results["add_slide"][-1]["error"]["code"] == "POINTS_TOO_LONG"

def test_an_unclear_request_once_answered_is_acted_on_with_its_tools(server, fake_model):
    """Measured: "Cria um slide", routed unclear, answered "Decide tu.": the turn offered ask_user alone after the answer,
    and the model asked the same question again 13 times, nothing made."""
    pid, did, cid = setup(server)
    fake_model.routes = [{"intent": "Create a slide", "kind": "unclear", "groups": ["structure", "text"]}]
    fake_model.script = [{"tools": [("ask_user", {"question": "Que conteúdo deve ter o slide?"})]}]
    chat = Chat(server, pid, cid)
    try:
        assert chat.send("Cria um slide")["status"] == "waiting"
        assert fake_model.offered[-1] == ["ask_user"]
        add = {"layout": "Title and Content", "content": {"title": "Próximos passos", "points": ["Rever a proposta"]}}
        fake_model.script = [{"tools": [("add_slide", add)]}, {"text": "Fiz um slide."}]
        assert chat.answer(answer="Decide tu.")["status"] == "done"
        offered = fake_model.offered[-2]
        assert "add_slide" in offered and len(offered) > 1  # its tools, after the answer
        assert len(read.open_deck(deck_bytes(server, pid, did)).slides) == 3  # in review, not yet saved
        assert chat.last("proposal_updated")["decks"][did]["slides"]
    finally:
        chat.close()


def test_a_new_slide_with_its_title_alone_on_a_text_layout_is_sent_back_once(server, fake_model):
    """Measured: "a slide with the conditions of the product" made with its title alone, its text box showing its
    prompt. Once per turn: the second call goes through; a layout with no text place is never sent back."""
    alone = {"layout": "Title and Content", "content": {"title": "Condições"}}
    _, results = _figure_turn(server, fake_model, [("add_slide", alone), ("add_slide", alone)])
    assert results["add_slide"][0]["error"]["code"] == "TITLE_ALONE" and results["add_slide"][-1].get("ok")
    _, results = _figure_turn(server, fake_model, [("add_slide", {"layout": "Title Only", "content": {"title": "Obrigado"}})])
    assert results["add_slide"][-1].get("ok")


def test_content_asked_for_the_empty_slide_the_request_is_about_goes_on_it(server, fake_model):
    """Measured: "Cria um slide" made an empty cover; "write an index of the topics" then made a second slide and left
    the first showing its prompts (2 runs in 3). Also: content without a layout gets the one it fits, and a refused
    call says nothing was changed."""
    pid = requests.post(f"{server}/api/projects", json={"name": "P"}, headers=h(), timeout=10).json()["id"]
    tpl = {"title": "Apresentação", "template": {"kind": "admin", "id": "default"}}
    did = requests.post(f"{server}/api/projects/{pid}/decks", json=tpl, headers=h(), timeout=60).json()["id"]
    cid = requests.post(f"{server}/api/projects/{pid}/conversations", json={"deck_id": did}, headers=h(), timeout=10).json()["id"]
    chat = Chat(server, pid, cid)
    try:
        fake_model.routes = [{"intent": "Create a slide", "kind": "change", "groups": ["structure", "text"]}]
        fake_model.script = [{"tools": [("add_slide", {"layout": "Title and Content"})]}, {"text": "Criei um slide."}]
        chat.send("Cria um slide")
        chat.decide(chat.last("proposal_updated")["id"], True)
        index = {"content": {"title": "Índice", "points": ["Oferta", "Fases do processo", "Servicing"]}}  # no layout
        fake_model.routes = [{"intent": "An index", "where": "1", "kind": "change", "groups": ["structure", "text"]}]
        # the same call sent again is refused (measured: it went on three new slides, the index twice in the deck)
        again = ("add_slide", {**index, "after_slide_id": 1})
        fake_model.script = [{"tools": [again]}, {"tools": [again]}, {"text": "Escrevi o índice."}]
        chat.send("Escreve um índice")
        chat.decide(chat.last("proposal_updated")["id"], True)
        from .test_kb import tool_results

        assert tool_results(fake_model, "add_slide")[-1]["error"]["code"] == "ALREADY_MADE"
        # said as what happened (measured: listed as "add_slide", the reply said "a new slide, slide 2", 8 runs in 8)
        now = [m[0]["content"] for m in fake_model.sent if m[0]["role"] == "system" and "Done in this turn" in m[0]["content"]]
        assert "Done in this turn: the content written on slide 1 (no slide added)" in now[-1]
        note = tool_results(fake_model, "add_slide")[-2]["note"]
        assert "slide 1, had no text yet: the content was written on it and no slide was added" in note
        prs = read.open_deck(deck_bytes(server, pid, did))
        assert len(prs.slides) == 1 and prs.slides[0].shapes.title.text == "Índice"
        fake_model.routes = [{"intent": "A slide", "kind": "change", "groups": ["structure", "text"]}]
        fake_model.script = [{"tools": [("add_slide", {"content": {"title": "X", "points": ["a" * 1200]}})]}, {"text": "?"}]
        chat.send("acrescenta um slide")
        assert tool_results(fake_model, "add_slide")[-1]["error"]["message"].startswith("Nothing was changed")
        # a place past the deck's end is answered with the slides that exist and the way to the end (measured: "after
        # slide 18" in a one-slide deck, told only "the deck has slides 1 to 1", was sent again unchanged, 2 runs in 8)
        sid = prs.slides[0].slide_id
        fake_model.script = [{"tools": [("add_slide", {**index, "after_slide_id": 18})]}, {"text": "?"}]
        chat.send("acrescenta um slide")
        error = tool_results(fake_model, "add_slide")[-1]["error"]
        assert error["code"] == "SLIDE_NOT_FOUND" and "There is no slide 18" in error["message"]
        assert f"Slide 1, the last, is ID {sid}." in error["hint"] and "leave after_slide_id out" in error["hint"]
    finally:
        chat.close()


def test_a_slide_written_from_an_earlier_turns_passages_cites_them(server, fake_model, fake_kb):
    """Measured: the cover's turn searched, the next slide was written from those passages without searching again, and
    its notes had no source (3 runs in 3). A slide whose text a passage does not hold is not given one."""
    pid, did, cid = setup(server)
    requests.patch(f"{server}/api/projects/{pid}", json={"settings": {"kb_domains": ["produtos"]}}, headers=h(), timeout=10)
    chat = Chat(server, pid, cid)
    try:
        fake_model.script = [{"tools": [("kb_search", {"query": "garantia"})]}, {"text": "A garantia cobre 2 anos."}]
        chat.send("qual é a garantia?")
        point = "A garantia de 2026 cobre todos os equipamentos durante 2 anos a partir da data de compra."
        fake_model.routes = [{"intent": "A slide on the warranty", "kind": "change", "groups": ["structure", "text"]}]
        add = {"layout": "Title and Content", "content": {"title": "Garantia", "points": [point]}}
        fake_model.script = [{"tools": [("add_slide", add)]}, {"text": "Fiz o diapositivo."}]
        chat.send("acrescenta um diapositivo sobre a garantia")
        made = chat.last("proposal_updated")
        sid = made["decks"][did]["slides"][-1]["slide_id"]
        chat.decide(made["id"], True)
        thanks = {"title": "Obrigado", "points": ["Obrigado pela vossa atenção e até breve."]}
        mine = {"layout": "Title and Content", "content": thanks}
        fake_model.script = [{"tools": [("add_slide", mine)]}, {"text": "Fiz o diapositivo."}]
        chat.send("acrescenta um diapositivo de agradecimento")
        other = chat.last("proposal_updated")["decks"][did]["slides"][-1]["slide_id"]
        chat.decide(chat.last("proposal_updated")["id"], True)
    finally:
        chat.close()
    prs = read.open_deck(deck_bytes(server, pid, did))
    assert "Fonte: Política de garantia 2026" in read.slide(prs, sid)["notes"]
    assert not read.slide(prs, other)["notes"]


def test_a_long_point_that_is_a_list_becomes_a_point_per_item(server, fake_model):
    """As the model wrote it from the knowledge base: seven criteria in one point, which did not fit even at 14 pt."""
    crit = ("Critérios de acesso (cumulativos): Idade entre 18 e 35 anos; domicílio fiscal em Portugal; rendimentos não "
            "ultrapassarem o 8.º escalão do IRS; não ser proprietário de prédio urbano; nunca ter usufruído da garantia "
            "pessoal do Estado; valor da transação não exceder 450.000,00 euros")
    add = {"layout": "Title and Content", "content": {"title": "Condições", "points": ["Regime especial para jovens", crit]}}
    prs, results = _figure_turn(server, fake_model, [("add_slide", add)], say="Condições: " + crit)
    assert "Point 2 held a list" in results["add_slide"][-1]["note"]
    lines = [p.text for sh in list(prs.slides)[-1].placeholders for p in sh.text_frame.paragraphs if p.text]
    assert "domicílio fiscal em Portugal" in lines and "Critérios de acesso (cumulativos): Idade entre 18 e 35 anos" in lines


def test_an_edit_that_writes_back_the_same_text_says_nothing_changed(server, fake_model):
    """Measured: "shorten the agenda's points" sent every point back unchanged, was told "ok", and the reply said they
    were shortened."""
    pid, did, cid = setup(server)
    data = deck_bytes(server, pid, did)
    agenda = outline(data)[1]["slide_id"]
    shape = body_id(data, agenda)
    body = next(x for x in read.slide(read.open_deck(data), agenda)["shapes"] if x["shape_id"] == shape)
    now = [p["runs"][0]["text"] for p in body["paragraphs"]]
    same = {"slide_id": 2, "shape_id": shape, "operations": [{"op": "set", "index": i, "text": x} for i, x in enumerate(now)]}
    shorter = {**same, "operations": [{"op": "set", "index": 0, "text": "Contexto"}]}
    fake_model.script = [{"tools": [("edit_paragraphs", same)]}, {"tools": [("edit_paragraphs", shorter)]}, {"text": "Feito."}]
    chat = Chat(server, pid, cid)
    try:
        chat.send("encurta os pontos da agenda")
    finally:
        chat.close()
    from .test_kb import tool_results

    results = tool_results(fake_model, "edit_paragraphs")
    assert results[0]["error"]["code"] == "SAME_TEXT" and results[-1].get("ok")


def test_a_step_left_out_is_asked_for_before_the_turn_ends(server, fake_model):
    """Measured (squad-16, 2 runs in 6): the box copied, "I now need to connect it with an arrow", and the turn ended.
    Changes made and the turn ending in words: the check names what is missing, and the model is asked to do it."""
    pid, did, cid = setup(server)
    data = deck_bytes(server, pid, did)
    sid = outline(data)[1]["slide_id"]
    shape = body_id(data, sid)
    fake_model.routes = [{"intent": "Add a box, linked by an arrow", "kind": "change", "groups": ["structure"]}]
    first = {"slide_id": sid, "shape_id": shape, "paragraphs": [{"runs": [{"text": "Caixa"}]}]}
    fake_model.script = [{"tools": [("update_text", first)]},
                         {"text": "Copiei a caixa. Preciso agora de a ligar com uma seta."},
                         {"tools": [("set_notes", {"slide_id": sid, "text": "ligada"})]},
                         {"text": "Copiei a caixa e liguei-a."}]  # fmt: skip
    fake_model.checks = ["connect the new box to the first with an arrow"]  # what the message says is still to do
    chat = Chat(server, pid, cid)
    try:
        assert chat.send("acrescenta uma caixa igual, ligada por uma seta")["status"] == "done"
    finally:
        chat.close()
    said = "Your answer says this is still to be done"
    asked = [m for call in fake_model.sent for m in call if m["role"] == "user" and said in str(m["content"])]
    assert asked and "connect the new box" in asked[0]["content"]
    check = fake_model.checked[0][0]["content"]
    assert "Preciso agora de a ligar com uma seta" in check  # the closing message, checked
    assert len(fake_model.checked) == 1  # once a turn
    from .test_kb import tool_results

    assert tool_results(fake_model, "set_notes")  # what it was asked for was then called
