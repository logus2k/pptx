"""Presentation files: checking an upload, opening a template, a new deck from a template, the slides of a deck
(technical design sections 9 and 11). Nothing here keeps state: bytes in, bytes or facts out."""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass

from lxml import etree
from pptx import Presentation

CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
PRESENTATION = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
TEMPLATE = "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
MACRO_TYPES = (
    "application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml",
    "application/vnd.ms-powerpoint.template.macroEnabled.main+xml",
    "application/vnd.ms-powerpoint.slideshow.macroEnabled.main+xml",
    "application/vnd.ms-office.vbaProject",
)
OLE_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # a compound file: an encrypted .pptx, or an old .ppt
MAX_ENTRIES = 20000
MAX_UNCOMPRESSED = 2 * 1024**3  # 2 GiB unpacked, whatever the packed size
MAX_RATIO = 200  # unpacked / packed: a zip bomb is far above it


class Rejected(Exception):
    """An upload that is not accepted; `code` for the client, the message for people."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class Checked:
    data: bytes  # the presentation (a template's content type already made a presentation's)
    slide_count: int
    is_template: bool


def _main_content_type(z: zipfile.ZipFile) -> str:
    try:
        root = etree.fromstring(z.read("[Content_Types].xml"), etree.XMLParser(resolve_entities=False, no_network=True))
    except KeyError:
        raise Rejected("not_pptx", "This file is not a PowerPoint presentation.") from None
    for override in root.findall(f"{{{CT_NS}}}Override"):
        if override.get("PartName") == "/ppt/presentation.xml":
            return override.get("ContentType") or ""
    raise Rejected("not_pptx", "This file is not a PowerPoint presentation.")


def _as_presentation(data: bytes) -> bytes:
    """A .potx with its main part's content type changed to a presentation's, the only change: python-pptx opens
    presentations only, and a template is otherwise the same package."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            blob = src.read(info.filename)
            if info.filename == "[Content_Types].xml":
                root = etree.fromstring(blob, etree.XMLParser(resolve_entities=False, no_network=True))
                for override in root.findall(f"{{{CT_NS}}}Override"):
                    if override.get("PartName") == "/ppt/presentation.xml":
                        override.set("ContentType", PRESENTATION)
                blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            dst.writestr(info, blob)
    return out.getvalue()


def check(data: bytes, max_bytes: int) -> Checked:
    """An uploaded .pptx or .potx, checked by content (never by its name): size, a zip with a presentation part,
    unpacked size and entry count (zip bombs), no macros, not encrypted, and it opens."""
    if len(data) > max_bytes:
        raise Rejected("too_large", f"The file is larger than {max_bytes // (1024 * 1024)} MB.")
    if data[:8] == OLE_SIGNATURE:
        raise Rejected("encrypted", "This file is password-protected or in the old .ppt format: save it as an unprotected .pptx.")
    if data[:4] != b"PK\x03\x04":
        raise Rejected("not_pptx", "This file is not a PowerPoint presentation.")
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
        infos = z.infolist()
    except zipfile.BadZipFile:
        raise Rejected("not_pptx", "This file is damaged or is not a PowerPoint presentation.") from None
    unpacked = sum(i.file_size for i in infos)
    if len(infos) > MAX_ENTRIES or unpacked > MAX_UNCOMPRESSED or unpacked > MAX_RATIO * max(len(data), 1):
        raise Rejected("too_large", "This file unpacks to too much data to be a presentation.")
    main = _main_content_type(z)
    names = {i.filename.lower() for i in infos}
    if main in MACRO_TYPES or any(n.endswith("vbaproject.bin") for n in names):
        raise Rejected("macros", "Presentations with macros (.pptm, .potm) are not accepted.")
    if main not in (PRESENTATION, TEMPLATE):
        raise Rejected("not_pptx", "This file is not a PowerPoint presentation.")
    is_template = main == TEMPLATE
    if is_template:
        data = _as_presentation(data)
    try:
        prs = Presentation(io.BytesIO(data))
    except Exception:  # noqa: BLE001 - python-pptx raises many kinds for a damaged package
        raise Rejected("damaged", "This presentation could not be opened: it may be damaged.") from None
    return Checked(data=data, slide_count=len(prs.slides), is_template=is_template)


def without_slides(template: bytes) -> bytes:
    """A new deck from a template: its masters, layouts and theme, none of its slides (spec PM-9). The slides' parts
    are dropped with their relationships, so they are not written."""
    prs = Presentation(io.BytesIO(template))
    ids = prs.slides._sldIdLst
    for sld in list(ids):
        prs.part.drop_rel(sld.rId)
        ids.remove(sld)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def slides(data: bytes) -> list[dict]:
    """The deck's slides in order: id (the file's slide ID, stable across reordering), index, hidden, title."""
    prs = Presentation(io.BytesIO(data))
    out = []
    for i, s in enumerate(prs.slides):
        title = ""
        try:
            if s.shapes.title is not None and s.shapes.title.has_text_frame:
                title = s.shapes.title.text_frame.text.strip()
        except (KeyError, AttributeError):
            title = ""
        out.append({"id": s.slide_id, "index": i, "hidden": s._element.get("show") == "0", "title": title[:200]})
    return out


def layouts(data: bytes) -> list[str]:
    prs = Presentation(io.BytesIO(data))
    return [layout.name for master in prs.slide_masters for layout in master.slide_layouts]


def fonts(data: bytes) -> list[str]:
    """The typefaces a template names: its themes' major and minor Latin fonts and any explicit Latin typeface in its
    masters and layouts (theme placeholders such as +mj-lt are references, not fonts)."""
    z = zipfile.ZipFile(io.BytesIO(data))
    found: set[str] = set()
    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    for name in z.namelist():
        if not name.startswith(("ppt/theme/", "ppt/slideMasters/", "ppt/slideLayouts/")) or not name.endswith(".xml"):
            continue
        root = etree.fromstring(z.read(name), parser)
        for el in root.iter(f"{{{A_NS}}}latin"):
            face = (el.get("typeface") or "").strip()
            if face and not face.startswith("+"):
                found.add(face)
    return sorted(found)
