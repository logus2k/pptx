"""The read model (technical design section 9; contracts/slide.schema.json): a deck's outline and a slide as the
assistant sees it. Shapes are addressed by the file's IDs; what the engine cannot safely change is `locked`."""

from __future__ import annotations

import hashlib
import io

from pptx import Presentation

from . import textfit
from pptx.enum.shapes import MSO_SHAPE_TYPE

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
DIAGRAM_URI = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
CHART_URI = "http://schemas.openxmlformats.org/drawingml/2006/chart"
TABLE_URI = "http://schemas.openxmlformats.org/drawingml/2006/table"


def open_deck(data: bytes) -> Presentation:
    return Presentation(io.BytesIO(data))


# what the XML says, read from it directly: python-pptx's accessors (Font.color, paragraph.level, run.font) add
# a:pPr, a:rPr and an empty a:solidFill to what they read, and edits then saved them (measured 2026-10-06)
SCHEME = {
    "tx1": "text1", "tx2": "text2", "bg1": "background1", "bg2": "background2",
    "dk1": "text1", "lt1": "background1", "dk2": "text2", "lt2": "background2",
    "accent1": "accent1", "accent2": "accent2", "accent3": "accent3",
    "accent4": "accent4", "accent5": "accent5", "accent6": "accent6",
    "hlink": "hyperlink", "folHlink": "followed_hyperlink",
}  # fmt: skip
ALGN = {"l": "left", "ctr": "center", "r": "right", "just": "justify"}


def _flag(value: str | None) -> bool | None:
    return None if value is None else value in ("1", "true")


def _run(r) -> dict:
    """An a:r: its text and the formatting it sets itself."""
    out = {"text": "".join(t.text or "" for t in r.findall(f"{A}t"))}
    rpr = r.find(f"{A}rPr")
    if rpr is None:
        return out
    if _flag(rpr.get("b")) is not None:
        out["bold"] = _flag(rpr.get("b"))
    if _flag(rpr.get("i")) is not None:
        out["italic"] = _flag(rpr.get("i"))
    if rpr.get("sz"):
        out["size_pt"] = round(int(rpr.get("sz")) / 100, 1)
    scheme = rpr.find(f"{A}solidFill/{A}schemeClr")
    if scheme is not None and scheme.get("val") in SCHEME:
        out["theme_color"] = SCHEME[scheme.get("val")]
    return out


def paragraphs(text_frame) -> list[dict]:
    """Paragraphs with their runs in document order; a line break (a:br) is a run of "\n", a field (a:fld: a slide
    number, a date) a run of its current text."""
    out = []
    for p in text_frame._txBody.findall(f"{A}p"):
        runs = []
        for child in p:
            if child.tag == f"{A}r":
                runs.append(_run(child))
            elif child.tag == f"{A}br":
                runs.append({"text": "\n"})
            elif child.tag == f"{A}fld":
                t = child.find(f"{A}t")
                runs.append({"text": t.text if t is not None and t.text else ""})
        ppr = p.find(f"{A}pPr")
        para = {"level": int(ppr.get("lvl", 0)) if ppr is not None else 0, "runs": runs}
        if ppr is not None and ppr.get("algn") in ALGN:
            para["alignment"] = ALGN[ppr.get("algn")]
        out.append(para)
    return out


def _alt(shape) -> str:
    nv = shape._element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr")
    return (nv.get("descr") or "") if nv is not None else ""


def _geometry(shape, width: int, height: int) -> dict:
    if shape.left is None or shape.width is None:
        return {}
    x, y, w, h = int(shape.left), int(shape.top), int(shape.width), int(shape.height)
    return {
        "emu": {"x": x, "y": y, "w": w, "h": h},
        "box": {"x": round(x / width, 4), "y": round(y / height, 4), "w": round(w / width, 4), "h": round(h / height, 4)},
    }


def _graphic_uri(shape) -> str:
    gd = shape._element.find(f".//{A}graphicData")
    return gd.get("uri") if gd is not None else ""


def shape(sh, width: int, height: int, measure: bool = True) -> dict:
    out = {"shape_id": sh.shape_id, "name": sh.name, "locked": False}
    st = sh.shape_type
    if sh.is_placeholder:
        out["type"] = "placeholder"
        pf = sh.placeholder_format
        out["placeholder"] = {"type": pf.type.name.lower() if pf.type is not None else "body", "idx": pf.idx}
    elif st == MSO_SHAPE_TYPE.GROUP:
        out["type"] = "group"
    elif st == MSO_SHAPE_TYPE.PICTURE:
        out["type"] = "picture"
    elif st == MSO_SHAPE_TYPE.TABLE or (st == MSO_SHAPE_TYPE.PLACEHOLDER and sh.has_table):
        out["type"] = "table"
    elif st == MSO_SHAPE_TYPE.TEXT_BOX:
        out["type"] = "text_box"
    elif st == MSO_SHAPE_TYPE.LINE or sh._element.tag.endswith("}cxnSp"):
        out["type"] = "connector"
        ends = [sh._element.find(f".//{A}{tag}") for tag in ("stCxn", "endCxn")]
        if all(e is not None for e in ends):  # attached: the shapes it links (a diagram's arrow)
            out["connects"] = [int(e.get("id")) for e in ends]
    elif st == MSO_SHAPE_TYPE.AUTO_SHAPE or st == MSO_SHAPE_TYPE.FREEFORM:
        out["type"] = "shape"
    else:
        out["type"] = "other"
    out.update(_geometry(sh, width, height))

    # what the engine cannot safely change: kept as it is, read-only
    uri = _graphic_uri(sh) if sh._element.tag.endswith("}graphicFrame") else ""
    if uri == CHART_URI:
        out.update(type="chart", subtype="chart", locked=True)
    elif uri == DIAGRAM_URI:
        out.update(type="other", subtype="smartart", locked=True)
    elif uri and uri != TABLE_URI:
        out.update(type="other", subtype="ole" if "ole" in uri.lower() else "graphic", locked=True)
    if (
        st == MSO_SHAPE_TYPE.MEDIA
        or sh._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}videoFile") is not None
    ):
        out.update(type="other", subtype="media", locked=True)
    if sh._element.tag.endswith("}contentPart"):
        out.update(type="other", subtype="content_part", locked=True)

    if out["type"] == "table" and sh.has_table:
        t = sh.table
        out["table"] = {"rows": len(t.rows), "cols": len(t.columns), "cells": [[c.text for c in row.cells] for row in t.rows]}
    elif out["type"] == "group":
        out["shapes"] = [shape(s, width, height, measure) for s in sh.shapes]
    elif getattr(sh, "has_text_frame", False) and sh.has_text_frame and not out["locked"]:
        out["paragraphs"] = paragraphs(sh.text_frame)
        if measure:  # (an outline's one-line summary does not need it: measured, reading a deck took 3 times as long)
            out["overflow"] = overflow(sh, out["paragraphs"])
    if out["type"] == "picture":
        out["alt_text"] = _alt(sh)
        out["description"] = None  # what it shows, by the vision model (describe.py), filled by whoever has it
        try:
            out["image_sha256"] = hashlib.sha256(sh.image.blob).hexdigest()
        except (AttributeError, KeyError, ValueError):  # a linked or missing image
            pass
    elif out["type"] in ("other", "chart", "group", "shape"):
        if out["type"] == "chart":
            data = chart_data(sh)
            if data:
                out["chart"] = data
        alt = _alt(sh)
        if alt:
            out["alt_text"] = alt
    return out


def chart_data(sh) -> dict | None:
    """A chart's kind, title, categories and series (edit_chart's arguments); None when python-pptx cannot read it."""
    try:
        chart = sh.chart
        plot = chart.plots[0]
        kind = {"BAR_CLUSTERED": "bar", "COLUMN_CLUSTERED": "column", "PIE": "pie"}.get(
            chart.chart_type.name, "line" if "LINE" in chart.chart_type.name else chart.chart_type.name.lower()
        )
        return {
            "kind": kind,
            "title": chart.chart_title.text_frame.text if chart.has_title and chart.chart_title.has_text_frame else None,
            "categories": [str(c) for c in plot.categories],
            "series": [{"name": se.name, "values": list(se.values)} for se in plot.series],
        }
    except Exception:  # noqa: BLE001 - a chart python-pptx cannot read (a combination chart, a missing workbook)
        return None


def _notes(slide) -> str:
    if not slide.has_notes_slide:
        return ""
    tf = slide.notes_slide.notes_text_frame
    return tf.text if tf is not None else ""


def slide(prs: Presentation, slide_id: int) -> dict:
    s = prs.slides.get(slide_id)
    if s is None:
        raise KeyError(slide_id)
    width, height = int(prs.slide_width), int(prs.slide_height)
    index = list(prs.slides).index(s)
    return {
        "slide_id": s.slide_id,
        "index": index,
        "layout": s.slide_layout.name,
        "hidden": s._element.get("show") == "0",
        "size": {"width_emu": width, "height_emu": height},
        "shapes": [shape(sh, width, height) for sh in s.shapes],
        "notes": _notes(s),
    }


def _title(s) -> str:
    try:
        t = s.shapes.title
        return t.text_frame.text.strip() if t is not None and t.has_text_frame else ""
    except (KeyError, AttributeError):
        return ""


SUMMARY_CHARS = 200


def _describe(d: dict) -> str:
    """What one shape (as shape() reads it) holds, in a few words: its text, or what kind of thing it is."""
    if d.get("paragraphs"):
        lines = (" ".join("".join(r.get("text", "") for r in p["runs"]).split()) for p in d["paragraphs"])
        return " / ".join(x for x in lines if x)  # one paragraph from the next: where each bullet ends
    if d["type"] == "table" and d.get("table"):
        t = d["table"]
        head = ", ".join(c for c in (t["cells"][0] if t["cells"] else []) if c.strip())
        return f"table, {t['rows']} rows by {t['cols']} columns" + (f": {head}" if head else "")
    if d["type"] == "group":
        return " · ".join(x for x in (_describe(c) for c in d.get("shapes", [])) if x)
    if d["type"] == "picture":
        return f"picture ({d['alt_text']})" if d.get("alt_text") else "picture"
    if d["type"] == "chart":
        return "chart"
    if d.get("subtype"):
        return d["subtype"]
    return ""


def _summary(s, width: int, height: int) -> str:
    """One line: what the slide holds after its title - its text, tables (size and header), pictures (alt text),
    charts, groups - shortened, with a visible "…", past SUMMARY_CHARS (get_slide has all of it)."""
    title = s.shapes.title
    title_id = title.shape_id if title is not None else None  # (python-pptx makes a new proxy on every access)
    parts = [_describe(shape(sh, width, height, measure=False)) for sh in s.shapes if sh.shape_id != title_id]
    text = " · ".join(p for p in parts if p)
    return text if len(text) <= SUMMARY_CHARS else text[: SUMMARY_CHARS - 1].rstrip() + "…"


def outline(prs: Presentation) -> list[dict]:
    """get_deck_outline: slide IDs, order, layout names, titles and a one-line summary per slide."""
    return [
        {
            "slide_id": s.slide_id,
            "index": i,
            "layout": s.slide_layout.name,
            "hidden": s._element.get("show") == "0",
            "title": _title(s),
            "summary": _summary(s, prs.slide_width, prs.slide_height),
        }
        for i, s in enumerate(prs.slides)
    ]


def layouts(prs: Presentation) -> list[dict]:
    """list_layouts: each layout's name and its placeholders (type and idx)."""
    out = []
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            out.append(
                {
                    "name": layout.name,
                    "placeholders": [
                        {
                            "type": ph.placeholder_format.type.name.lower() if ph.placeholder_format.type is not None else "body",
                            "idx": ph.placeholder_format.idx,
                        }
                        for ph in layout.placeholders
                    ],
                }
            )
    return out


def pictures(data: bytes) -> dict[str, bytes]:
    """Every picture's image in a deck, once each: {sha256: bytes}."""
    out: dict[str, bytes] = {}
    for slide in open_deck(data).slides:
        stack = list(slide.shapes)
        while stack:
            sh = stack.pop()
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                stack += list(sh.shapes)
            elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    blob = sh.image.blob
                except (AttributeError, KeyError, ValueError):
                    continue
                out.setdefault(hashlib.sha256(blob).hexdigest(), blob)
    return out


def overflow(sh, paras: list[dict]) -> bool | None:
    """Do the shape's glyphs go past its box, as the deck's own fonts lay them out (textfit.py)? None when the box is
    unknown. `paras` (the read model's) is kept for the callers; textfit reads the shape itself."""
    try:
        if not sh.width or not sh.height:
            return None
    except AttributeError:
        return None
    if not any(r.get("text", "").strip() for p in paras for r in p["runs"]):
        return False
    return textfit.overflows(sh)
