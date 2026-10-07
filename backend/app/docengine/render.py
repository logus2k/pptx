"""Slide images (technical design section 9): per slide, so the cost of an edit does not grow with the deck.

Each slide's render key is a hash of what decides its look: its own XML, its layout's, its master's, its theme's,
the bytes of the media they reference, and the fonts installed (a slide rendered with a stand-in font is rendered
again once the template's own font is there). Only slides whose key has no images yet are rendered: a temporary copy
of the deck holding just those slides (made visible), LibreOffice converts it to PDF, pypdfium2 rasterises each page
to a preview and a thumbnail, stored under the slide's key. Each conversion uses a fresh LibreOffice profile and a
timeout (60 s plus 1 s per slide); at most two run at once."""

from __future__ import annotations

import asyncio
import hashlib
import io
import logging
import shutil
import subprocess
import tempfile
import uuid
from pathlib import Path

import pypdfium2 as pdfium
from pptx import Presentation

from . import emf

log = logging.getLogger("slides.render")

PREVIEW_WIDTH = 2400  # px: the stage, up to 1100 css px wide: 2200 device px on a 2x screen (at 1600, text was soft)
THUMB_WIDTH = 320  # px: the slide strip and the project home
_gate = asyncio.Semaphore(2)  # LibreOffice conversions at once
# (cache, render key) -> the future of the render making it: a second request for the same slides (the draft of a proposal and
# the version it started from share every unchanged slide) waits for that render instead of making them again
_inflight: dict[tuple[str, str], asyncio.Future] = {}


class RenderError(Exception):
    pass


RENDER_VERSION = "3"  # what the renderer does, as part of the key (2: EMF pictures, emf.py; 3: previews 2400 px wide)
_fonts: str | None = None


def fonts_digest() -> str:
    """The fonts LibreOffice can use, as one hash (fontconfig's list: each file and its size, and its configuration);
    computed once, as the fonts change only when the container is restarted with others mounted (fonts/:
    docs/licenses.md)."""
    global _fonts
    if _fonts is None:
        try:
            run = subprocess.run(["fc-list", "--format", "%{file}\n"], capture_output=True, timeout=30, check=False)  # noqa: S607
            files = sorted(set(run.stdout.decode("utf-8", "replace").split("\n")) - {""})
        except (OSError, subprocess.TimeoutExpired):
            files = []
        h = hashlib.sha256()
        for f in files:
            try:
                h.update(f"{f}:{Path(f).stat().st_size}\n".encode())
            except OSError:
                continue
        # and which font a name gets (config/fonts.conf among them): a mapping changed, the same files draw otherwise
        for conf in sorted(Path("/etc/fonts/conf.d").glob("*.conf")):
            try:
                h.update(conf.name.encode() + conf.read_bytes())
            except OSError:
                continue
        _fonts = h.hexdigest()
    return _fonts


def _part_digest(part, h, seen: set) -> None:
    """The part's bytes and, recursively, those of the parts it relates to that shape its look (layout, master,
    theme, media) - never other slides or notes (a notes change must not re-render the slide)."""
    if part.partname in seen:
        return
    seen.add(part.partname)
    h.update(str(part.partname).encode())
    h.update(part.blob)
    for rel in part.rels.values():
        if rel.is_external:
            h.update(rel.target_ref.encode())
            continue
        kind = rel.reltype.rsplit("/", 1)[-1]
        if kind in (
            "slideLayout",
            "slideMaster",
            "theme",
            "image",
            "media",
            "video",
            "audio",
            "chart",
            "package",
            "oleObject",
            "diagramData",
            "diagramLayout",
            "diagramColors",
            "diagramQuickStyle",
            "diagramDrawing",
        ):
            _part_digest(rel.target_part, h, seen)


def keys(data: bytes) -> list[dict]:
    """[{id, index, key}] for every slide of the deck, in order."""
    prs = Presentation(io.BytesIO(data))
    out = []
    for i, s in enumerate(prs.slides):
        h = hashlib.sha256()
        h.update(f"{prs.slide_width}x{prs.slide_height}".encode())
        h.update(fonts_digest().encode())
        h.update(RENDER_VERSION.encode())
        _part_digest(s.part, h, set())
        out.append({"id": s.slide_id, "index": i, "key": h.hexdigest()})
    return out


def _only(data: bytes, keep: set[int]) -> bytes:
    """The deck with only the slides whose ids are in `keep`, all visible."""
    prs = Presentation(io.BytesIO(data))
    ids = prs.slides._sldIdLst
    for sld in list(ids):
        if int(sld.get("id")) not in keep:
            prs.part.drop_rel(sld.rId)
            ids.remove(sld)
    for s in prs.slides:
        if s._element.get("show") == "0":
            s._element.attrib.pop("show")
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def _convert(pptx: bytes, count: int) -> bytes:
    """PDF bytes from LibreOffice, in a fresh profile; RenderError with LibreOffice's own words if it fails. Its EMF
    pictures are drawn as PowerPoint draws them (emf.py: vector, not clipped at their frame)."""
    try:
        pptx = emf.normalised(pptx)
    except Exception as e:  # noqa: BLE001 - a picture it cannot prepare is rendered as it is
        log.warning("EMF pictures not prepared", extra={"reason": type(e).__name__})
    with tempfile.TemporaryDirectory(prefix="slides-render-") as tmp:
        src = Path(tmp) / "deck.pptx"
        src.write_bytes(pptx)
        profile = Path(tmp) / "profile"
        cmd = [
            "soffice",
            "--headless",
            "--norestore",
            f"-env:UserInstallation=file://{profile}",
            "--convert-to",
            "pdf",
            "--outdir",
            tmp,
            str(src),
        ]
        try:
            run = subprocess.run(cmd, capture_output=True, timeout=60 + count, check=False)  # noqa: S603 - fixed argv
        except subprocess.TimeoutExpired:
            raise RenderError(f"LibreOffice took longer than {60 + count} s") from None
        pdf = Path(tmp) / "deck.pdf"
        if not pdf.exists():
            words = (run.stderr or run.stdout or b"").decode("utf-8", "replace").strip()[-500:]
            raise RenderError(f"LibreOffice made no PDF (exit {run.returncode}): {words}")
        return pdf.read_bytes()


def _rasterise(pdf: bytes, out_dirs: list[Path]) -> None:
    doc = pdfium.PdfDocument(pdf)
    if len(doc) != len(out_dirs):
        raise RenderError(f"the PDF has {len(doc)} pages for {len(out_dirs)} slides")
    for page, target in zip(doc, out_dirs, strict=True):
        width = page.get_width()
        image = page.render(scale=PREVIEW_WIDTH / width).to_pil().convert("RGB")
        target.mkdir(parents=True, exist_ok=True)
        tag = uuid.uuid4().hex  # this render's own temporary names: another render of the same slide never shares them
        tmp_preview, tmp_thumb = target / f".preview.{tag}.tmp", target / f".thumb.{tag}.tmp"
        image.save(tmp_preview, "PNG", optimize=True)
        thumb = image.resize((THUMB_WIDTH, round(image.height * THUMB_WIDTH / image.width)))
        thumb.save(tmp_thumb, "PNG", optimize=True)
        tmp_thumb.replace(target / "thumb.png")
        tmp_preview.replace(target / "preview.png")  # last: its presence means the slide is done


def missing(data: bytes, cache: Path) -> list[dict]:
    return [k for k in keys(data) if not (cache / k["key"] / "preview.png").exists()]


async def ensure(data: bytes, cache: Path, only: set[int] | None = None) -> list[dict]:
    """Render the slides (all, or the ids in `only`) whose images are not in `cache` yet; returns [{id, index, key}]
    for every slide."""
    all_keys = await asyncio.to_thread(keys, data)
    while True:
        todo = [k for k in all_keys if (only is None or k["id"] in only) and not (cache / k["key"] / "preview.png").exists()]
        if not todo:
            return all_keys
        mine = [k for k in todo if (str(cache), k["key"]) not in _inflight]
        if not mine:  # every one is being made by another request: wait for those, then look again
            await asyncio.gather(*{id(f): f for f in (_inflight[str(cache), k["key"]] for k in todo)}.values())
            continue
        done = asyncio.get_running_loop().create_future()
        for k in mine:
            _inflight[str(cache), k["key"]] = done
        try:
            async with _gate:
                subset = await asyncio.to_thread(_only, data, {k["id"] for k in mine})
                pdf = await asyncio.to_thread(_convert, subset, len(mine))
                await asyncio.to_thread(_rasterise, pdf, [cache / k["key"] for k in mine])
                log.info("rendered slides", extra={"slideCount": len(mine)})
        finally:  # those waiting look again: made, or (this render failed) theirs to make
            for k in mine:
                _inflight.pop((str(cache), k["key"]), None)
            done.set_result(None)


def soffice_available() -> bool:
    return shutil.which("soffice") is not None
